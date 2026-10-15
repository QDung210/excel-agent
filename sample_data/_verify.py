"""Runs every sample through the real API in a throw-away data dir and prints what was filled.

    .venv-ws/Scripts/python.exe sample_data/_verify.py          # mock mode (no network)
    .venv-ws/Scripts/python.exe sample_data/_verify.py azure    # uses AZURE_OPENAI_* from the environment
"""
from __future__ import annotations

import io
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent))
os.environ["WORKSPACE_DATA_DIR"] = tempfile.mkdtemp(prefix="ws-sample-")
os.environ["WORKSPACE_LLM"] = "azure" if sys.argv[1:2] == ["azure"] else "mock"

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from workspace import runs  # noqa: E402
from workspace.app import create_app  # noqa: E402

ANSWERS = "tra_loi_cau_hoi.txt"            # optional answer to the clarification question
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def template_case(c: TestClient, folder: str) -> None:
    d = ROOT / folder
    ws = c.post("/api/workspaces", json={"mode": "template"}).json()["id"]
    tpl = next(d.glob("template_*.xlsx"))
    assert c.post(f"/api/workspaces/{ws}/files", files={"file": (tpl.name, tpl.read_bytes(), XLSX)},
                  data={"role": "template"}).status_code == 200
    for src in sorted(p for p in d.iterdir() if p.suffix in (".csv", ".txt") and p.name not in ("prompt.txt", ANSWERS)):
        r = c.post(f"/api/workspaces/{ws}/files", files={"file": (src.name, src.read_bytes(), "text/plain")},
                   data={"role": "source"})
        assert r.status_code == 200, r.text
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": (d / "prompt.txt").read_text("utf-8")}).json()
    for m in r.get("mapping") or []:
        print(f"  map {m['target']}: {m['field']} <- {m['source'] or '-'} = {m['value']!r} [{m['status']}/{m['via']}]")
    if r["status"] == "needs_clarification":
        qs = r.get("questions", [])
        print("  questions:", [q.get("question") or q.get("id") for q in qs])
        # map_* questions: accept the AI's suggestion (the default), as a user clicking "next" would
        answers = {q["id"]: q.get("default") for q in qs if q["id"].startswith("map_")}
        if (d / ANSWERS).exists():
            answers["figures"] = (d / ANSWERS).read_text("utf-8")
        print("  answers:", answers or "(để trống)")
        r = c.post(f"/api/workspaces/{ws}/clarify", json={"clarification_id": r["clarification_id"],
                                                           "answers": answers}).json()
    if r["status"] != "ready":
        print("  status:", r["status"], r)
        return
    run = runs.wait(c.post(f"/api/workspaces/{ws}/generate", json={}).json()["run_id"])
    if run.status != "done":
        print("  run failed:", run.events[-1])
        return
    res = run.result
    wb = load_workbook(io.BytesIO(c.get(f"/api/workbooks/{res['workbook_id']}/export").content))
    filled = []
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for cell in row:
                if cell.value is not None and not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    filled.append(f"{cell.coordinate}={cell.value!r}")
    print(f"  workbook {res['workbook_id']} revision {res.get('revision')}")
    for w in res.get("warnings", []):
        print("  WARN:", w if isinstance(w, str) else w.get("message", w))
    print(f"  {len(filled)} non-formula cells, e.g.:", ", ".join(filled[:40]))


def chat_case(c: TestClient) -> None:
    p = ROOT / "05_chat_agent" / "theo_doi_ban_hang.xlsx"
    wid = c.post("/api/workbooks/upload", files={"file": (p.name, p.read_bytes(), XLSX)}).json()["id"]
    rev, thread = 1, None
    for msg in ["tóm tắt workbook", "đọc A3:E9", "đặt B4 = 15", "xóa A20:F30", "đồng ý",
                "điền doanh thu tháng 10", "thêm công thức tổng cho cột B", "kiểm tra workbook"]:
        body = {"workbook_id": wid, "revision": rev, "message": msg, "active_sheet": "Bán hàng"}
        if thread:
            body["thread_id"] = thread
        started = c.post("/api/agent/turns", json=body).json()
        thread = started["thread_id"]
        run = runs.wait(started["turn_id"])
        done = next(e["data"] for e in run.events if e["event"] == "done")
        rev = done["revision"]
        print(f"  > {msg}\n    {done['text'][:160]}  [revision {rev}]")


if __name__ == "__main__":
    c = TestClient(create_app())
    print("mode:", os.environ["WORKSPACE_LLM"], "| data dir:", os.environ["WORKSPACE_DATA_DIR"])
    for folder in sys.argv[2:] or ["01_bang_luong", "02_bao_gia", "03_bao_cao_doanh_thu", "04_phieu_nhap_kho",
                                  "06_so_ban_hang_nam", "07_luong_thue_tncn"]:
        print(f"\n== {folder}")
        template_case(c, folder)
    print("\n== 05_chat_agent")
    chat_case(c)
