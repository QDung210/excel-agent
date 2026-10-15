"""Stage 1-2: parse workbook with ks-xlsx-parser, then group constant cells into regions.

A region = constant (non-formula) cells in the same parser chunk, same column, same
parser annotation (label/data), same value type and same style (bold/fill). The LLM
labels regions by id, so it can never invent a range.
"""
from __future__ import annotations

import logging
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from html import unescape
from typing import Any

from openpyxl.utils import get_column_letter, range_boundaries

from . import ooxml

logging.getLogger("excel_parser").setLevel(logging.WARNING)

from excel_parser import parse_workbook  # noqa: E402
from excel_parser.analysis.header_detector import find_header_span  # noqa: E402


@dataclass
class Region:
    id: int
    sheet: str
    chunk: int
    col: int
    cells: list[tuple[int, int]]          # (row, col), sorted
    annotation: str                        # label | data | ...
    vtype: str                             # text | number | date | bool
    bold: bool
    fill: str | None
    in_header: bool                        # inside the parser's detected header band
    samples: list[str]
    header_above: str = ""
    label_left: str = ""
    refs: int = 0                          # formulas reading these cells
    refs_cross_sheet: int = 0
    dropdown_source: bool = False
    hidden_sheet: bool = False
    sequence: str = ""                     # e.g. "1..13" or "S1..S7" when values form a series

    @property
    def range_text(self) -> str:
        rows = {r for r, _ in self.cells}
        if len(rows) == 1:                 # one row, several columns (merged run)
            cols = sorted(c for _, c in self.cells)
            a, b = get_column_letter(cols[0]), get_column_letter(cols[-1])
            return f"{a}{min(rows)}" if a == b else f"{a}{min(rows)}:{b}{min(rows)}"
        return f"{get_column_letter(self.col)}[{compress_rows(list(rows))}]"

    def bbox(self) -> tuple[int, int, int, int]:
        rows = [r for r, _ in self.cells]
        cols = [c for _, c in self.cells]
        return min(rows), min(cols), max(rows), max(cols)


@dataclass
class Chunk:
    idx: int
    sheet: str
    block_type: str
    r1: int
    c1: int
    r2: int
    c2: int
    preview: str
    header_rows: tuple[int, int] | None = None

    @property
    def a1(self) -> str:
        return f"{get_column_letter(self.c1)}{self.r1}:{get_column_letter(self.c2)}{self.r2}"


@dataclass
class ParsedWorkbook:
    path: str
    sheets: list[str]
    hidden_sheets: set[str]
    chunks: list[Chunk]
    regions: list[Region]
    formula_count: dict[str, int] = field(default_factory=dict)
    parse_ms: float = 0.0


def compress_rows(rows: list[int]) -> str:
    rows = sorted(set(rows))
    out, start, prev = [], None, None
    for r in rows:
        if start is None:
            start = prev = r
        elif r == prev + 1:
            prev = r
        else:
            out.append(f"{start}" if start == prev else f"{start}-{prev}")
            start = prev = r
    if start is not None:
        out.append(f"{start}" if start == prev else f"{start}-{prev}")
    return ",".join(out)


def _vtype(cell) -> str:
    dt = getattr(cell, "data_type", "") or ""
    if dt == "n":
        return "number"
    if dt == "d":
        return "date"
    if dt == "b":
        return "bool"
    return "text"


def _style(cell) -> tuple[bool, str | None]:
    st = getattr(cell, "style", None)
    if st is None:
        return False, None
    font = getattr(st, "font", None)
    fill = getattr(st, "fill", None)
    bold = bool(getattr(font, "bold", False)) if font is not None else False
    color = None
    if fill is not None:
        color = getattr(fill, "fg_color", None) or getattr(fill, "bg_color", None)
        if getattr(fill, "pattern_type", None) in (None, "none"):
            color = None
    return bold, color


def _is_constant(cell) -> bool:
    if cell is None or cell.formula or cell.is_merged_slave:
        return False
    v = cell.raw_value
    return v is not None and not (isinstance(v, str) and v.strip() == "")


def _ann(cell) -> str:
    a = cell.annotation
    return getattr(a, "value", str(a)) if a is not None else "none"


def _short(v: Any, n: int = 40) -> str:
    s = str(v).replace("\n", " ").strip()
    if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}T00:00:00", s):
        s = s[:10]
    return s if len(s) <= n else s[: n - 1] + "…"


def dropdown_sources(z: zipfile.ZipFile) -> dict[str, list[tuple[int, int, int, int]]]:
    """Ranges used as list sources by data validation (standard + x14 extension).

    Read from raw XML because openpyxl drops the x14 extension.
    """
    out: dict[str, list[tuple[int, int, int, int]]] = defaultdict(list)
    for name, part in ooxml.sheet_parts(z).items():
        try:
            xml = z.read(part).decode("utf-8")
        except KeyError:
            continue
        formulas = re.findall(r"<formula1>(.*?)</formula1>", xml, re.S)
        formulas += re.findall(r"<x14:formula1>\s*<xm:f>(.*?)</xm:f>", xml, re.S)
        for f in formulas:
            f = unescape(f).strip().lstrip("=")
            m = re.fullmatch(r"(?:'?([^'!]+)'?!)?(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)", f)
            if not m:
                continue
            sh = m.group(1) or name
            c1, r1, c2, r2 = range_boundaries(m.group(2).replace("$", ""))
            out[sh].append((r1, c1, r2, c2))
    return out


def _intersects(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> bool:
    return not (a[2] < b[0] or a[0] > b[2] or a[3] < b[1] or a[1] > b[3])


def build(path: str, preview_lines: int = 6, preview_width: int = 260) -> ParsedWorkbook:
    # Whole-table chunks (no token windowing): every table is one chunk, so the
    # header and the body end up in the same chunk.
    res = parse_workbook(path=path, max_chunk_tokens=None)
    wb = res.workbook
    with zipfile.ZipFile(path) as z:
        hidden = ooxml.hidden_sheets(z)
        dv = dropdown_sources(z)
    sheet_by_name = {s.sheet_name: s for s in wb.sheets}

    chunks: list[Chunk] = []
    for i, c in enumerate(res.chunks):
        c1, r1, c2, r2 = range_boundaries(f"{c.top_left_cell}:{c.bottom_right_cell}")
        lines = [ln[:preview_width] for ln in (c.render_text or "").splitlines() if ln.strip() and not set(ln) <= set("|-: ")]
        bt = c.block_type if isinstance(c.block_type, str) else c.block_type.value
        hs = find_header_span(sheet_by_name[c.sheet_name], c.cell_range) if c.cell_range else None
        chunks.append(Chunk(i, c.sheet_name, bt, r1, c1, r2, c2, "\n".join(lines[:preview_lines]),
                            (hs.top, hs.bottom) if hs else None))

    # formula edges -> target bounding boxes, per target sheet
    edges: dict[str, list[tuple[tuple[int, int, int, int], bool]]] = defaultdict(list)
    for e in wb.dependency_graph.edges:
        tsheet = e.target_sheet or e.source_sheet
        if e.target_range is not None:
            tr = e.target_range
            box = (tr.top_left.row, tr.top_left.col, tr.bottom_right.row, tr.bottom_right.col)
        elif e.target_coord is not None:
            box = (e.target_coord.row, e.target_coord.col, e.target_coord.row, e.target_coord.col)
        else:
            continue
        edges[tsheet].append((box, tsheet != e.source_sheet))

    groups: dict[tuple, list] = defaultdict(list)
    formula_count: dict[str, int] = {}
    for sheet in wb.sheets:
        name = sheet.sheet_name
        sheet_chunks = [ch for ch in chunks if ch.sheet == name]
        formula_count[name] = sum(1 for c in sheet.cells.values() if c.formula)
        for cell in sheet.cells.values():
            if not _is_constant(cell):
                continue
            r, c = cell.coord.row, cell.coord.col
            k = next((k for k in sheet_chunks if k.r1 <= r <= k.r2 and k.c1 <= c <= k.c2), None)
            ch = k.idx if k else -1
            in_header = bool(k and k.header_rows and k.header_rows[0] <= r <= k.header_rows[1])
            bold, fill = _style(cell)
            key = (name, ch, c, _ann(cell), _vtype(cell), bold, fill, in_header)
            groups[key].append(cell)

    groups = _split_key_value(groups, sheet_by_name)
    items = _merge_rows(groups)

    regions: list[Region] = []
    order = {s.sheet_name: i for i, s in enumerate(wb.sheets)}
    for key, cells in sorted(items, key=lambda it: (order[it[0][0]], it[0][1], it[0][2], min(c.coord.row for c in it[1]))):
        name, ch, col, ann, vtype, bold, fill, in_header = key[:8]
        cells = sorted(cells, key=lambda c: (c.coord.row, c.coord.col))
        seen, samples = set(), []
        for c in cells:
            s = _short(c.raw_value)
            if s not in seen:
                seen.add(s)
                samples.append(s)
            if len(samples) >= 5:
                break
        reg = Region(
            id=len(regions), sheet=name, chunk=ch, col=col,
            cells=[(c.coord.row, c.coord.col) for c in cells],
            annotation=ann, vtype=vtype, bold=bold, fill=fill, in_header=in_header, samples=samples,
            hidden_sheet=name in hidden, sequence=_sequence([c.raw_value for c in cells]),
        )
        _context(reg, sheet_by_name[name], chunks[ch] if ch >= 0 else None)
        box = reg.bbox()
        for ebox, cross in edges.get(name, []):
            if _intersects(box, ebox):
                reg.refs += 1
                reg.refs_cross_sheet += cross
        reg.dropdown_source = any(_intersects(box, b) for b in dv.get(name, []))
        regions.append(reg)

    return ParsedWorkbook(
        path=path, sheets=[s.sheet_name for s in wb.sheets], hidden_sheets=hidden,
        chunks=chunks, regions=regions, formula_count=formula_count,
        parse_ms=wb.parse_duration_ms or 0.0,
    )


def _split_key_value(groups: dict[tuple, list], sheets: dict) -> dict[tuple, list]:
    """A column of numbers/dates where every row has its own distinct text label on the
    left is a key/value parameter table (e.g. 'Hours per MD | 8'): each value gets its own
    region so the LLM can keep generic settings and clear instance-specific ones."""
    out: dict[tuple, list] = {}
    for key, cells in groups.items():
        name, ch, col, ann, vtype, *rest = key
        if vtype in ("number", "date") and 2 <= len(cells) <= 30 and col > 1:
            sheet = sheets[name]
            labels = []
            for c in cells:
                left = sheet.get_cell(c.coord.row, col - 1)
                labels.append(left.raw_value if left is not None and _is_constant(left)
                              and _vtype(left) == "text" else None)
            if all(labels) and len(set(labels)) == len(labels):
                for c in cells:
                    out[(name, ch, col, ann, vtype, *rest, c.coord.row)] = [c]
                continue
        out[key] = cells
    return out


def _merge_rows(groups: dict[tuple, list]) -> list[tuple[tuple, list]]:
    """Single-cell groups that sit side by side on one row with the same key become one
    region (e.g. a header row, or week numbers 1..13 across a timeline)."""
    items: list[tuple[tuple, list]] = []
    buckets: dict[tuple, list[tuple[tuple, list]]] = defaultdict(list)
    for key, cells in groups.items():
        if len(cells) == 1:
            name, ch, col, *rest = key
            buckets[(name, ch, cells[0].coord.row, *rest)].append((key, cells))
        else:
            items.append((key, cells))
    for run_items in buckets.values():
        run_items.sort(key=lambda it: it[0][2])
        run = [run_items[0]]
        for it in run_items[1:]:
            if it[0][2] == run[-1][0][2] + 1:        # strictly adjacent columns
                run.append(it)
            else:
                items.append((run[0][0], [c for _, cs in run for c in cs]))
                run = [it]
        items.append((run[0][0], [c for _, cs in run for c in cs]))
    return items


def _sequence(values: list[Any]) -> str:
    """'1..13' / 'S1..S7' when the values are a consecutive integer series (len >= 3)."""
    if len(values) < 3:
        return ""
    prefix, nums = None, []
    for v in values:
        if isinstance(v, bool):
            return ""
        if isinstance(v, (int, float)) and float(v).is_integer():
            p, n = "", int(v)
        else:
            m = re.fullmatch(r"([A-Za-z#\- ]*)(\d+)", str(v).strip())
            if not m:
                return ""
            p, n = m.group(1), int(m.group(2))
        if prefix is not None and p != prefix:
            return ""
        prefix = p
        nums.append(n)
    if sorted(nums) != list(range(min(nums), min(nums) + len(nums))):
        return ""
    return f"{prefix}{min(nums)}..{prefix}{max(nums)}"


def _context(reg: Region, sheet, chunk: Chunk | None) -> None:
    """Nearest label above (same column) and to the left (same row) of the first cell."""
    r0, c = reg.cells[0]
    top = chunk.r1 if chunk else 1
    for r in range(r0 - 1, max(top, r0 - 40) - 1, -1):
        x = sheet.get_cell(r, c)
        if x is not None and _is_constant(x) and _vtype(x) == "text" and (_ann(x) == "label" or _style(x)[0]):
            reg.header_above = _short(x.raw_value, 30)
            break
    for cc in range(c - 1, max(0, c - 4), -1):
        x = sheet.get_cell(r0, cc)
        if x is not None and _is_constant(x) and _vtype(x) == "text":
            # numbers/dates: the text on the left names the value (key/value tables)
            if _ann(x) == "label" or _style(x)[0] or reg.vtype in ("number", "date"):
                reg.label_left = _short(x.raw_value, 30)
            break
