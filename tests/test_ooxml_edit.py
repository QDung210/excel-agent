"""Value edits on existing workbooks keep everything else intact (spec 07 L0-L4)."""
from __future__ import annotations

import zipfile

import pytest
from openpyxl import load_workbook

from workspace.errors import InvalidOperation
from workspace.ooxml_edit import apply_actions
from workspace.schemas import ClearValue, SetValue
from workspace.validate_edit import validate


def sv(sheet, cell, value, **kw):
    return SetValue(type="set_value", sheet=sheet, cell=cell, value=value, **kw)


def _members(p):
    with zipfile.ZipFile(p) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_set_values_validate_and_preserve(rich_template, tmp_path):
    dst = tmp_path / "out.xlsx"
    st = apply_actions(str(rich_template), str(dst), [sv("Report", "A3", "Task A"), sv("Report", "C3", 4),
                                                       sv("Report", "B3", "Open"), sv("Report", "F20", "new cell")])
    rep = validate(str(rich_template), str(dst), st)
    assert rep["pass"], rep["errors"]
    a, b = _members(rich_template), _members(dst)
    assert set(a) == set(b)
    for n in a:
        if n not in st.touched_parts:
            assert a[n] == b[n], n
    assert st.touched_parts <= {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml", "xl/sharedStrings.xml",
                                "xl/workbook.xml"}
    ws = load_workbook(dst)["Report"]
    assert ws["A3"].value == "Task A" and ws["C3"].value == 4 and ws["F20"].value == "new cell"
    assert ws["D3"].value == "=C3*25"                  # formulas untouched
    orig = load_workbook(rich_template)["Report"]
    assert ws["A3"].style_id == orig["A3"].style_id    # style kept


def test_x14_extension_survives_edit(x14_template, tmp_path):
    dst = tmp_path / "out.xlsx"
    st = apply_actions(str(x14_template), str(dst), [sv("Report", "B4", "Done")])
    assert validate(str(x14_template), str(dst), st)["pass"]
    with zipfile.ZipFile(dst) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
    assert "x14:dataValidation" in xml and "Lists!$A$1:$A$3" in xml


def test_formula_target_rejected(rich_template, tmp_path):
    with pytest.raises(InvalidOperation, match="formula"):
        apply_actions(str(rich_template), str(tmp_path / "o.xlsx"), [sv("Report", "D3", 5)])


def test_merged_non_anchor_rejected(rich_template, tmp_path):
    with pytest.raises(InvalidOperation, match="merged"):
        apply_actions(str(rich_template), str(tmp_path / "o.xlsx"), [sv("Report", "B1", "x")])


def test_unknown_sheet_rejected(rich_template, tmp_path):
    with pytest.raises(InvalidOperation, match="does not exist"):
        apply_actions(str(rich_template), str(tmp_path / "o.xlsx"), [sv("Nope", "A1", 1)])


def test_equals_text_stored_as_literal_not_formula(simple_sales, tmp_path):
    dst = tmp_path / "o.xlsx"
    st = apply_actions(str(simple_sales), str(dst), [sv("Sales", "A12", '=HYPERLINK("http://x")')])
    assert validate(str(simple_sales), str(dst), st)["pass"]
    c = load_workbook(dst)["Sales"]["A12"]
    assert c.data_type == "s" and c.value.startswith("=HYPERLINK")


def test_clear_and_types(simple_sales, tmp_path):
    dst = tmp_path / "o.xlsx"
    acts = [sv("Sales", "B3", "21"), ClearValue(type="clear_value", sheet="Sales", cell="A5"),
            sv("Sales", "A10", True), sv("Sales", "B10", "2026-10-01", value_type="date"),
            sv("Sales", "C10", "007")]
    st = apply_actions(str(simple_sales), str(dst), acts)
    rep = validate(str(simple_sales), str(dst), st)
    assert rep["pass"], rep["errors"]
    ws = load_workbook(dst)["Sales"]
    assert ws["B3"].value == 21 and ws["A5"].value is None and ws["A10"].value is True
    assert ws["C10"].value == "007"                  # leading zeros stay text
    assert ws["B10"].value == 46296                  # date serial (cell has no date format)
    assert st.formula_cache_stripped >= 1            # dependent formula caches invalidated


def test_dependent_sheet_caches_stripped(rich_template, tmp_path):
    dst = tmp_path / "o.xlsx"
    st = apply_actions(str(rich_template), str(dst), [sv("Report", "C4", 3)])
    assert "xl/worksheets/sheet2.xml" in st.touched_parts      # Summary references Report
    assert "xl/worksheets/sheet3.xml" not in st.touched_parts  # Lists has no formulas


def test_validator_catches_untargeted_change(simple_sales, tmp_path):
    dst = tmp_path / "o.xlsx"
    st = apply_actions(str(simple_sales), str(dst), [sv("Sales", "A10", "x")])
    st.intended.clear()
    st.intended[("Sales", "A11")] = "y"            # lie about the target: L3 must fail
    assert not validate(str(simple_sales), str(dst), st)["pass"]


def test_real_sample_edit(wbs, tmp_path):
    dst = tmp_path / "o.xlsx"
    st = apply_actions(str(wbs), str(dst), [sv("WBS", "A2", "Unit: person-day"), sv("WBS", "B200", 1.5)])
    rep = validate(str(wbs), str(dst), st)
    assert rep["pass"], rep["errors"]
    assert rep["checks"]["L2_formulas_checked"] > 1000
    with zipfile.ZipFile(wbs) as a, zipfile.ZipFile(dst) as b:
        for n in a.namelist():
            if "x14:dataValidation" in a.read(n).decode("utf-8", "ignore"):
                assert "x14:dataValidation" in b.read(n).decode("utf-8", "ignore")


def test_stripped_cache_leaves_no_text_in_numeric_cells(wbs, tmp_path):
    """Regression: Resource!L26 had <v xml:space="preserve">text</v>; t="str" was dropped but the text kept."""
    import re
    dst = tmp_path / "o.xlsx"
    apply_actions(str(wbs), str(dst), [sv("WBS", "E8", "QA manual edit")])
    from xlsx_template import ooxml
    with zipfile.ZipFile(dst) as z:
        for n in z.namelist():
            if not n.startswith("xl/worksheets/sheet"):
                continue
            _, inner, _ = ooxml.split_sheet_data(z.read(n).decode("utf-8"))
            for m in ooxml.CELL_RE.finditer(inner):
                a, body = ooxml.attrs(m.group(1)), m.group(2) or ""
                v = re.search(r"<v(?:\s[^>]*)?>(.*?)</v>", body, re.S)
                if v and a.get("t", "n") == "n":
                    float(v.group(1))                     # raises if a numeric cell carries text
    from openpyxl import load_workbook
    assert load_workbook(dst, data_only=True)["Daily Report"]["L26"].value is None


def test_explicit_style_zero_is_kept(tmp_path):
    """Cells written as s="0" (common in exported forms) keep the attribute, so validation passes."""
    import re
    from openpyxl import Workbook
    wb = Workbook()
    wb.active["A1"], wb.active["B1"] = "N", 1
    raw, src = tmp_path / "raw.xlsx", tmp_path / "s0.xlsx"
    wb.save(raw)
    with zipfile.ZipFile(raw) as zi, zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in zi.namelist():
            data = zi.read(n)
            if n == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<c r="(A1|B1)"(?![^>]*\bs=)', rb'<c r="\1" s="0"', data)
                assert data.count(b's="0"') == 2
            zo.writestr(n, data)
    dst = tmp_path / "out.xlsx"
    st = apply_actions(str(src), str(dst), [sv("Sheet", "A1", "Y"), sv("Sheet", "B1", 2)])
    rep = validate(str(src), str(dst), st)
    assert rep["pass"], rep["errors"]
    assert load_workbook(dst)["Sheet"]["A1"].value == "Y"
