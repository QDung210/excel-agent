"""Phase 2: analyze -> clarify -> generate (scratch). Core rule under test: no invented numbers."""
from __future__ import annotations

import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from workspace import config, runs, sources, spec_gen
from workspace.app import create_app
from workspace.formula_eval import Workbook, UNSUPPORTED, evaluate

SALES_CSV = "Sản phẩm,Số lượng,Đơn giá\nBút bi,10,\"5.000\"\nVở,4,12000\n".encode()
NOTES_TXT = "Doanh thu tháng 9/2026: 150.000.000 đ\nChi phí: 90 triệu\nGhi chú: khách hàng mới tăng\n".encode()


def _ws(c: TestClient, mode: str = "scratch", files: list[tuple[str, bytes]] = ()) -> str:
    ws = c.post("/api/workspaces", json={"mode": mode}).json()["id"]
    for name, data in files:
        r = c.post(f"/api/workspaces/{ws}/files", files={"file": (name, data, "text/plain")}, data={"role": "source"})
        assert r.status_code == 200, r.text
    return ws


def _generate(c: TestClient, ws: str) -> tuple[runs.Run, bytes]:
    rid = c.post(f"/api/workspaces/{ws}/generate", json={}).json()["run_id"]
    run = runs.wait(rid)
    assert run.status == "done", [e for e in run.events if e["event"] in ("error", "validation_failed")]
    wid = run.result["workbook_id"]
    data = c.get(f"/api/workbooks/{wid}/export").content
    return run, data


def _numbers(xlsx: bytes) -> list[float]:
    """Every numeric constant cell (formulas excluded)."""
    wb = load_workbook(io.BytesIO(xlsx))
    out = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
           if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
    wb.close()
    return out


@pytest.mark.parametrize("text,value,unit", [
    ("15.000.000 đ", 15_000_000, "VND"), ("1,5", 1.5, None), ("12%", 0.12, "%"), ("2 tỷ", 2e9, "VND"),
    ("1,234.56", 1234.56, None), ("$1,200", 1200, "USD"), ("1.234,5", 1234.5, None), ("abc", None, None),
    ("1.23.4", None, None), ("15/09/2026", None, None), ("Quý 3/2026", None, None),
    ("1 234 567", 1234567, None), ("0100000000", None, None), ("0901 234 567", None, None), ("0", 0, None),
    ("0,5", 0.5, None)])
def test_parse_number(text, value, unit):
    v, u = sources.parse_number(text)
    assert (v == pytest.approx(value) if value is not None else v is None) and (u == unit or value is None)


def test_generate_from_sources_uses_only_source_numbers():
    c = TestClient(create_app())
    ws = _ws(c, files=[("ban_hang.csv", SALES_CSV), ("notes.txt", NOTES_TXT)])
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Tạo báo cáo bán hàng tháng 9 kèm biểu đồ"}).json()
    assert r["status"] == "ready" and r["facts"] == 7
    assert any("Biểu đồ" in w for w in r["warnings"])                 # unsupported feature is announced
    run, data = _generate(c, ws)
    names = [e["event"] for e in run.events]
    assert names[:2] == ["run_started", "analysis_ready"] and names[-2:] == ["workbook_revision", "done"]
    assert [e["data"]["tool"] for e in run.events if e["event"] == "tool_completed"] == \
        ["read_sources", "build_spec", "build_workbook", "validate_workbook"]
    assert set(_numbers(data)) <= {10, 4, 5000, 12000, 150_000_000, 90_000_000}
    wb = load_workbook(io.BytesIO(data))
    ws1 = wb.worksheets[0]
    assert ws1["D5"].value == "=B5*C5" and ws1["D7"].value == "=SUM(D5:D6)"
    assert "Thành tiền" in [c.value for c in ws1[4]]
    assert wb.worksheets[1].title == "Tổng hợp"
    cached = load_workbook(io.BytesIO(data), data_only=True).worksheets[0]
    assert cached["D5"].value == 50000 and cached["D7"].value == 98000
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert b"fullCalcOnLoad" in z.read("xl/workbook.xml")
    meta = c.get(f"/api/workbooks/{run.result['workbook_id']}/metadata").json()
    assert meta["origin"] == "generated" and meta["workspace_id"] == ws and meta["llm_mode"] == "mock"
    assert c.get(f"/api/workspaces/{ws}").json()["workbook_id"] == meta["id"]


def test_no_numbers_asks_then_leaves_inputs_blank():
    c = TestClient(create_app())
    ws = _ws(c)
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Tạo báo cáo doanh thu theo tháng"}).json()
    assert r["status"] == "needs_clarification"
    ids = {q["id"] for q in r["questions"]}
    assert {"currency", "missing_numbers"} <= ids and len(ids) <= 5
    assert c.post(f"/api/workspaces/{ws}/generate", json={}).status_code == 422     # must answer first
    bad = c.post(f"/api/workspaces/{ws}/clarify", json={"clarification_id": r["clarification_id"],
                                                         "answers": {"currency": "JPY"}})
    assert bad.status_code == 422
    r = c.post(f"/api/workspaces/{ws}/clarify", json={"clarification_id": r["clarification_id"],
                                                       "answers": {"currency": "VND", "missing_numbers": "blank"}})
    assert r.json()["status"] == "ready"
    run, data = _generate(c, ws)
    assert _numbers(data) == []                                      # nothing invented
    assert any("blank" in w.lower() or "trống" in w for w in run.result["warnings"])


def test_user_figures_are_used_with_provenance():
    c = TestClient(create_app())
    ws = _ws(c)
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Lập báo cáo doanh thu quý 3"}).json()
    r = c.post(f"/api/workspaces/{ws}/clarify", json={
        "clarification_id": r["clarification_id"],
        "answers": {"currency": "VND", "missing_numbers": "provide",
                    "figures": "Doanh thu tháng 7: 120.000.000\nDoanh thu tháng 8: 135.000.000"}}).json()
    run, data = _generate(c, ws)
    assert sorted(_numbers(data)) == [120_000_000, 135_000_000]
    wb = load_workbook(io.BytesIO(data))
    sources_col = [c.value for row in wb.worksheets[0].iter_rows() for c in row
                   if isinstance(c.value, str) and "answer:figures" in c.value]
    assert len(sources_col) == 2


def test_blocked_states():
    c = TestClient(create_app())
    ws = _ws(c, mode="template")
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Điền mẫu"}).json()
    assert r["status"] == "blocked" and "template" in r["reason"].lower()
    ws2 = _ws(c)
    assert c.post(f"/api/workspaces/{ws2}/analyze", json={"prompt": " "}).json()["status"] == "blocked"
    assert c.post(f"/api/workspaces/{ws2}/analyze",
                  json={"prompt": "x", "source_file_ids": ["file_0000000000000000"]}).status_code == 422


def test_formula_guard():
    allowed = {0.08}
    keys = {"qty", "price"}
    assert spec_gen.check_formula("={qty}*{price}", keys, allowed) is None
    assert spec_gen.check_formula("=ROUND({qty}*{price},2)", keys, allowed) is None
    assert spec_gen.check_formula("={price}*0.08", keys, allowed) is None              # grounded constant
    assert "ungrounded" in spec_gen.check_formula("={price}*1.15", keys, allowed)      # invented rate
    assert "unsupported" in spec_gen.check_formula("=VLOOKUP({qty},{price},2)", keys, allowed)
    assert "unknown" in spec_gen.check_formula("={qty}*{tax}", keys, allowed)
    assert "references" in spec_gen.check_formula("={qty}*Sheet2!B3", keys, allowed)


def test_llm_layout_never_types_data(monkeypatch):
    b = sources.SourceBundle()
    sources._read_csv(b, "file_x", "sales.csv", SALES_CSV.decode())
    layout = spec_gen.LlmLayout(title="Sales", sheets=[spec_gen.LlmSheet(
        name="Sales", source_table=0, total_headers=["Amount", "Product"],
        columns=[spec_gen.LlmColumn(header="Product", source_header="Sản phẩm", type="text"),
                 spec_gen.LlmColumn(header="Qty", source_header="Số lượng", type="integer"),
                 spec_gen.LlmColumn(header="Price", source_header="Đơn giá", type="currency"),
                 spec_gen.LlmColumn(header="Amount", type="currency", formula="={Qty}*{Price}"),
                 spec_gen.LlmColumn(header="With VAT", type="currency", formula="={Amount}*1.1"),
                 spec_gen.LlmColumn(header="Region", source_header="Khu vực", type="text")])])
    spec, warns = spec_gen.from_layout(layout, b, "sales", {})
    s = spec.sheets[0]
    assert [r["qty"] for r in s.rows] == [10, 4] and s.totals == ["amount"]
    assert next(c for c in s.columns if c.key == "amount").formula == "={qty}*{price}"
    assert next(c for c in s.columns if c.key == "with_vat").formula is None
    assert any("1.1" in w for w in warns) and any("Khu vực" in w for w in warns)

    monkeypatch.setattr(config, "llm_mode", lambda: "azure")
    monkeypatch.setattr(spec_gen, "llm_layout", lambda *a: (_ for _ in ()).throw(RuntimeError("down")))
    spec, warns, mode = spec_gen.build_spec(b, "sales", {})
    assert mode == "mock_fallback" and "unavailable" in warns[0]


def test_ground_check_blanks_unknown_numbers():
    b = sources.SourceBundle()
    b.add("x", 5, None, "prompt")
    spec, _ = spec_gen.mock_spec(b, "t", {})
    spec.sheets[0].rows.append({"item": "fake", "value": 999})
    assert spec_gen.ground_check(spec, b) == ["1 value(s) without a source were removed."]
    assert spec.sheets[0].rows[-1]["value"] is None


def test_formula_eval_basics():
    wb = Workbook({"S": {(1, 1): 2, (1, 2): 3, (2, 1): "a"}}, {"S": {(1, 3): "=A1*B1", (1, 4): "=SUM(A1:C1)"}})
    assert wb.value("S", 1, 3) == 6 and wb.value("S", 1, 4) == 11
    assert evaluate("=IFERROR(1/0,\"x\")", "S", lambda *a: None) == "x"
    assert evaluate("=1/0", "S", lambda *a: None) == "#DIV/0!"
    assert evaluate("=VLOOKUP(1,A1:B2,2,FALSE)", "S", lambda *a: None) is UNSUPPORTED
    cyc = Workbook({}, {"S": {(1, 1): "=B1", (1, 2): "=A1"}})
    assert cyc.value("S", 1, 1) is UNSUPPORTED


REGION_CSV = "Sản phẩm,Số lượng,Đơn giá,Khu vực\nBút bi,10,5000,Bắc\nVở,4,12000,Nam\nThước,7,8000,Bắc\n"


def test_group_sheet_sumif_mock_and_layout():
    b = sources.SourceBundle()
    sources._read_csv(b, "file_x", "ban_hang.csv", REGION_CSV)
    spec, _ = spec_gen.mock_spec(b, "Tạo báo cáo bán hàng, tổng theo khu vực", {})
    names = [s.name for s in spec.sheets]
    assert len(names) == 2 and "Khu vực" in names[1]
    g = spec.sheets[1]
    assert [r[g.columns[0].key] for r in g.rows] == ["Bắc", "Nam"]
    assert all("SUMIF(" in f for f in g.formulas.values())
    assert "don_gia" not in [c.key for c in g.columns]                # unit price is never summed
    ev = __import__("workspace.builder", fromlist=["x"])
    wb = ev.compute(spec, ev.plan_cells(spec))
    L = ev._Layout(g)
    assert wb.value(g.name, L.first + 1, 2) == 17 and wb.value(g.name, L.first + 2, 2) == 4

    layout = spec_gen.LlmLayout(title="Sales", sheets=[
        spec_gen.LlmSheet(name="Detail", source_table=0, columns=[
            spec_gen.LlmColumn(header="Product", source_header="Sản phẩm", type="text"),
            spec_gen.LlmColumn(header="Qty", source_header="Số lượng", type="integer"),
            spec_gen.LlmColumn(header="Đơn giá", source_header="Đơn giá", type="currency"),
            spec_gen.LlmColumn(header="Thành tiền", type="currency", formula="={Qty}*{Đơn giá}")]),
        spec_gen.LlmSheet(name="By product", group_by_sheet="Detail", group_by_header="Product",
                          sum_headers=["Qty", "Đơn giá", "Thành tiền"]),
        spec_gen.LlmSheet(name="Bad", group_by_sheet="Detail", group_by_header="Nope")])
    spec, warns = spec_gen.from_layout(layout, b, "sales", {})
    assert [s.name for s in spec.sheets] == ["Detail", "By product"]
    assert any("Nope" in w for w in warns)
    keys = [c.key for c in spec.sheets[1].columns]
    assert "qty" in keys and "don_gia" not in keys and len(spec.sheets[1].rows) == 3
