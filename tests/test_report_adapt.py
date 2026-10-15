"""Adapting a template to a report described in prose (invoice -> monthly sales report by branch x product):
sequence column, renamed + reformulated columns, new columns, implied zeros, label edits, totals, an extra summary
sheet, the dry-run preview, the reorder safety block and the approval gate."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from workspace import formula_eval, sources, templates
from workspace import restructure as R
from workspace.errors import InvalidOperation

ROOT = Path(__file__).resolve().parent.parent
DOC = (ROOT / "bao_cao_kinh_doanh_thang_09_2026.txt").read_text(encoding="utf-8")
v = R.LlmValue

PRICES = {"LT14": ("15900000", "giá niêm yết 15,9 triệu một chiếc"), "MN27": ("3450000", "bán 3.450.000đ"),
          "KB87": ("1250000", "có giá bán 1 triệu 250 nghìn"), "MS02": ("390000", "đang để giá 390k")}
NAMES = {"LT14": ("laptop Ultrabook A14", "LT14 là laptop Ultrabook A14"),
         "MN27": ("màn hình 27 inch", "MN27 là màn hình 27 inch"),
         "KB87": ("Bàn phím cơ", "Bàn phím cơ mã KB87"), "MS02": ("Chuột không dây", "Chuột không dây MS02")}
HN = {"LT14": ("18", "18 laptop LT14", "0"), "MN27": ("30", "30 màn hình MN27", "1"),
      "KB87": ("42", "42 bàn phím KB87", "0"), "MS02": ("90", "90 chuột MS02", "3")}
HP = {"LT14": ("12", "12 chiếc LT14", "1"), "MN27": ("22", "22 chiếc MN27", "0"),
      "KB87": ("36", "36 chiếc KB87", "2"), "MS02": ("68", "68 chiếc MS02", "0")}
RET_Q = {("HN", "MN27"): "có một màn hình", ("HN", "MS02"): "ba chuột bị khách trả lại",
         ("HP", "LT14"): "Có một chiếc LT14 bị trả", ("HP", "KB87"): "hai bàn phím KB87 hoàn về"}
RET_ALL = {"HN": "có một màn hình và ba chuột bị khách trả lại", "HP": "Có một chiếc LT14 bị trả và hai bàn phím"}
CK = {"HN": ("1", "chương trình giảm giá 1%"), "HP": ("2", "áp dụng chiết khấu 2%")}
BR = {"HN": ("Hà Nội", "Về cửa hàng Hà Nội"), "HP": ("Hải Phòng", "Ở Hải Phòng")}


def _col(src, header, kind, fmt, formula=None, total=False):
    return R.LlmColumn(source_header=src, header=header, kind=kind, format=fmt, formula=formula, total=total,
                       reason="lý do")


def _rows() -> list[R.LlmRow]:
    out = []
    for br, data in (("HN", HN), ("HP", HP)):
        for code, (qty, qq, ret) in data.items():
            rq = RET_Q.get((br, code))
            cells = [v(header="Chi nhánh", value=BR[br][0], quote=BR[br][1]),
                     v(header="Mã hàng", value=code, quote=NAMES[code][1]),
                     v(header="Sản phẩm", value=NAMES[code][0], quote=NAMES[code][1]),
                     v(header="SL xuất", value=qty, quote=qq),
                     v(header="SL trả", value=ret, quote=rq) if rq else
                     v(header="SL trả", value="0", quote=RET_ALL[br], basis="implied",
                       note="không có trong danh sách hàng trả"),
                     v(header="Đơn giá", value=PRICES[code][0], quote=PRICES[code][1]),
                     v(header="CK (%)", value=CK[br][0], quote=CK[br][1])]
            out.append(R.LlmRow(cells=cells))
    return out


def answer(sheets: bool = True, order: list[str] | None = None) -> R.LlmStructure:
    cols = [_col("STT", "STT", "sequence", "number"),
            _col(None, "Chi nhánh", "input", "text"),
            _col(None, "Mã hàng", "input", "text"),
            _col("Sản phẩm", "Sản phẩm", "keep", "text"),
            _col(None, "SL xuất", "input", "number", total=True),
            _col(None, "SL trả", "input", "number", total=True),
            _col("Số lượng", "SL bán ròng", "formula", "number", "{SL xuất}-{SL trả}", total=True),
            _col("Đơn giá", "Đơn giá", "keep", "number"),
            _col("Thành tiền", "Thành tiền", "keep", "number", total=True),
            _col(None, "CK (%)", "input", "percent_number"),
            _col(None, "Tiền CK", "formula", "number", "{Thành tiền}*{CK (%)}/100", total=True),
            _col(None, "Doanh thu ròng", "formula", "number", "{Thành tiền}-{Tiền CK}", total=True)]
    if order:
        cols = sorted(cols, key=lambda c: order.index(c.header) if c.header in order else 99)
    sh = [R.LlmSheet(name="Tổng hợp chi nhánh", purpose="Doanh thu ròng theo chi nhánh", reason="tài liệu yêu cầu",
                     columns=[R.LlmSheetColumn(header="Chi nhánh", kind="input", format="text"),
                              R.LlmSheetColumn(header="Doanh thu ròng", kind="formula", total=True,
                                               formula="SUMIF({Hóa đơn!Chi nhánh},{Chi nhánh},"
                                                       "{Hóa đơn!Doanh thu ròng})")],
                     rows=[R.LlmRow(cells=[v(header="Chi nhánh", value="Hà Nội", quote="Về cửa hàng Hà Nội")]),
                           R.LlmRow(cells=[v(header="Chi nhánh", value="Hải Phòng", quote="Ở Hải Phòng")])])]
    und = R.LlmUnderstanding(template_purpose="hóa đơn một khách", fit="partial", missing=["thuế suất VAT"])
    labels = [R.LlmLabel(current="HÓA ĐƠN BÁN HÀNG", new="BÁO CÁO BÁN HÀNG THÁNG 9/2026", reason="báo cáo tháng"),
              R.LlmLabel(current="Khách hàng:", new="Đơn vị", reason="không có một khách hàng")]
    kv = [R.LlmKv(label="Khách hàng", value="Công ty TNHH Nam An Tech", quote="Công ty TNHH Nam An Tech")]
    return R.LlmStructure(understanding=und, columns=cols, labels=labels, kv=kv, rows=_rows(),
                          sheets=sh if sheets else [])


@pytest.fixture
def tpl(tmp_path) -> Path:
    p = tmp_path / "invoice.xlsx"
    templates._invoice(p)
    return p


def _propose(tpl, ans):
    b = sources.SourceBundle(texts={"f1": DOC}, names={"f1": "bao_cao.txt"})
    return R.propose(tpl, templates.detect_slots(tpl), b, "lập báo cáo từ file", llm=lambda ctx: (ans, "fake"))


def _values(path: Path, sheet: str) -> formula_eval.Workbook:
    cells, formulas = R._grid(path)
    return formula_eval.Workbook(cells, formulas)


def test_invoice_to_branch_report(tpl, tmp_path):
    plan = _propose(tpl, answer())
    assert plan["preview"]["ok"], plan["preview"]
    assert len(plan["data"]) == 8 and plan["overflow"] == 0
    # implied zeros are marked as assumed, stated values as explicit
    provs = {(r["Chi nhánh"]["value"], r["Mã hàng"]["value"]): r["SL trả"]["prov"] for r in plan["data"]}
    assert provs[("Hà Nội", "LT14")] == "assumed" and provs[("Hà Nội", "MN27")] == "explicit"
    out = tmp_path / "out.xlsx"
    R.apply_plan(tpl, out, plan)
    assert R.validate(tpl, out, plan)["pass"]
    ws = load_workbook(out)["Hóa đơn"]
    heads = [ws.cell(8, c).value for c in range(1, 13)]
    assert heads == ["STT", "Chi nhánh", "Mã hàng", "Sản phẩm", "SL xuất", "SL trả", "SL bán ròng", "Đơn giá",
                     "Thành tiền", "CK (%)", "Tiền CK", "Doanh thu ròng"]
    assert ws["A1"].value == "BÁO CÁO BÁN HÀNG THÁNG 9/2026"
    assert ws["A3"].value == "Đơn vị:" and ws["B3"].value == "Công ty TNHH Nam An Tech"     # fields stay put
    assert [ws.cell(r, 1).value for r in range(9, 18)] == [1, 2, 3, 4, 5, 6, 7, 8, None]
    assert ws["I9"].value == '=IF(G9="","",G9*H9)'                                       # kept formula, remapped
    assert ws["G9"].value.startswith("=") and "E9-F9" in ws["G9"].value
    assert ws["I21"].value == "=SUM(I9:I20)" and ws["L21"].value == "=SUM(L9:L20)"
    wb = _values(out, "Hóa đơn")
    # HN MN27: (30-1)*3.45m = 100.05m; HP LT14: (12-1)*15.9m*(1-2%)
    assert wb.value("Hóa đơn", 10, 9) == 29 * 3450000
    assert wb.value("Hóa đơn", 13, 12) == pytest.approx(11 * 15900000 * 0.98)
    total = sum((int(HN[k][0]) - int(HN[k][2])) * float(PRICES[k][0]) * 0.99 +
                (int(HP[k][0]) - int(HP[k][2])) * float(PRICES[k][0]) * 0.98 for k in HN)
    assert wb.value("Hóa đơn", 21, 12) == pytest.approx(total)
    # extra summary sheet aggregates the main table with formulas
    s = load_workbook(out)["Tổng hợp chi nhánh"]
    assert s["A3"].value == "Chi nhánh" and "SUMIF('Hóa đơn'!$B$9:$B$20" in s["B4"].value
    assert wb.value("Tổng hợp chi nhánh", 6, 2) == pytest.approx(total)
    # no VAT invented: the template's VAT input stays blank
    assert not any(isinstance(c, str) and "VAT" in c for c in heads)
    # the dry-run preview is what the file contains
    prev = R.table_preview(out, plan)
    assert prev["rows"] == plan["preview"]["rows"] and prev["sheets"] == plan["preview"]["sheets"]


def test_matching_template_no_needless_changes(tpl, tmp_path):
    cols = [_col("STT", "STT", "keep", "number"), _col("Sản phẩm", "Sản phẩm", "keep", "text"),
            _col("Số lượng", "Số lượng", "keep", "number"), _col("Đơn giá", "Đơn giá", "keep", "number"),
            _col("Thành tiền", "Thành tiền", "keep", "number")]
    rows = [R.LlmRow(cells=[v(header="Sản phẩm", value="LT14", quote="18 laptop LT14"),
                            v(header="Số lượng", value="18", quote="18 laptop LT14"),
                            v(header="Đơn giá", value="15900000", quote="giá niêm yết 15,9 triệu một chiếc")])]
    plan = _propose(tpl, R.LlmStructure(columns=cols, rows=rows))
    assert all(c["changes"] == ["keep"] for c in plan["columns"]) and not plan["dropped"]
    assert plan["preview"]["ok"]


def test_missing_and_fabricated_values_left_blank(tpl):
    ans = answer(sheets=False)
    ans.rows[0].cells.append(v(header="CK (%)", value="5", quote="chiết khấu 5% cho mọi đơn"))   # not in source
    ans.kv.append(R.LlmKv(label="Số hóa đơn", value="HD999", quote="số hóa đơn HD999"))
    plan = _propose(tpl, ans)
    assert all(k["cell"] != "B5" for k in plan["kv"])
    assert {"kind": "ai", "what": "thuế suất VAT"} in plan["missing"]
    assert any(m["kind"] == "field" for m in plan["missing"])
    # a non-zero value may not be "implied"
    bad = R._ground_cell(v(header="SL trả", value="2", quote=RET_ALL["HN"], basis="implied"), "number",
                         sources.SourceBundle(texts={"f1": DOC}, names={}), "")
    assert isinstance(bad, str)


def test_reorder_blocked_when_ranges_span_columns(tpl, tmp_path):
    plan = _propose(tpl, answer(sheets=False))
    # move Đơn giá before Sản phẩm, then add a multi-column formula over the table
    cols = plan["columns"]
    i, j = next(k for k, c in enumerate(cols) if c["header"] == "Sản phẩm"), \
        next(k for k, c in enumerate(cols) if c["header"] == "Đơn giá")
    cols.insert(i, cols.pop(j))
    src = tmp_path / "risky.xlsx"
    with zipfile.ZipFile(tpl) as zi, zipfile.ZipFile(src, "w") as zo:
        for n in zi.namelist():
            d = zi.read(n)
            if n == "xl/worksheets/sheet1.xml":
                d = d.decode().replace("</sheetData>", '<row r="30"><c r="A30"><f>SUM(B9:D9)</f></c></row>'
                                                       "</sheetData>").encode()
            zo.writestr(n, d)
    with pytest.raises(InvalidOperation, match="order"):
        R.apply_plan(src, tmp_path / "o.xlsx", plan)
    # without such ranges the move is applied
    out = tmp_path / "ok.xlsx"
    R.apply_plan(tpl, out, plan)
    assert R.validate(tpl, out, plan)["pass"]


def test_fill_only_keeps_template_structure(tpl, tmp_path):
    plan = _propose(tpl, answer())
    plan["apply"] = False
    eff = R.effective(plan)
    out = tmp_path / "f.xlsx"
    R.apply_plan(tpl, out, eff)
    assert R.validate(tpl, out, eff)["pass"]
    ws = load_workbook(out)["Hóa đơn"]
    assert [ws.cell(8, c).value for c in range(1, 6)] == ["STT", "Sản phẩm", "Số lượng", "Đơn giá", "Thành tiền"]
    assert ws["A1"].value == "HÓA ĐƠN BÁN HÀNG" and "Tổng hợp chi nhánh" not in load_workbook(out).sheetnames


def _api(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    from workspace import config
    from workspace.app import create_app
    monkeypatch.setattr(config, "llm_mode", lambda: "azure")
    monkeypatch.setattr(R, "_llm", lambda ctx: (answer(), "fake-model"))
    c = TestClient(create_app())
    ws = c.post("/api/workspaces", json={"mode": "template"}).json()["id"]
    tpl = tmp_path / "hoa_don.xlsx"
    templates._invoice(tpl)
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    c.post(f"/api/workspaces/{ws}/files", files={"file": ("bao_cao.txt", DOC.encode(), "text/plain")},
           data={"role": "source"})
    c.post(f"/api/workspaces/{ws}/files", files={"file": ("hoa_don.xlsx", tpl.read_bytes(), mime)},
           data={"role": "template"})
    return c, ws


def test_generate_needs_approval_then_matches_preview(monkeypatch, tmp_path):
    from workspace import runs, service
    c, ws = _api(monkeypatch, tmp_path)
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Lập báo cáo theo tài liệu"}).json()
    assert r["status"] == "needs_clarification" and not r["structure"]["approved"]
    ids = [q["id"] for q in r["questions"]]
    assert ids[:2] == ["structure_apply", "structure_sheets"]
    assert "thuế suất VAT" in r["questions"][0]["explanation"]
    assert r["structure"]["preview"]["ok"] and len(r["structure"]["preview"]["rows"]) == 8
    assert c.post(f"/api/workspaces/{ws}/generate", json={}).status_code >= 400          # not approved yet
    r = c.post(f"/api/workspaces/{ws}/clarify", json={"clarification_id": r["clarification_id"], "answers": {}}).json()
    assert r["status"] == "ready" and r["structure"]["approved"]
    run = runs.wait(c.post(f"/api/workspaces/{ws}/generate", json={}).json()["run_id"])
    assert run.status == "done", run.events
    assert run.result["matches_preview"] is True
    wb = load_workbook(service.revision_path(run.result["workbook_id"], run.result["revision"]))
    assert wb["Hóa đơn"]["G8"].value == "SL bán ròng" and "Tổng hợp chi nhánh" in wb.sheetnames
    meta = service.get_metadata(run.result["workbook_id"])
    assert any(x["via"] == "assumed" for x in meta["fill_report"])


def test_auto_approve_and_declined_sheets(monkeypatch, tmp_path):
    from workspace import runs, service
    c, ws = _api(monkeypatch, tmp_path)
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Lập báo cáo", "auto_approve": True}).json()
    assert r["status"] == "ready" and r["structure"]["approved"]
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Lập báo cáo"}).json()
    r = c.post(f"/api/workspaces/{ws}/clarify", json={"clarification_id": r["clarification_id"],
                                                       "answers": {"structure_sheets": "Không thêm sheet mới"}}).json()
    assert r["structure"]["approved"] and r["structure"]["preview"]["sheets"] == []
    run = runs.wait(c.post(f"/api/workspaces/{ws}/generate", json={}).json()["run_id"])
    assert run.status == "done" and run.result["matches_preview"] is True
    wb = load_workbook(service.revision_path(run.result["workbook_id"], run.result["revision"]))
    assert wb.sheetnames == ["Hóa đơn"]


def test_extra_sheets_of_different_lengths_get_their_own_totals(tpl, tmp_path):
    ans = answer()
    one = R.LlmSheet(name="Hà Nội", purpose="Doanh thu Hà Nội", reason="tài liệu yêu cầu",
                     columns=[R.LlmSheetColumn(header="Chi nhánh", kind="input", format="text"),
                              R.LlmSheetColumn(header="Doanh thu ròng", kind="formula", total=True,
                                               formula="SUMIF({Hóa đơn!Chi nhánh},{Chi nhánh},"
                                                       "{Hóa đơn!Doanh thu ròng})")],
                     rows=[R.LlmRow(cells=[v(header="Chi nhánh", value="Hà Nội", quote="Về cửa hàng Hà Nội")])])
    ans = ans.model_copy(update={"sheets": ans.sheets + [one]})
    plan = _propose(tpl, ans)
    assert plan["preview"]["ok"], plan["preview"]
    out = tmp_path / "out.xlsx"
    R.apply_plan(tpl, out, plan)
    assert R.validate(tpl, out, plan)["pass"]
    wb = load_workbook(out)
    a, b = wb["Tổng hợp chi nhánh"], wb["Hà Nội"]
    assert a["B6"].value == "=SUM(B4:B5)" and b["B5"].value == "=SUM(B4:B4)"     # each sheet's own row count
    assert a.max_row == 6 and b.max_row == 5
    # a same-row formula is blank-guarded: an empty row shows blank, never a computed 0
    assert a["B4"].value.startswith('=IF(A4="","",SUMIF(')
    with zipfile.ZipFile(out) as z:
        for n in z.namelist():
            if n.startswith("xl/worksheets/sheet"):
                rows = re.findall(r'<row r="(\d+)"', z.read(n).decode("utf-8"))
                assert len(rows) == len(set(rows)), n                            # no duplicate <row> elements


def test_extra_sheet_row_blank_when_any_input_missing_and_quotes_ignore_case(tpl, tmp_path):
    ans = answer()
    stock = R.LlmSheet(name="Tồn kho", purpose="Tồn kho", reason="đối chiếu",
                       columns=[R.LlmSheetColumn(header="Chi nhánh", kind="input", format="text"),
                                R.LlmSheetColumn(header="Mã hàng", kind="input", format="text"),
                                R.LlmSheetColumn(header="Xuất ròng", kind="formula",
                                                 formula="SUMIFS({Hóa đơn!SL bán ròng},{Hóa đơn!Chi nhánh},{Chi nhánh},"
                                                         "{Hóa đơn!Mã hàng},{Mã hàng})")],
                       rows=[R.LlmRow(cells=[v(header="Chi nhánh", value="Hà Nội", quote="về cửa hàng hà nội"),
                                             v(header="Mã hàng", value="LT14", quote="LT14: 25")]),     # invented quote
                             R.LlmRow(cells=[v(header="Chi nhánh", value="Hà Nội", quote="Về cửa hàng Hà Nội"),
                                             v(header="Mã hàng", value="MN27", quote="MN27 là màn hình 27 inch")])])
    ans = ans.model_copy(update={"sheets": [stock]})
    plan = _propose(tpl, ans)
    assert plan["preview"]["ok"], plan["preview"]
    rows = plan["preview"]["sheets"][0]["rows"]
    assert rows[0]["cells"][0]["v"] == "Hà Nội"                     # lower-case quote still grounded
    assert rows[0]["cells"][1]["prov"] == "missing"                  # the invented quote is rejected
    assert rows[0]["cells"][2]["v"] in ("", None)                    # ... so the row's formula stays blank, not 0
    assert rows[1]["cells"][2]["v"] == 29
    out = tmp_path / "out.xlsx"
    R.apply_plan(tpl, out, plan)
    assert load_workbook(out)["Tồn kho"]["C4"].value.startswith('=IF(OR(A4="",B4=""),"",SUMIFS(')


def test_stt_is_numbered_and_paraphrased_quotes_keep_only_the_verbatim_span(tpl, tmp_path):
    ans = answer(sheets=False)
    ans.columns[0] = _col(None, "STT", "input", "number")             # the model forgot it is a row number
    hp_lt14 = ans.rows[4].cells
    i = next(k for k, c in enumerate(hp_lt14) if c.header == "SL xuất")
    hp_lt14[i] = v(header="SL xuất", value="12", quote="12 chiếc LT14 đã giao cho khách hàng")   # not verbatim
    plan = _propose(tpl, ans)
    stt = next(c for c in plan["columns"] if c["header"] == "STT")
    assert stt["action"] == "sequence"
    cell = plan["data"][4]["SL xuất"]
    assert cell["value"] == 12 and cell["quote"] == "12 chiếc LT14" and cell["locator"] == "bao_cao.txt:line:13"
    assert not any("SL xuất" in w for w in plan["warnings"])
    out = tmp_path / "out.xlsx"
    R.apply_plan(tpl, out, plan)
    ws = load_workbook(out)["Hóa đơn"]
    assert [ws.cell(r, 1).value for r in range(9, 17)] == [1, 2, 3, 4, 5, 6, 7, 8]
    # a paraphrase whose number is not in any verbatim span stays blank
    b = sources.SourceBundle(texts={"f1": DOC}, names={"f1": "bao_cao.txt"})
    assert R.ground("12", "Hải Phòng bán 12 laptop", "number", b, "") == "quote not found in the sources"


def test_unused_returns_column_triggers_one_review_round(tpl):
    bad = answer(sheets=False)
    bad.columns[6] = _col(None, "SL bán ròng", "input", "number")    # returns no longer feed any formula
    bad.columns[8] = _col(None, "Thành tiền", "formula", "number", "{SL xuất}*{Đơn giá}", total=True)
    good = answer(sheets=False)
    seen: list[str] = []

    def llm(ctx):
        seen.append(ctx)
        return (bad if len(seen) == 1 else good), "fake"

    b = sources.SourceBundle(texts={"f1": DOC}, names={"f1": "bao_cao.txt"})
    assert R.unused_inputs(bad) and not R.unused_inputs(good)
    plan = R.propose(tpl, templates.detect_slots(tpl), b, "lập báo cáo", llm=llm)
    assert len(seen) == 2 and "SL trả" in seen[1] and '"problems"' in seen[1]
    net = next(c for c in plan["columns"] if c["header"] == "SL bán ròng")
    assert net["action"] == "formula"
    # a review answer that is not better is ignored; a kept template formula counts as a use
    seen.clear()
    plan = R.propose(tpl, templates.detect_slots(tpl), b, "lập báo cáo", llm=lambda ctx: (seen.append(ctx), (bad, "f"))[1])
    assert len(seen) == 2 and any("SL trả" in w and "not used" in w for w in plan["warnings"])
    kept = R.LlmStructure(columns=[_col("Số lượng", "Số lượng", "keep", "number"), _col("Đơn giá", "Đơn giá", "keep", "number"),
                                   _col("Thành tiền", "Thành tiền", "keep", "number")])
    assert R.unused_inputs(kept, {"so luong": "", "thanh tien": "IF({Số lượng}=\"\",\"\",{Số lượng}*{Đơn giá})"}) == []


def test_implied_zero_keeps_the_verbatim_part_of_a_shortened_quote():
    b = sources.SourceBundle(texts={"f1": DOC}, names={"f1": "bao_cao.txt"})
    out = R._ground_cell(v(header="SL trả", value="0", quote="một chiếc LT14 và hai bàn phím KB87", basis="implied",
                           note="không có trả hàng"), "number", b, "")
    assert out["value"] == 0 and out["prov"] == "assumed"
    assert out["quote"] in ("một chiếc LT14", "và hai bàn phím KB87") and R.find_quote(b, "", out["quote"])
    assert R._ground_cell(v(header="SL trả", value="0", quote="không ai trả hàng cả", basis="implied"),
                          "number", b, "") == "quote not found in the sources"


def test_quote_with_an_ellipsis_is_grounded_when_every_part_is_verbatim_on_one_line():
    src = "MN27 là màn hình 27 inch, bán 3.450.000đ, giá vốn 2,62 triệu.\nKB87 bán 1 triệu 250 nghìn."
    b = sources.SourceBundle(texts={"f1": src}, names={"f1": "bc.txt"})
    assert R.ground("3450000", "MN27 ... bán 3.450.000đ", "number", b, "")[:3] == (3450000, "number", "bc.txt:line:1")
    assert R.ground("3450000", "MN27 … bán 3.450.000đ", "number", b, "")[0] == 3450000
    assert R.ground("3450000", "bán 3.450.000đ ... MN27", "number", b, "") == "quote not found in the sources"  # order
    assert R.ground("1250000", "MN27 ... bán 1 triệu 250 nghìn", "number", b, "") == "quote not found in the sources"
    assert R.ground("3450000", "MN27 ... bán 3.450.000đ ... giảm giá", "number", b, "") == "quote not found in the sources"
