"""Phase 4: chat agent (Deep Agents + tool allowlist) in mock mode, plans, grounding, revision guard, gate flow."""
from __future__ import annotations

import io
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from workspace import runs, service
from workspace.agent import runtime, store, tools
from workspace.agent.mock_model import MockChatModel
from workspace.app import create_app

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(c: TestClient, path) -> str:
    r = c.post("/api/workbooks/upload", files={"file": ("sales.xlsx", path.read_bytes(), XLSX)})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _turn(c: TestClient, wid: str, rev: int, message: str, thread_id: str | None = None, **kw) -> tuple[dict, runs.Run]:
    body = {"workbook_id": wid, "revision": rev, "message": message, **kw}
    if thread_id:
        body["thread_id"] = thread_id
    r = c.post("/api/agent/turns", json=body)
    assert r.status_code == 200, r.text
    started = r.json()
    run = runs.wait(started["turn_id"])
    assert run.status == "done", run.events
    return started, run


def _events(run: runs.Run, name: str) -> list[dict]:
    return [e["data"] for e in run.events if e["event"] == name]


def _cell(c: TestClient, wid: str, sheet: str, ref: str):
    return load_workbook(io.BytesIO(c.get(f"/api/workbooks/{wid}/export").content))[sheet][ref].value


def test_chat_sets_cell_and_streams_events(simple_sales):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    started, run = _turn(c, wid, 1, "Đặt ô B3 = 25", active_sheet="Sales")
    names = [e["event"] for e in run.events]
    assert names[0] == "run_started" and names[-1] == "done"
    assert ("tool_started", "workbook_revision", "tool_completed") == tuple(
        n for n in names if n in ("tool_started", "tool_completed", "workbook_revision"))
    assert _events(run, "tool_completed")[0] == {"tool": "propose_workbook_edits", "status": "ok",
                                                 "summary": "revision 2"}
    done = _events(run, "done")[0]
    assert done["revision"] == 2 and "revision 2" in done["text"] and _events(run, "text_delta")
    assert _cell(c, wid, "Sales", "B3") == 25 and _cell(c, wid, "Sales", "D3") == "=B3*C3"
    meta = service.get_metadata(wid)
    assert meta["revisions"][-1]["origin"] == "agent"
    # the thread remembers the conversation; history log has no cell values
    t = store.get_thread(started["thread_id"], wid)
    assert [m["role"] for m in t["messages"]] == ["user", "assistant"]
    log = [json.loads(x) for x in (store._dir("logs") / f"{started['thread_id']}.jsonl").read_text(
        encoding="utf-8").splitlines()]
    assert [e["tool"] for e in log] == ["propose_workbook_edits"]
    assert log[0]["args"] == {"edits": 1, "cells": ["B3"]}               # addresses only, not the value 25


def test_no_filesystem_or_shell_tools_are_exposed(simple_sales):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    _turn(c, wid, 1, "tóm tắt workbook")
    assert MockChatModel.bound_tools == sorted(["apply_workbook_edits", "build_workbook", "inspect_workbook",
                                                "propose_workbook_edits", "read_range", "read_source", "transform_range",
                                                "run_template_extractor", "validate_workbook"])
    assert not set(MockChatModel.bound_tools) & runtime.EXCLUDED_TOOLS


def test_read_and_inspect_are_read_only(simple_sales):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    _, run = _turn(c, wid, 1, "đọc A2:D6")
    text = _events(run, "done")[0]["text"]
    assert "B3=20" in text and "D6=" in text
    _, run = _turn(c, wid, 1, "summarize the workbook")
    assert "Sales" in _events(run, "done")[0]["text"]
    _, run = _turn(c, wid, 1, "xin chào")                              # not understood: honest help text
    assert "mock" in _events(run, "done")[0]["text"].lower()
    assert service.get_metadata(wid)["current_revision"] == 1


def test_large_edit_needs_confirmation_then_applies(simple_sales):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    started, run = _turn(c, wid, 1, "xóa A10:F20")                    # 66 cells > 25
    conf = _events(run, "confirmation_required")[0]
    assert conf["cells"] == 66 and conf["base_revision"] == 1 and len(conf["diff"]) == 66
    assert service.get_metadata(wid)["current_revision"] == 1         # nothing written yet
    # explicit confirmation in chat applies the pending plan
    _, run2 = _turn(c, wid, 1, "đồng ý", thread_id=started["thread_id"])
    assert _events(run2, "tool_completed")[0]["tool"] == "apply_workbook_edits"
    assert _events(run2, "done")[0]["revision"] == 2
    assert c.post(f"/api/agent/plans/{conf['plan_id']}/apply", json={"base_revision": 1}).status_code == 422


def test_plan_apply_and_reject_over_http(simple_sales):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    _, run = _turn(c, wid, 1, "clear A10:F20")
    pid = _events(run, "confirmation_required")[0]["plan_id"]
    r = c.post(f"/api/agent/plans/{pid}/apply", json={"base_revision": 1})
    assert r.status_code == 200 and r.json()["revision"] == 2
    _, run = _turn(c, wid, 2, "clear A30:F40")
    pid2 = _events(run, "confirmation_required")[0]["plan_id"]
    assert c.post(f"/api/agent/plans/{pid2}/reject", json={}).json() == {"ok": True, "status": "rejected"}
    assert c.post(f"/api/agent/plans/{pid2}/apply", json={"base_revision": 2}).status_code == 422
    assert c.post("/api/agent/plans/plan_0000000000000000/apply", json={"base_revision": 2}).status_code == 404
    # a stale plan (workbook moved on) is refused with 409
    _, run = _turn(c, wid, 2, "clear A50:F60")
    pid3 = _events(run, "confirmation_required")[0]["plan_id"]
    _turn(c, wid, 2, "set A1 = 'Báo cáo'")                             # -> revision 3
    assert c.post(f"/api/agent/plans/{pid3}/apply", json={"base_revision": 2}).status_code == 409


def test_ungrounded_numbers_and_formulas_are_refused(simple_sales):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    run = runs.Run("turn_0000000000000001", "agent")
    ctx = tools.Ctx(run=run, workbook_id=wid, revision=1, thread_id="thr_0000000000000001",
                    message="điền doanh thu tháng 10", user_text="điền doanh thu tháng 10")
    by = {t.name: t for t in tools.make_tools(ctx)}
    out = json.loads(by["propose_workbook_edits"].invoke({"edits": [{"op": "set", "cell": "B3", "value": 777}]}))
    assert "not grounded" in out["error"] and out["ungrounded"] == ["Sales!B3=777"]
    out = json.loads(by["propose_workbook_edits"].invoke({"edits": [{"cell": "E3", "value": "=B3*2"}]}))
    assert "formulas is not supported" in out["error"]
    # a number read from the workbook in this turn is grounded
    by["read_range"].invoke({"range": "B3:C5"})
    out = json.loads(by["propose_workbook_edits"].invoke({"edits": [{"cell": "B10", "value": 450000}]}))
    assert out["revision"] == 2
    assert json.loads(by["read_range"].invoke({"range": "A1:Z100"}))["error"].startswith("Range has 2600")
    assert service.get_metadata(wid)["current_revision"] == 2


def test_stale_revision_is_409(simple_sales):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    _turn(c, wid, 1, "set B3 = 21")
    r = c.post("/api/agent/turns", json={"workbook_id": wid, "revision": 1, "message": "set B4 = 30"})
    assert r.status_code == 409 and r.json()["current_revision"] == 2
    assert c.post("/api/agent/turns", json={"workbook_id": wid, "revision": 2, "message": "x",
                                            "thread_id": "thr_0000000000000000"}).status_code == 404


def test_gate_manual_then_agent_then_export_and_undo(simple_sales):
    """Phase 4 gate: manual edit -> agent edit -> export matches both; undo returns to the manual revision."""
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    r = c.post(f"/api/workbooks/{wid}/operations", json={
        "base_revision": 1, "operation_id": "ui-manual-0001", "origin": "manual",
        "actions": [{"type": "set_value", "sheet": "Sales", "cell": "A3", "value": "Laptop Pro"}]})
    assert r.json()["revision"] == 2
    _, run = _turn(c, wid, 2, "đặt C4 = 210000 và ghi A10 = 'Ghi chú: giá mới'", active_sheet="Sales")
    assert _events(run, "done")[0]["revision"] == 3
    s = load_workbook(io.BytesIO(c.get(f"/api/workbooks/{wid}/export").content))["Sales"]
    assert (s["A3"].value, s["C4"].value, s["A10"].value) == ("Laptop Pro", 210000, "Ghi chú: giá mới")
    assert s["D6"].value == "=SUM(D3:D5)"
    assert [r["origin"] for r in service.get_metadata(wid)["revisions"]] == ["upload", "manual", "agent"]
    prev = c.get(f"/api/workbooks/{wid}/preview")
    assert prev.status_code == 200 and "Laptop Pro" in prev.text and "Ghi chú: giá mới" in prev.text
    u = c.post(f"/api/workbooks/{wid}/undo", json={"base_revision": 3}).json()
    assert u["revision"] == 4
    s = load_workbook(io.BytesIO(c.get(f"/api/workbooks/{wid}/export").content))["Sales"]
    assert (s["A3"].value, s["C4"].value, s["A10"].value) == ("Laptop Pro", 200000, None)


class _ClaimingModel(MockChatModel):
    """Simulates the live failure: the model claims an edit in text without calling any tool."""

    def _next(self, messages):
        from langchain_core.messages import AIMessage
        return AIMessage(content="Đã xóa toàn bộ nội dung vùng Sales!A10:F20 (revision 6).")


class _FailsAfterWriteModel(MockChatModel):
    """Simulates the live failure: the edit tool succeeds, then the provider rate-limits the follow-up call."""

    def _next(self, messages):
        from langchain_core.messages import ToolMessage
        if isinstance(messages[-1], ToolMessage):
            raise RuntimeError("rate limited")
        return super()._next(messages)


def test_model_failure_after_write_reports_the_new_revision(simple_sales, monkeypatch):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    monkeypatch.setattr(runtime, "_models", lambda: [("a", _FailsAfterWriteModel()), ("b", MockChatModel())])
    _, run = _turn(c, wid, 1, "đặt B3 = 25")
    done = _events(run, "done")[0]
    assert done["revision"] == 2 and "revision 2" in done["text"] and done["model"] is None
    assert not _events(run, "model_fallback")                         # no retry once something was written
    assert _cell(c, wid, "Sales", "B3") == 25


def test_unbacked_edit_claim_is_flagged(simple_sales, monkeypatch):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    real = runtime._models
    monkeypatch.setattr(runtime, "_models", lambda: [("fake", _ClaimingModel())])
    started, run = _turn(c, wid, 1, "Xóa toàn bộ nội dung vùng Sales!A10:F20.")
    done = _events(run, "done")[0]
    assert not _events(run, "tool_started") and done["revision"] == 1
    assert "KHÔNG có thay đổi" in done["text"] and "revision 1" in done["text"]
    assert "KHÔNG có thay đổi" in "".join(e["text"] for e in _events(run, "text_delta"))
    assert "KHÔNG có thay đổi" in store.get_thread(started["thread_id"], wid)["messages"][-1]["content"]
    assert service.get_metadata(wid)["current_revision"] == 1
    # a real edit is not flagged
    monkeypatch.setattr(runtime, "_models", real)
    _, run = _turn(c, wid, 1, "đặt B3 = 25")
    assert "⚠️" not in _events(run, "done")[0]["text"]
    assert not runtime.CLAIM_RE.search("Sales!A2:D6: B3=20, C3=150000") and not runtime.CLAIM_RE.search(
        "Workbook có 1 sheet: Sales (A1:D6).")


def test_mock_asks_back_instead_of_guessing(simple_sales):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    for msg, expect in [("điền doanh thu tháng 10", "ô nào"), ("thêm công thức tổng cho cột B", "chưa được hỗ trợ")]:
        _, run = _turn(c, wid, 1, msg)
        done = _events(run, "done")[0]
        assert expect in done["text"] and done["revision"] == 1 and not _events(run, "tool_started")


def test_transform_range_converts_with_the_users_rate_in_code(simple_sales):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)
    run = runs.Run("turn_0000000000000002", "agent")
    text = "Cột này thành tiền đô\n1USD = 25000VND"
    ctx = tools.Ctx(run=run, workbook_id=wid, revision=1, thread_id="thr_0000000000000002", message="1USD = 25000VND",
                    user_text=text, vi=True)
    by = {t.name: t for t in tools.make_tools(ctx)}
    # the model's own arithmetic is still refused, with a pointer to transform_range
    out = json.loads(by["propose_workbook_edits"].invoke({"edits": [{"cell": "C3", "value": 600}]}))
    assert out["ungrounded"] == ["Sales!C3=600"] and "transform_range" in out["error"]
    # a rate the user never gave is refused
    out = json.loads(by["transform_range"].invoke({"range": "C3:C5", "op": "divide", "factor": 24000}))
    assert out["ungrounded"] == ["factor=24000"] and service.get_metadata(wid)["current_revision"] == 1
    # header text and formula cells are skipped; the values are computed by code
    out = json.loads(by["transform_range"].invoke({"range": "C2:D5", "op": "divide", "factor": 25000}))
    assert out["revision"] == 2 and out["skipped_formula_cells"] == ["D3", "D4", "D5"]
    assert out["computed_by"].startswith("code: value / 25000")
    wb = load_workbook(io.BytesIO(c.get(f"/api/workbooks/{wid}/export").content))["Sales"]
    assert [wb[f"C{r}"].value for r in (2, 3, 4, 5)] == ["Unit price", 600, 8, 18]
    assert wb["D3"].value == "=B3*C3" and wb["D6"].value == "=SUM(D3:D5)"
    # multiplying by the reciprocal of the user's rate is grounded too
    out = json.loads(by["transform_range"].invoke({"range": "B3", "op": "multiply", "factor": 0.00004,
                                                   "round_digits": 6}))
    assert out["revision"] == 3
    assert json.loads(by["transform_range"].invoke({"range": "D3:D5", "op": "divide", "factor": 25000}))[
        "skipped_formula_cells"] == ["D3", "D4", "D5"]


def test_turn_that_ends_in_a_loop_reports_the_tool_refusal(simple_sales, monkeypatch):
    c = TestClient(create_app())
    wid = _upload(c, simple_sales)

    def looping(run, model, tool_list, system, msgs):
        by = {t.name: t for t in tool_list}
        for _ in range(3):
            by["propose_workbook_edits"].invoke({"edits": [{"cell": "C3", "value": 600}]})
        raise RecursionError("recursion limit")

    monkeypatch.setattr(runtime, "_invoke", looping)
    _, run = _turn(c, wid, 1, "xác nhận, convert đi")
    text = _events(run, "done")[0]["text"]
    assert "Model không phản hồi" not in text and "không có nguồn" in text and "revision 1" in text
