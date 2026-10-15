"""Template restructuring: AI-proposed column structure from prose sources, grounded values, OOXML column moves."""
from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from workspace import restructure as R
from workspace import sources, templates

ROOT = Path(__file__).resolve().parent.parent
TPL = ROOT / "sample_data" / "06_so_ban_hang_nam" / "template_so_ban_hang.xlsx"

DOC = """Cột "Mã KH" có thể bỏ hẳn. Cột "Tháng" ở cuối cũng không cần nữa.
Thuế suất mặc định có thể để 8% cho dòng mới, nhưng vẫn phải sửa được.
Phần thông tin đầu trang, điền "Năm: 2026" và "Đơn vị: Cửa hàng Minh Phát"; trường "Người lập" để trống.
1) Ngày 05/09/2026, hóa đơn HD001, khách Công ty An Phúc, mã hàng SP001, tên hàng Bàn phím cơ, đơn vị cái, số lượng 3, đơn giá 750.000 đồng, chiết khấu 5%, VAT 8%, ghi chú "Đã thanh toán".
2) Ngày 05/09/2026, hóa đơn HD001, khách Công ty An Phúc, mã hàng SP002, tên hàng Chuột không dây, đơn vị cái, số lượng 5, đơn giá 320.000 đồng, không chiết khấu, VAT 8%, ghi chú "Giao cùng HD001".
"""
L1 = DOC.splitlines()[3]
L2 = DOC.splitlines()[4]


def _col(src, header, kind, fmt, formula=None, default=None, dq=None):
    return R.LlmColumn(source_header=src, header=header, kind=kind, format=fmt, formula=formula, default=default,
                       default_quote=dq, reason="")


def _llm_answer() -> R.LlmStructure:
    cols = [
        _col("Ngày", "Ngày", "keep", "date"),
        _col("Số HĐ", "Số hóa đơn", "input", "text"),
        _col(None, "Tên khách hàng", "input", "text"),
        _col("Mã hàng", "Mã hàng", "keep", "text"),
        _col("Tên hàng", "Tên hàng", "input", "text"),
        _col(None, "Đơn vị tính", "input", "text"),
        _col("Số lượng", "Số lượng", "keep", "number"),
        _col("Đơn giá", "Đơn giá", "input", "number"),
        _col("Thành tiền", "Thành tiền", "formula", "number", "{Số lượng}*{Đơn giá}"),
        _col("CK (%)", "Chiết khấu (%)", "input", "percent_number"),
        _col("Tiền CK", "Tiền chiết khấu", "formula", "number", "{Thành tiền}*{Chiết khấu (%)}/100"),
        _col(None, "Thuế suất VAT (%)", "input", "percent_number", default="8%",
             dq="Thuế suất mặc định có thể để 8% cho dòng mới"),
        _col("VAT", "Tiền VAT", "formula", "number", "({Thành tiền}-{Tiền chiết khấu})*{Thuế suất VAT (%)}/100"),
        _col("Phải thu", "Tổng thanh toán", "formula", "number", "{Thành tiền}-{Tiền chiết khấu}+{Tiền VAT}"),
        _col(None, "Ghi chú", "input", "text"),
    ]
    v = R.LlmValue
    rows = [R.LlmRow(cells=[v(header="Ngày", value="05/09/2026", quote="Ngày 05/09/2026"),
                            v(header="Số hóa đơn", value="HD001", quote="hóa đơn HD001"),
                            v(header="Tên khách hàng", value="Công ty An Phúc", quote="khách Công ty An Phúc"),
                            v(header="Mã hàng", value="SP001", quote="mã hàng SP001"),
                            v(header="Tên hàng", value="Bàn phím cơ", quote="tên hàng Bàn phím cơ"),
                            v(header="Đơn vị tính", value="cái", quote="đơn vị cái"),
                            v(header="Số lượng", value="3", quote="số lượng 3"),
                            v(header="Đơn giá", value="750000", quote="đơn giá 750.000 đồng"),
                            v(header="Chiết khấu (%)", value="5", quote="chiết khấu 5%"),
                            v(header="Thuế suất VAT (%)", value="8", quote="VAT 8%"),
                            v(header="Ghi chú", value="Đã thanh toán", quote='ghi chú "Đã thanh toán"'),
                            v(header="Thành tiền", value="2250000", quote="số lượng 3")]),      # formula: ignored
            R.LlmRow(cells=[v(header="Số hóa đơn", value="HD001", quote="hóa đơn HD001"),
                            v(header="Chiết khấu (%)", value="0", quote="không chiết khấu"),
                            v(header="Đơn giá", value="999000", quote="đơn giá 320.000 đồng"),      # not grounded
                            v(header="Tên khách hàng", value="Công ty Bịa", quote="khách Công ty Bịa")])]
    kv = [R.LlmKv(label="Năm", value="2026", quote="Năm: 2026"),
          R.LlmKv(label="Đơn vị", value="Cửa hàng Minh Phát", quote="Đơn vị: Cửa hàng Minh Phát")]
    return R.LlmStructure(columns=cols, kv=kv, rows=rows)


@pytest.fixture
def plan(tmp_path):
    if not TPL.exists():
        pytest.skip("sample template not present")
    b = sources.SourceBundle(texts={"f1": DOC}, names={"f1": "yeu_cau.txt"})
    slots = templates.detect_slots(TPL)
    return R.propose(TPL, slots, b, "sửa template theo file", llm=lambda ctx: (_llm_answer(), "fake"))


def test_remap_formula_moves_deletes_and_shrinks():
    cm = R.ColMap({1: 1, 2: 2, 3: None, 4: 5}, 1, 4, 1)
    assert R.remap_formula("SUM(A1:D1)+C2+D3+E4", "S", "S", cm) == "SUM(A1:E1)+#REF!+E3+F4"
    assert R.remap_formula("'S'!$D$1+T!D1+\"D1\"", "T", "S", cm) == "'S'!$E$1+T!D1+\"D1\""
    assert R.remap_formula("SUM(S!C:D)", None, "S", cm) == "SUM(S!E:E)"


def test_ground_rejects_values_not_in_quote():
    b = sources.SourceBundle(texts={"f1": L1}, names={"f1": "a.txt"})
    assert R.ground("750000", "đơn giá 750.000 đồng", "number", b, "")[0] == 750000
    assert R.ground("5", "chiết khấu 5%", "percent_number", b, "")[0] == 5
    assert isinstance(R.ground("999", "đơn giá 750.000 đồng", "number", b, ""), str)
    assert isinstance(R.ground("750000", "đơn giá 750.000 USD", "number", b, ""), str)      # quote not in source
    assert R.ground("05/09/2026", "Ngày 05/09/2026", "date", b, "")[0] == "2026-09-05"


def test_plan_grounding_conflicts_and_formula_checks(plan):
    heads = [c["header"] for c in plan["columns"]]
    assert heads[:3] == ["Ngày", "Số hóa đơn", "Tên khách hàng"] and heads[-1] == "Ghi chú"
    dropped = {d["header"]: d for d in plan["dropped"]}
    assert set(dropped) == {"Mã KH", "Tháng"}
    assert all(d["mode"] == "hide" and d["used_by"] for d in dropped.values())   # Tổng hợp still uses both
    vat = next(c for c in plan["columns"] if c["header"] == "Thuế suất VAT (%)")
    assert vat["default"] == 8 and vat["changes"] == ["add", "default"]
    assert plan["kv"][0]["value"] == "2026" and plan["kv"][1]["value"] == "Cửa hàng Minh Phát"
    r1, r2 = plan["data"]
    assert r1["Đơn giá"]["value"] == 750000 and "Thành tiền" not in r1
    assert r2["Chiết khấu (%)"]["value"] == 0 and "Đơn giá" not in r2 and "Tên khách hàng" not in r2
    assert any("Đơn giá" in w for w in plan["warnings"])
    bad = R.check_formula("{Số lượng}*1.1", {"Số lượng"}, set())
    assert bad and "constant" in bad
    assert R.check_formula("{X}*2", {"Số lượng"}, set())
    assert R.check_formula("A1*2", set(), set())
    assert R.check_formula("SUM({Số lượng})", {"Số lượng"}, set()) is None


def test_apply_plan_moves_columns_and_keeps_workbook_consistent(plan, tmp_path):
    src = tmp_path / "t.xlsx"
    shutil.copyfile(TPL, src)
    out = tmp_path / "o.xlsx"
    res = R.apply_plan(src, out, plan)
    rep = R.validate(src, out, plan)
    assert rep["pass"], rep["errors"]
    assert res["changed"]
    wb = load_workbook(out)
    ws = wb["Sổ bán hàng"]
    heads = [ws.cell(7, c).value for c in range(1, 18)]
    assert heads == ["Ngày", "Số hóa đơn", "Mã KH", "Tên khách hàng", "Mã hàng", "Tên hàng", "Đơn vị tính",
                     "Số lượng", "Đơn giá", "Thành tiền", "Chiết khấu (%)", "Tiền chiết khấu", "Thuế suất VAT (%)",
                     "Tiền VAT", "Tổng thanh toán", "Tháng", "Ghi chú"]
    assert ws.column_dimensions["C"].hidden and ws.column_dimensions["P"].hidden
    assert ws["J8"].value == '=IF(OR(H8="",I8=""),"",H8*I8)'
    assert ws["L9"].value == '=IF(OR(J9="",K9=""),"",J9*K9/100)'
    assert ws["N8"].value == '=IF(OR(J8="",L8="",M8=""),"",(J8-L8)*M8/100)'
    assert ws["O8"].value == '=IF(OR(J8="",L8="",N8=""),"",J8-L8+N8)'
    assert ws["P8"].value == '=IF(A8="","",MONTH(A8))'               # hidden Tháng still follows the date
    assert ws["I8"].value == 750000 and ws["K8"].value == 5 and ws["M8"].value == 8 and ws["M20"].value == 8
    assert ws["Q8"].value == "Đã thanh toán" and ws["D8"].value == "Công ty An Phúc" and ws["D9"].value is None
    assert ws["A8"].number_format == "dd/mm/yyyy" and ws["A8"].value.day == 5
    assert ws["I8"].number_format == "#,##0" and ws["F8"].value == "Bàn phím cơ"
    assert ws["B3"].value == "2026" and ws["B4"].value == "Cửa hàng Minh Phát"
    assert ws.auto_filter.ref == "A7:Q407" and ws.freeze_panes == "A8"
    assert ws["O408"].value == "=SUBTOTAL(9,O8:O407)"                 # totals follow Phải thu (L -> O)
    assert ws["H7"].fill.fgColor.rgb == ws["D7"].fill.fgColor.rgb      # new headers copy the header look
    th = wb["Tổng hợp"]
    fs = [c.value for row in th.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("'Sổ bán hàng'!$P$8:$P$407" in f and "'Sổ bán hàng'!$O$8:$O$407" in f for f in fs)
    assert not any("#REF!" in f for f in fs)
    wb.close()
    with zipfile.ZipFile(out) as z:
        wbx = z.read("xl/workbook.xml").decode()
        assert "$A$7:$Q$407" in wbx and 'fullCalcOnLoad="1"' in wbx
        assert z.read("xl/styles.xml") == zipfile.ZipFile(src).read("xl/styles.xml")


def test_fill_only_keeps_original_columns(plan, tmp_path):
    plan["apply"] = False
    eff = R.effective(plan)
    out = tmp_path / "o.xlsx"
    R.apply_plan(TPL, out, eff)
    assert R.validate(TPL, out, eff)["pass"]
    ws = load_workbook(out)["Sổ bán hàng"]
    assert [ws.cell(7, c).value for c in range(1, 14)][:3] == ["Ngày", "Số HĐ", "Mã KH"]
    assert ws["B8"].value == "HD001" and ws["F8"].value == 3
    assert isinstance(ws["G8"].value, str) and ws["G8"].value.startswith("=IF(")      # formula column untouched
    assert ws.cell(7, 14).value is None


def test_analyze_clarify_generate_restructures_via_api(monkeypatch):
    if not TPL.exists():
        pytest.skip("sample template not present")
    from fastapi.testclient import TestClient

    from workspace import config, runs, service
    from workspace.app import create_app
    monkeypatch.setattr(config, "llm_mode", lambda: "azure")
    monkeypatch.setattr(R, "_llm", lambda ctx: (_llm_answer(), "fake-model"))
    c = TestClient(create_app())
    ws = c.post("/api/workspaces", json={"mode": "template"}).json()["id"]
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert c.post(f"/api/workspaces/{ws}/files", files={"file": ("yeu_cau.txt", DOC.encode(), "text/plain")},
                  data={"role": "source"}).status_code == 200
    assert c.post(f"/api/workspaces/{ws}/files", files={"file": ("so.xlsx", TPL.read_bytes(), mime)},
                  data={"role": "template"}).status_code == 200
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Sửa mẫu theo tài liệu"}).json()
    assert r["status"] == "needs_clarification", r
    ids = [q["id"] for q in r["questions"]]
    assert ids[0] == "structure_apply" and {"drop_3", "drop_13"} <= set(ids)
    assert r["structure"]["columns"][2]["header"] == "Mã KH" and r["structure"]["columns"][2]["changes"] == ["hidden"]
    assert not r["mapping"]
    r = c.post(f"/api/workspaces/{ws}/clarify", json={"clarification_id": r["clarification_id"],
                                                       "answers": {"drop_3": "Giữ cột hiển thị"}}).json()
    assert r["status"] == "ready" and r["structure"]["columns"][2]["changes"] == ["keep"]
    run = runs.wait(c.post(f"/api/workspaces/{ws}/generate", json={}).json()["run_id"])
    assert run.status == "done", run.events
    wb_id = run.result["workbook_id"]
    assert run.result["revision"] == 2
    meta = service.get_metadata(wb_id)
    assert meta["structure"]["columns"][-1]["header"] == "Ghi chú"
    assert any(x["cell"] == "I8" and x["value"] == 750000 for x in meta["fill_report"])
    ws_ = load_workbook(service.revision_path(wb_id, 2))["Sổ bán hàng"]
    assert ws_["Q7"].value == "Ghi chú" and not ws_.column_dimensions["C"].hidden
    assert ws_.column_dimensions["P"].hidden
    ws1 = load_workbook(service.revision_path(wb_id, 1))["Sổ bán hàng"]
    assert ws1["B7"].value == "Số HĐ"                                  # revision 1 = untouched template
    assert c.get(f"/api/workbooks/{wb_id}/preview?revision=2").status_code == 200


def test_analyze_falls_back_when_llm_fails(monkeypatch):
    if not TPL.exists():
        pytest.skip("sample template not present")
    from fastapi.testclient import TestClient

    from workspace import config
    from workspace.app import create_app

    def boom(ctx):
        raise RuntimeError("model unavailable")
    monkeypatch.setattr(config, "llm_mode", lambda: "azure")
    monkeypatch.setattr(R, "_llm", boom)
    from workspace import template_fill
    orig = template_fill.propose
    monkeypatch.setattr(template_fill, "propose", lambda slots, b, llm: orig(slots, b, False))   # no network
    c = TestClient(create_app())
    ws = c.post("/api/workspaces", json={"mode": "template"}).json()["id"]
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    c.post(f"/api/workspaces/{ws}/files", files={"file": ("yeu_cau.txt", DOC.encode(), "text/plain")},
           data={"role": "source"})
    c.post(f"/api/workspaces/{ws}/files", files={"file": ("so.xlsx", TPL.read_bytes(), mime)},
           data={"role": "template"})
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Sửa mẫu"}).json()
    assert r["structure"] is None and any("could not propose" in w for w in r["warnings"])


def test_fix_percent_scales_each_percent_column_once():
    fm = {"T": "number", "P": "percent", "N": "percent_number"}
    assert R.fix_percent("ROUND({T}*{P}/100,0)", fm) == "ROUND({T}*{P},0)"      # 5% is stored as 0.05
    assert R.fix_percent("{T}*{N}", fm) == "{T}*({N}/100)"                   # 5 means 5%
    assert R.fix_percent("{T}*{N}/100", fm) == "{T}*{N}/100"
    assert R.fix_percent("{T}/{P}/100", fm) == "{T}/{P}/100"                 # divisor: left alone


def test_row_formula_does_not_double_the_blank_guard():
    letters = {"Số lượng": "H", "Đơn giá": "I"}
    assert R.row_formula('IF({Số lượng}="","",{Số lượng}*{Đơn giá})', letters, 8) == 'IF(OR(H8="",I8=""),"",H8*I8)'
    assert R.row_formula("{Số lượng}*{Đơn giá}", letters, 8) == 'IF(OR(H8="",I8=""),"",H8*I8)'
    # an inner IF that is not a whole-formula guard is kept as is
    assert R.row_formula('IF({Số lượng}="","",1)+{Đơn giá}', letters, 8) == \
        'IF(OR(H8="",I8=""),"",IF(H8="","",1)+I8)'


def test_lookup_column_becomes_input_when_documents_state_values():
    assert R.LOOKUP_RE.search('IF(D8="","",IFERROR(VLOOKUP(D8,\'Danh mục\'!$A$4:$C$13,2,FALSE),""))')
    assert not R.LOOKUP_RE.search('IF(F8="","",F8*G8)')
