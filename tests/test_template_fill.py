"""Phase 3: template gallery, slot detection, template fill via set_value, extractor adapter gating."""
from __future__ import annotations

import io
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from workspace import config, runs, service, sources, storage, template_fill, templates
from workspace.app import create_app

ITEMS_CSV = "Sản phẩm,Số lượng,Đơn giá\nBút bi,10,5000\nVở,4,12000\nThước,7,8000\n".encode()
HEAD_TXT = "Khách hàng: Công ty ABC\nSố hóa đơn: HD-001\nNgày: 15/09/2026\n".encode()


def _ws(c: TestClient, files=(), template: tuple[str, bytes] | None = None) -> str:
    ws = c.post("/api/workspaces", json={"mode": "template"}).json()["id"]
    for name, data in files:
        assert c.post(f"/api/workspaces/{ws}/files", files={"file": (name, data, "text/plain")},
                      data={"role": "source"}).status_code == 200
    if template:
        r = c.post(f"/api/workspaces/{ws}/files", files={"file": (template[0], template[1], config_mime())},
                   data={"role": "template"})
        assert r.status_code == 200, r.text
    return ws


def config_mime() -> str:
    return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _run(c: TestClient, ws: str, body: dict, answers: dict | None = None) -> runs.Run:
    r = c.post(f"/api/workspaces/{ws}/analyze", json=body).json()
    if r["status"] == "needs_clarification":
        r = c.post(f"/api/workspaces/{ws}/clarify", json={"clarification_id": r["clarification_id"],
                                                           "answers": answers or {}}).json()
    assert r["status"] == "ready", r
    run = runs.wait(c.post(f"/api/workspaces/{ws}/generate", json={}).json()["run_id"])
    assert run.status == "done", run.events
    return run


def _builtin(c: TestClient, key_name: str) -> str:
    return next(t["id"] for t in c.get("/api/templates").json() if t["name"] == key_name)


def test_gallery_builtins_preview_and_upload(x14_template):
    c = TestClient(create_app())
    gal = c.get("/api/templates").json()
    assert {t["name"] for t in gal if t["origin"] == "builtin"} == {"Hóa đơn bán hàng", "Ngân sách chi phí"}
    assert len(c.get("/api/templates").json()) == 2                      # built-ins are created once
    tid = gal[0]["id"]
    assert c.get(f"/api/templates/{tid}/preview").json()["sheets"]
    sl = c.get(f"/api/templates/{tid}/slots").json()
    assert sl["tables"][0]["input_rows"] >= 12 and sl["kv"] and sl["hidden_sheets"] == []
    r = c.post("/api/templates", files={"file": ("Status.xlsx", x14_template.read_bytes(), config_mime())},
               data={"name": "Status"})
    assert r.status_code == 200 and r.json()["origin"] == "user" and "Lists" in r.json()["sheets"]
    bad = c.post("/api/templates", files={"file": ("x.csv", b"a,b\n1,2\n", "text/csv")})
    assert bad.status_code == 400
    assert c.get("/api/templates/tpl_0000000000000000/preview").status_code == 404


def test_fill_builtin_invoice_from_sources():
    c = TestClient(create_app())
    ws = _ws(c, files=[("items.csv", ITEMS_CSV), ("head.txt", HEAD_TXT)])
    run = _run(c, ws, {"prompt": "Điền hóa đơn", "template_id": _builtin(c, "Hóa đơn bán hàng")})
    tools = [e["data"]["tool"] for e in run.events if e["event"] == "tool_completed"]
    assert tools == ["read_sources", "inspect_template", "map_sources", "apply_workbook_edits"]
    wid = run.result["workbook_id"]
    meta = c.get(f"/api/workbooks/{wid}/metadata").json()
    assert run.result["revision"] == 2 and meta["origin"] == "template"
    data = c.get(f"/api/workbooks/{wid}/export").content
    s = load_workbook(io.BytesIO(data))["Hóa đơn"]
    assert (s["B3"].value, s["B5"].value) == ("Công ty ABC", "HD-001")
    assert s["B6"].value.date().isoformat() == "2026-09-15"
    assert [s[f"B{r}"].value for r in (9, 10, 11)] == ["Bút bi", "Vở", "Thước"]
    assert [s[f"C{r}"].value for r in (9, 10, 11)] == [10, 4, 7] and s["D10"].value == 12000
    assert [s[f"A{r}"].value for r in (9, 10, 11, 12)] == [1, 2, 3, None]   # derived row numbers only
    assert s["E9"].value == '=IF(C9="","",C9*D9)' and s["E21"].value == "=SUM(E9:E20)"   # formulas intact
    assert s["B12"].value is None and s["C22"].value is None                # no invented VAT / extra rows
    # styles kept: the filled cell keeps the template's input style
    orig = load_workbook(service.revision_path(wid, 1))["Hóa đơn"]
    assert s["C9"].number_format == orig["C9"].number_format == "#,##0"
    assert s["C9"].fill.fgColor.rgb == orig["C9"].fill.fgColor.rgb
    rep = service.get_metadata(wid)["fill_report"]
    assert any(r["cell"] == "B3" and r["locator"] == "line:1" for r in rep)
    assert any(r["cell"] == "C9" and r["fact_id"] for r in rep)
    assert any("VAT" in w for w in run.result["warnings"])                 # unfilled label is reported


def test_fill_uploaded_x14_template_preserves_parts(x14_template):
    c = TestClient(create_app())
    raw = x14_template.read_bytes()
    tasks = "Task|Status|Hours\nDesign|Done|12\nBuild|Open|30\n".encode()
    ws = _ws(c, files=[("tasks.txt", tasks), ("owner.txt", "Owner: Lan\n".encode())],
             template=("status.xlsx", raw))
    run = _run(c, ws, {"prompt": "Fill the status report"})
    wid = run.result["workbook_id"]
    data = c.get(f"/api/workbooks/{wid}/export").content
    wb = load_workbook(io.BytesIO(data))
    r = wb["Report"]
    assert [r["A3"].value, r["B3"].value, r["C3"].value, r["D3"].value] == ["Design", "Done", 12, "=C3*25"]
    assert r["C8"].value == "=SUM(C3:C7)" and wb["Summary"]["B3"].value == "Lan"
    touched = set(service.get_metadata(wid)["revisions"][-1]["touched_parts"])
    with zipfile.ZipFile(io.BytesIO(raw)) as a, zipfile.ZipFile(io.BytesIO(data)) as b:
        for n in a.namelist():
            if n not in touched:
                assert a.read(n) == b.read(n), n                          # untouched members byte-equal
        assert b"x14:dataValidation" in b.read("xl/worksheets/sheet1.xml")
    fid = c.get(f"/api/workspaces/{ws}").json()["files"][-1]["id"]
    assert storage.sha256((storage.uploads_dir() / fid / "original.xlsx").read_bytes()) == storage.sha256(raw)


def test_template_without_sources_asks_and_uses_answers():
    c = TestClient(create_app())
    ws = _ws(c)
    tid = _builtin(c, "Ngân sách chi phí")
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Điền mẫu", "template_id": tid}).json()
    assert r["status"] == "needs_clarification" and [q["id"] for q in r["questions"]] == ["figures"]
    run = _run(c, ws, {"prompt": "Điền mẫu", "template_id": tid}, {"figures": "Kỳ: Quý 3/2026"})
    s = load_workbook(io.BytesIO(c.get(f"/api/workbooks/{run.result['workbook_id']}/export").content))["Ngân sách"]
    assert s["B3"].value == "Quý 3/2026" and s["B6"].value is None
    # empty answer -> unchanged template, honestly reported
    run = _run(c, ws, {"prompt": "Điền mẫu", "template_id": tid}, {})
    assert run.result["revision"] == 1 and any("unchanged template" in w for w in run.result["warnings"])
    assert c.post(f"/api/workspaces/{ws}/analyze",
                  json={"prompt": "x", "template_id": "tpl_0000000000000000"}).status_code == 404


def test_overflow_and_unmatched_are_warned(tmp_path):
    p = tmp_path / "inv.xlsx"
    templates._invoice(p, rows=2)
    b = sources.SourceBundle()
    sources._read_csv(b, "file_x", "items.csv", ITEMS_CSV.decode())
    sources._read_csv(b, "file_y", "other.csv", "Mã,Kho\nA1,HN\n")
    fills, warns, mode = template_fill.plan_fill(templates.detect_slots(p), b, use_llm=False)
    assert mode == "deterministic"
    assert {f.cell for f in fills if f.label == "Sản phẩm"} == {"B9", "B10"}
    assert any("1 row(s) did not fit" in w for w in warns)
    assert any("other.csv" in w and "matched no template table" in w for w in warns)


def test_llm_mapping_pairs_ids_only(tmp_path, monkeypatch):
    p = tmp_path / "inv.xlsx"
    templates._invoice(p)
    slots = templates.detect_slots(p)
    b = sources.SourceBundle()
    b.add("Tên công ty mua", "Công ty XYZ", "file_x", "line:1")
    b.add("Ghi chú", "999", "file_x", "line:2")
    k = next(i for i, s in enumerate(slots.kv) if s.label == "Khách hàng")
    monkeypatch.setattr(template_fill, "_llm_pairs", lambda *a: template_fill.LlmPairs(pairs=[
        template_fill.LlmPair(slot=f"k{k}", target="f-001"),
        template_fill.LlmPair(slot="k99", target="f-002"),          # unknown slot: ignored
        template_fill.LlmPair(slot="k0", target="f-777")]))         # invented fact id: ignored
    fills, warns, mode = template_fill.plan_fill(slots, b, use_llm=True)
    assert mode.startswith("azure:")
    assert [(f.cell, f.value, f.via) for f in fills] == [("B3", "Công ty XYZ", "llm")]
    monkeypatch.setattr(template_fill, "_llm_pairs", lambda *a: (_ for _ in ()).throw(RuntimeError("down")))
    fills, warns, mode = template_fill.plan_fill(slots, b, use_llm=True)
    assert mode == "deterministic" and fills == [] and "Model unavailable" in warns[0]


def test_populated_template_is_detected_and_extractor_gated(wbs):
    s = templates.detect_slots(wbs)
    assert any(t.prefilled_rows > 100 and not t.rows for t in s.tables) and "Lists" in s.hidden_sheets
    c = TestClient(create_app())
    assert c.get("/api/health").json()["extractor_available"] is False
    r = c.post("/api/templates/extract", json={"template_id": _builtin(c, "Ngân sách chi phí")})
    assert r.status_code == 503 and r.json()["code"] == "MODEL_UNAVAILABLE"


def test_template_with_sources_asks_for_labels_the_sources_lack():
    c = TestClient(create_app())
    ws = _ws(c, files=[("items.csv", ITEMS_CSV), ("head.txt", HEAD_TXT)])
    tid = _builtin(c, "Hóa đơn bán hàng")
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Điền hóa đơn", "template_id": tid}).json()
    assert r["status"] == "needs_clarification" and [q["id"] for q in r["questions"]] == ["figures"]
    q = r["questions"][0]["question"]
    assert "VAT" in q and "Khách hàng" not in q                       # only labels the sources cannot fill
    vat = next(s for s in templates.detect_slots(templates.path(tid)).kv if "VAT" in s.label)
    run = _run(c, ws, {"prompt": "Điền hóa đơn", "template_id": tid}, {"figures": f"{vat.label}: 10%"})
    s = load_workbook(io.BytesIO(c.get(f"/api/workbooks/{run.result['workbook_id']}/export").content))[vat.sheet]
    assert s[vat.cell].value in (0.1, 10) and s["B3"].value == "Công ty ABC"
    rep = service.get_metadata(run.result["workbook_id"])["fill_report"]
    assert any(x["cell"] == vat.cell and x["locator"].startswith("answer:figures") for x in rep)


def test_ai_proposes_mapping_asks_guesses_and_generate_reuses_it(monkeypatch):
    """Arbitrary content vs arbitrary template: the AI points at facts / verbatim quotes, the user confirms guesses."""
    monkeypatch.setattr(config, "llm_mode", lambda: "azure")
    note = "Bên mua là Công ty TNHH Hoà Phát.\nNơi giao: 12 Lê Lợi, Q1\nSố HĐ: HD-009\nNgày: 15/09/2026\n".encode()

    def fake(kv, tabs, b, facts):
        k, fid = {s.label: i for i, s in enumerate(kv)}, {f.field: f.id for f in facts}
        return template_fill.LlmPairs(pairs=[
            template_fill.LlmPair(slot=f"k{k['Khách hàng']}", target="quote:Công ty TNHH Hoà Phát", reason="Bên mua"),
            template_fill.LlmPair(slot=f"k{k['Địa chỉ']}", target=fid["Nơi giao"], confident=False,
                                  reason="Nơi giao có thể là địa chỉ"),
            template_fill.LlmPair(slot=f"k{k['VAT (%)']}", target="quote:10%")])     # not in the sources: dropped
    monkeypatch.setattr(template_fill, "_llm_pairs", fake)
    c = TestClient(create_app())
    ws = _ws(c, files=[("items.csv", ITEMS_CSV), ("note.txt", note)])
    tid = _builtin(c, "Hóa đơn bán hàng")
    kv = {s.label: i for i, s in enumerate(templates.detect_slots(templates.path(tid)).kv)}
    r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Điền hóa đơn", "template_id": tid}).json()
    assert r["status"] == "needs_clarification"
    ids = [q["id"] for q in r["questions"]]
    assert ids == [f"map_k{kv['Địa chỉ']}", f"map_k{kv['Số hóa đơn']}", "figures"]
    q_addr = r["questions"][0]
    assert q_addr["default"].startswith("Nơi giao = 12 Lê Lợi") and "(để trống)" in q_addr["options"]
    assert q_addr["explanation"] == "Nơi giao có thể là địa chỉ"
    rows = {m["field"]: m for m in r["mapping"]}
    assert rows["Khách hàng"]["status"] == "auto" and rows["Khách hàng"]["via"] == "llm"
    assert rows["Khách hàng"]["value"] == "Công ty TNHH Hoà Phát" and "note.txt line 1" in rows["Khách hàng"]["source"]
    assert rows["Địa chỉ"]["status"] == "confirm" and rows["VAT (%)"]["status"] == "missing"
    assert rows["Sản phẩm"]["status"] == "auto" and "items.csv" in rows["Sản phẩm"]["source"]
    assert "VAT" in r["questions"][2]["question"]
    monkeypatch.setattr(template_fill, "_llm_pairs", lambda *a: (_ for _ in ()).throw(AssertionError("no 2nd call")))
    r = c.post(f"/api/workspaces/{ws}/clarify", json={"clarification_id": r["clarification_id"], "answers": {
        ids[0]: q_addr["default"], ids[1]: "(để trống)"}}).json()
    assert r["status"] == "ready"
    rows = {m["field"]: m for m in r["mapping"]}
    assert rows["Địa chỉ"]["status"] == "auto" and rows["Số hóa đơn"]["status"] == "missing"
    run = runs.wait(c.post(f"/api/workspaces/{ws}/generate", json={}).json()["run_id"])
    assert run.status == "done", run.events
    s = load_workbook(io.BytesIO(c.get(f"/api/workbooks/{run.result['workbook_id']}/export").content))["Hóa đơn"]
    assert (s["B3"].value, s["B4"].value, s["B5"].value) == ("Công ty TNHH Hoà Phát", "12 Lê Lợi, Q1", None)
    assert s["B9"].value == "Bút bi" and run.result["llm_mode"].startswith("azure")
    rep = {x["cell"]: x for x in service.get_metadata(run.result["workbook_id"])["fill_report"]}
    assert rep["B3"]["locator"].startswith("quote:") and rep["B3"]["fact_id"] is None


def test_quote_must_be_verbatim_in_sources():
    b = sources.SourceBundle()
    b.texts["f"], b.names["f"] = "Bên mua là  Công ty ABC.\n", "a.txt"
    assert template_fill.find_quote(b, "Công ty ABC") == ("f", 1)
    assert template_fill.find_quote(b, "là Công ty ABC") == ("f", 1)              # whitespace-normalized
    assert template_fill.find_quote(b, "Công ty XYZ") is None and template_fill.find_quote(b, "") is None
    assert template_fill.find_quote(b, "ty AB") is None                             # whole words only
    b.texts["c"], b.names["c"] = "Mã,Thuế\nA,0\n", "x.csv"
    assert template_fill.find_quote(b, "0") is None                                 # tables are not quotable


def test_table_with_different_headers_is_guessed_and_columns_asked():
    """No source column name matches: the leftover table is only a guess; nothing is filled until columns are chosen."""
    odd = "Mặt hàng,SL,Giá\nBút bi,10,5000\nVở,4,12000\n".encode()
    c = TestClient(create_app())
    tid = _builtin(c, "Hóa đơn bán hàng")
    for pick in (True, False):
        ws = _ws(c, files=[("odd.csv", odd), ("head.txt", HEAD_TXT)])
        r = c.post(f"/api/workspaces/{ws}/analyze", json={"prompt": "Điền hóa đơn", "template_id": tid}).json()
        assert r["status"] == "needs_clarification"
        cols = {q["question"].split("'")[1]: q for q in r["questions"] if q["id"].startswith("map_t")}
        assert {"Sản phẩm", "Số lượng", "Đơn giá"} <= set(cols)
        assert cols["Đơn giá"]["default"].startswith("Giá") and cols["Sản phẩm"]["default"] == "(để trống)"
        assert "đoán" in cols["Sản phẩm"]["explanation"]
        rows = {m["field"]: m for m in r["mapping"]}
        assert rows["Đơn giá"]["status"] == "confirm" and rows["Sản phẩm"]["status"] == "missing"
        opt = lambda q, h: next(o for o in q["options"] if o.startswith(h + " "))        # noqa: E731
        answers = {q["id"]: "(để trống)" for q in cols.values()}
        if pick:
            answers.update({cols["Sản phẩm"]["id"]: opt(cols["Sản phẩm"], "Mặt hàng"),
                            cols["Số lượng"]["id"]: opt(cols["Số lượng"], "SL"),
                            cols["Đơn giá"]["id"]: cols["Đơn giá"]["default"]})
        r = c.post(f"/api/workspaces/{ws}/clarify", json={"clarification_id": r["clarification_id"],
                                                           "answers": answers}).json()
        assert r["status"] == "ready", r
        run = runs.wait(c.post(f"/api/workspaces/{ws}/generate", json={}).json()["run_id"])
        assert run.status == "done", run.events
        s = load_workbook(io.BytesIO(c.get(f"/api/workbooks/{run.result['workbook_id']}/export").content))["Hóa đơn"]
        row9 = [s.cell(9, col).value for col in range(1, 6)]
        if pick:
            assert "Bút bi" in row9 and 10 in row9 and 5000 in row9
        else:
            assert "Bút bi" not in row9 and any("odd.csv" in w for w in run.result["warnings"])
        assert s["B3"].value == "Công ty ABC"
