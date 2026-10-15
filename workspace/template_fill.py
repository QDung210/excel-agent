"""Template mode: fill a template's input cells from source facts with `set_value` (spec 07: formulas/styles intact).

Flow: revision 1 = the template bytes as uploaded/picked (never modified), revision 2 = the filled copy made by
WorkbookService.apply_operations (OOXML patch + L0-L4 validation). Mapping is deterministic (normalized label/header
match); in azure mode an LLM may additionally pair *ids* (slot <-> fact / column <-> source header) or point at a
verbatim quote in a source text. Values always come from the SourceBundle, never from the model. At analyze time the
mapping is built as a proposal (propose/describe), shown to the user, uncertain items are asked, and generate fills
from the confirmed proposal (fills_from).
"""
from __future__ import annotations

import datetime as dt
import difflib
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field

from . import builder, config, mock_fill, runs, service, sources, templates, upload, workspaces
from .errors import InvalidOperation
from .schemas import OperationBatch, SetValue
from .storage import new_id

KV_MIN, HEAD_MIN = 0.8, 0.75
SEQ_HEADERS = {"stt", "no", "tt", "so tt", "#"}


def score(a: str, b: str) -> float:
    na, nb = templates.norm(a), templates.norm(b)
    if not na or not nb:
        return 0.0
    if na == nb:
        return 1.0
    ta, tb = set(na.split()), set(nb.split())
    jac = len(ta & tb) / len(ta | tb)
    return max(jac, difflib.SequenceMatcher(None, na, nb).ratio())


def _share_word(a: str, b: str) -> bool:
    return bool(set(templates.norm(a).split()) & set(templates.norm(b).split()))


@dataclass
class Fill:
    sheet: str
    cell: str
    label: str
    value: Any
    value_type: str
    fact_id: str | None
    locator: str
    via: str                      # label | header | sequence | llm | mock

    def action(self) -> SetValue:
        return SetValue(type="set_value", sheet=self.sheet, cell=self.cell, value=self.value,
                        value_type=self.value_type, source_fact_id=self.fact_id)

    def report(self) -> dict:
        return {"sheet": self.sheet, "cell": self.cell, "label": self.label, "value": self.value,
                "fact_id": self.fact_id, "locator": self.locator, "via": self.via}


def _typed(v: Any) -> tuple[Any, str]:
    if isinstance(v, bool):
        return v, "bool"
    if isinstance(v, (int, float)):
        return v, "number"
    if isinstance(v, str) and (d := builder.to_date(v)):
        return d.date().isoformat(), "date"
    return str(v), "string"


def _col_cells(t: templates.TableSlot, col: int) -> list[str]:
    from openpyxl.utils import get_column_letter
    return [f"{get_column_letter(col)}{r}" for r in t.rows]


def _match_table(t: templates.TableSlot, b: sources.SourceBundle, used: set[int],
                 pairs: dict[tuple[int, int], int] | None = None,
                 loose: bool = False) -> tuple[int | None, dict[int, int]]:
    """-> (source table index, {template col: source column index}). loose: accept any overlap."""
    best: tuple[int, int | None, dict[int, int]] = (0, None, {})
    for ti, st in enumerate(b.tables):
        if ti in used:
            continue
        m: dict[int, int] = {}
        taken: set[int] = set()
        cand = sorted(((score(c.header, h), c.col, hi) for c in t.inputs for hi, h in enumerate(st.headers)),
                      reverse=True)
        for s, col, hi in cand:
            if s >= HEAD_MIN and col not in m and hi not in taken:
                m[col], taken = hi, taken | {hi}
        for (tcol, hi), sti in (pairs or {}).items():
            if sti == ti and tcol not in m and hi not in taken and any(c.col == tcol for c in t.inputs):
                m[tcol], taken = hi, taken | {hi}
        need = 1 if loose else max(1, -(-min(len(t.inputs), len(st.headers)) // 2))
        if len(m) >= need and len(m) > best[0]:
            best = (len(m), ti, m)
    return best[1], best[2]


class LlmPair(BaseModel):
    slot: str = Field(description="slot id, e.g. k3 or t0c2")
    target: str = Field(description="for k* slots: a fact id (f-012), or 'quote:' + text copied verbatim from a source "
                                    "text when no fact fits; for t* slots: a source header id (s1h2)")
    confident: bool = Field(default=True, description="false when the pairing is plausible but the user should confirm")
    reason: str = Field(default="", max_length=200, description="short reason in the user's language (<= 15 words)")


class LlmPairs(BaseModel):
    pairs: list[LlmPair] = Field(default_factory=list, max_length=200)


LLM_SYSTEM = (
    "You map an Excel template's input slots to the user's source content (any layout, any language). For each slot, "
    "pick the source item that holds its value: a fact id, a source table header id, or for k* slots with no fitting "
    "fact a 'quote:' with the value copied EXACTLY (character for character) from one line of a source text. Only the "
    "value itself, never a label or a sentence around it. Never invent or compute values and never invent ids; leave a "
    "slot out when the sources do not contain its value. Set confident=false when the match is a guess the user "
    "should confirm (similar but not clearly the same meaning). Give a short reason. Return JSON.")


def _excerpts(b: sources.SourceBundle, limit: int = 9000) -> list[str]:
    out: list[str] = []
    budget = limit
    for fid, text in b.texts.items():
        if b.names.get(fid, "").lower().endswith(".csv") or budget <= 0:
            continue
        chunk = text[:min(4000, budget)]
        budget -= len(chunk)
        out.append(f"--- {b.names.get(fid, fid)}\n{chunk}")
    return out


def _llm_pairs(kv: list[templates.KvSlot], tabs: list[templates.TableSlot], b: sources.SourceBundle,
               facts: list) -> LlmPairs:
    lines = ["Template input slots:"]
    lines += [f"k{i}: label '{s.label}' ({s.sheet}!{s.cell})" for i, s in enumerate(kv)]
    lines += [f"t{ti}c{c.col}: table column '{c.header}' ({t.sheet})" for ti, t in enumerate(tabs) for c in t.inputs]
    lines.append("Source facts:")
    lines += [f"{f.id}: {f.field} = {str(f.value)[:40]}" for f in facts[:150]]
    lines.append("Source table headers (with the first row's value):")
    lines += [f"s{si}h{hi}: {h}" + (f" (e.g. {str(st.rows[0][hi])[:30]})" if st.rows and st.rows[0][hi] not in (None, "")
                                     else "") for si, st in enumerate(b.tables) for hi, h in enumerate(st.headers)]
    if ex := _excerpts(b):
        lines.append("Source texts (for quote: targets):")
        lines += ex
    last: Exception | None = None
    for _, llm in config.chat_models(timeout=60):
        try:
            return llm.with_structured_output(LlmPairs, method="json_schema").invoke(
                [("system", LLM_SYSTEM), ("user", "\n".join(lines))])
        except Exception as e:                # noqa: BLE001 - fall back to next model / deterministic only
            last = e
    raise RuntimeError(str(last))


# ---------------------------------------------------------------- proposal (analyze) -> fills (generate)
# A proposal is the JSON-able mapping decision for every input slot, shown to the user before anything is written:
#   kv[i]     = {sheet, cell, label, target: fact id | "quote:..." | None, via, confident, reason, candidates}
#   tables[x] = {sheet, header_row, source: source table index | None, cols: {col: {header, hi, via, confident, reason}}}
# Targets are references into the SourceBundle only; fills_from() re-validates them and reads values from the sources.
FUZZY_MIN = 0.5                         # below KV_MIN: suggested for confirmation, never filled unconfirmed
CAND_MIN = 0.45                         # weaker label similarity: offered as an option only


def _kv_facts(b: sources.SourceBundle) -> list:
    in_tables = {fid for t in b.tables for row in t.fact_ids for fid in row if fid}
    return [f for f in b.facts if f.id not in in_tables and not f.field.startswith(("number on line",
                                                                                    "number in prompt"))]


def find_quote(b: sources.SourceBundle, quote: str) -> tuple[str, int] | None:
    """(file id, 1-based line) of a source text line that contains `quote` verbatim (whitespace-normalized,
    whole words). CSV files are tables, not prose, and are not shown to the model, so they cannot be quoted."""
    q = " ".join((quote or "").split())
    if not q or len(q) > 300:
        return None
    pat = re.compile(rf"(?<!\w){re.escape(q)}(?!\w)")
    for fid, text in b.texts.items():
        if b.names.get(fid, "").lower().endswith(".csv"):
            continue
        for n, line in enumerate(text.splitlines(), start=1):
            if pat.search(" ".join(line.split())):
                return fid, n
    return None


def _quote_value(q: str) -> Any:
    return sources._cell_value(" ".join(q.split()))


def resolve(b: sources.SourceBundle, target: str | None, facts: dict) -> tuple[Any, str | None, str] | None:
    """target -> (raw value, fact id, locator), or None when it no longer points at the sources."""
    if not target:
        return None
    if target.startswith("quote:"):
        hit = find_quote(b, target[6:])
        return (_quote_value(target[6:]), None, f"quote:{hit[0]}:line:{hit[1]}") if hit else None
    f = facts.get(target)
    return (f.value, f.id, f.provenance.locator) if f else None


def _tabs(slots: templates.Slots) -> list[templates.TableSlot]:
    return [t for t in slots.tables if t.rows]


def propose(slots: templates.Slots, b: sources.SourceBundle, use_llm: bool) -> dict:
    warnings: list[str] = []
    mode = "deterministic"
    kv_facts = _kv_facts(b)
    byid = {f.id: f for f in kv_facts}
    tabs = _tabs(slots)
    llm_kv: dict[int, LlmPair] = {}
    llm_tab: dict[tuple[int, int, int], tuple[int, LlmPair]] = {}
    if use_llm and (slots.kv or tabs):
        try:
            res = _llm_pairs(slots.kv, tabs, b, kv_facts)
            for p in res.pairs:
                if (m := re.fullmatch(r"k(\d+)", p.slot)) and int(m[1]) < len(slots.kv):
                    if p.target in byid or (p.target.startswith("quote:") and find_quote(b, p.target[6:])):
                        llm_kv[int(m[1])] = p
                elif (m := re.fullmatch(r"t(\d+)c(\d+)", p.slot)) and (n := re.fullmatch(r"s(\d+)h(\d+)", p.target)):
                    si, hi = int(n[1]), int(n[2])
                    if int(m[1]) < len(tabs) and si < len(b.tables) and hi < len(b.tables[si].headers):
                        llm_tab[(int(m[1]), int(m[2]), hi)] = (si, p)
            mode = f"azure:{config.azure_deployment()}"
        except Exception as e:                # noqa: BLE001
            warnings.append(f"Model unavailable for mapping ({str(e)[:120]}); used exact label matching only.")
    # --- tables
    out_tabs: list[dict] = []
    used_tables: set[int] = set()
    matched: dict[int, tuple[int | None, dict[int, int], bool]] = {}
    for x, t in enumerate(tabs):                         # pass 1: tables whose headers clearly match
        pairs = {(c, hi): si for (tx, c, hi), (si, _) in llm_tab.items() if tx == x}
        ti, m = _match_table(t, b, used_tables, pairs)
        if ti is not None:
            used_tables.add(ti)
        matched[x] = (ti, m, False)
    for x, t in enumerate(tabs):                         # pass 2: guess a leftover source table, to confirm
        if matched[x][0] is not None:
            continue
        ti, m = _match_table(t, b, used_tables, loose=True)
        left = [i for i in range(len(b.tables)) if i not in used_tables]
        if ti is None and len(left) == 1 and sum(1 for y in matched.values() if y[0] is None) == 1:
            ti, m = left[0], {}
        if ti is not None:
            used_tables.add(ti)
            matched[x] = (ti, m, True)
    for x, t in enumerate(tabs):
        ti, m, guessed = matched[x]
        if guessed:                                      # weak column guesses by similarity, all to be confirmed
            taken = set(m.values())
            for sc, col, hi in sorted(((score(c.header, h), c.col, hi) for c in t.inputs
                                       for hi, h in enumerate(b.tables[ti].headers)), reverse=True):
                if sc >= FUZZY_MIN and col not in m and hi not in taken and _share_word(
                        next(c.header for c in t.inputs if c.col == col), b.tables[ti].headers[hi]):
                    m[col], taken = hi, taken | {hi}
        seq_col = next((c.col for c in t.inputs if templates.norm(c.header) in SEQ_HEADERS and c.col not in m), None)
        cols: dict[str, dict] = {}
        for c in t.inputs:
            d = {"header": c.header, "hi": None, "via": None, "confident": True, "reason": ""}
            if c.col in m:
                hi = m[c.col]
                if guessed:
                    d.update(hi=hi, via="fuzzy", confident=False)
                elif score(c.header, b.tables[ti].headers[hi]) >= HEAD_MIN:
                    d.update(hi=hi, via="header")
                else:
                    lp = llm_tab.get((x, c.col, hi), (None, None))[1]
                    d.update(hi=hi, via="llm", confident=bool(lp is None or lp.confident),
                             reason=(lp.reason if lp else "")[:200])
            elif c.col == seq_col:
                d.update(via="sequence")
            cols[str(c.col)] = d
        out_tabs.append({"sheet": t.sheet, "header_row": t.header_row, "source": ti, "cols": cols,
                         "confident": not guessed})
    # --- key/value: exact label match, then the model's pairs, then fuzzy suggestions to confirm
    used: set[str] = set()
    chosen: dict[int, dict] = {}
    cand = sorted(((score(s.label, f.field), i, f.id) for i, s in enumerate(slots.kv) for f in kv_facts), reverse=True)
    for sc, i, fid in cand:
        if sc >= KV_MIN and i not in chosen and fid not in used:
            chosen[i], used = {"target": fid, "via": "label", "confident": True, "reason": ""}, used | {fid}
    for i, p in sorted(llm_kv.items()):
        if i not in chosen and (p.target.startswith("quote:") or p.target not in used):
            chosen[i] = {"target": p.target, "via": "llm", "confident": p.confident, "reason": p.reason[:200]}
            used.add(p.target)
    for sc, i, fid in cand:
        if FUZZY_MIN <= sc < KV_MIN and i not in chosen and fid not in used:
            chosen[i], used = {"target": fid, "via": "fuzzy", "confident": False, "reason": ""}, used | {fid}
    out_kv: list[dict] = []
    for i, s in enumerate(slots.kv):
        d = {"sheet": s.sheet, "cell": s.cell, "label": s.label, "target": None, "via": None, "confident": True,
             "reason": ""} | chosen.get(i, {})
        taken = {c["target"] for j, c in chosen.items() if j != i and c["confident"]}
        top = [fid for sc, j, fid in cand if j == i and sc >= CAND_MIN and fid not in taken][:3]
        d["candidates"] = list(dict.fromkeys(([d["target"]] if d["target"] else []) + top))[:4]
        out_kv.append(d)
    return {"mode": mode, "warnings": warnings, "kv": out_kv, "tables": out_tabs}


def fills_from(slots: templates.Slots, b: sources.SourceBundle, prop: dict) -> tuple[list[Fill], list[str]]:
    """Turn a (possibly user-confirmed) proposal into fills. Unconfirmed guesses are not filled."""
    from openpyxl.utils import get_column_letter
    fills: list[Fill] = []
    warnings: list[str] = []
    kv_facts = _kv_facts(b)
    byid = {f.id: f for f in kv_facts}
    tabs = _tabs(slots)
    used_tables: set[int] = set()
    for x, t in enumerate(tabs):
        pts = prop.get("tables") or []
        pt = pts[x] if x < len(pts) else None
        if not pt or pt["sheet"] != t.sheet or pt["header_row"] != t.header_row:
            continue
        ti = pt.get("source")
        if ti is None or ti >= len(b.tables) or ti in used_tables or not pt.get("confident", True):
            continue
        used_tables.add(ti)
        st = b.tables[ti]
        m = {c.col: pt["cols"][str(c.col)]["hi"] for c in t.inputs
             if (d := pt["cols"].get(str(c.col))) and d.get("hi") is not None and d.get("confident", True)
             and d["hi"] < len(st.headers)}
        seq_col = next((c.col for c in t.inputs if (pt["cols"].get(str(c.col)) or {}).get("via") == "sequence"), None)
        n = min(len(st.rows), len(t.rows))
        if len(st.rows) > len(t.rows):
            warnings.append(f"{st.name}: {len(st.rows) - len(t.rows)} row(s) did not fit into '{t.sheet}' "
                            f"({len(t.rows)} input rows; inserting rows is not supported yet).")
        for i in range(n):
            r = t.rows[i]
            for col, hi in m.items():
                v = st.rows[i][hi]
                if v in (None, ""):
                    continue
                v, vt = _typed(v)
                fills.append(Fill(t.sheet, f"{get_column_letter(col)}{r}", st.headers[hi], v, vt,
                                  st.fact_ids[i][hi], f"{st.locators[i]},col:{st.headers[hi]}",
                                  pt["cols"][str(col)].get("via") or "header"))
            if seq_col:
                fills.append(Fill(t.sheet, f"{get_column_letter(seq_col)}{r}", "STT", i + 1, "number", None,
                                  "derived:row_number", "sequence"))
        unmapped = [c.header for c in t.inputs if c.col not in m and c.col != seq_col]
        if unmapped:
            warnings.append(f"'{t.sheet}' columns left blank (no matching source column): {unmapped}")
    for ti, st in enumerate(b.tables):
        if ti not in used_tables:
            warnings.append(f"Source table '{st.name}' matched no template table; its data was not used.")
    used: set[str] = set()
    done: set[int] = set()
    pkv = prop.get("kv") or []
    for i, s in enumerate(slots.kv):
        d = pkv[i] if i < len(pkv) else None
        if not d or (d["sheet"], d["cell"]) != (s.sheet, s.cell) or not d.get("confident", True):
            continue
        hit = resolve(b, d.get("target"), byid)
        if hit and (hit[1] is None or hit[1] not in used):
            v, vt = _typed(hit[0])
            fills.append(Fill(s.sheet, s.cell, s.label, v, vt, hit[1], hit[2], d.get("via") or "label"))
            used.add(hit[1] or "")
            done.add(i)
    # values the user typed as answers ('label: value'), for slots still empty
    typed = [f for f in kv_facts if f.status == "user_confirmed" and f.id not in used]
    for sc, i, fid in sorted(((score(s.label, f.field), i, f.id) for i, s in enumerate(slots.kv) if i not in done
                              for f in typed), reverse=True):
        if sc >= KV_MIN and i not in done and fid not in used:
            s, f = slots.kv[i], byid[fid]
            v, vt = _typed(f.value)
            fills.append(Fill(s.sheet, s.cell, s.label, v, vt, f.id, f.provenance.locator, "label"))
            done.add(i)
            used.add(fid)
    blank = [s.label for i, s in enumerate(slots.kv) if i not in done]
    if blank:
        warnings.append(f"{len(blank)} labelled input cell(s) left blank (no source value): " + ", ".join(blank[:8]))
    left = [f for f in kv_facts if f.id not in used]
    if left:
        warnings.append(f"{len(left)} source fact(s) were not placed: " + ", ".join(f.field for f in left[:8]))
    prefilled = [t for t in slots.tables if t.prefilled_rows and not t.rows]
    if prefilled:
        warnings.append(f"{len(prefilled)} table(s) in the template already contain data and were not changed. "
                        "Use 'Make blank template' to clear sample data first.")
    if slots.hidden_sheets:
        warnings.append(f"Hidden sheets were not filled: {slots.hidden_sheets}")
    return fills, warnings


def plan_fill(slots: templates.Slots, b: sources.SourceBundle, use_llm: bool) -> tuple[list[Fill], list[str], str]:
    prop = propose(slots, b, use_llm)
    fills, warnings = fills_from(slots, b, prop)
    return fills, prop["warnings"] + warnings, prop["mode"]


def _short(v: Any, n: int = 60) -> str:
    s = str(v)
    return s if len(s) <= n else s[:n - 1] + "…"


def describe(slots: templates.Slots, b: sources.SourceBundle, prop: dict) -> list[dict]:
    """One row per template input (label cell or table column) for the review screen."""
    byid = {f.id: f for f in _kv_facts(b)}
    rows: list[dict] = []
    for d in prop.get("kv") or []:
        hit = resolve(b, d.get("target"), byid)
        src = None
        if hit and hit[1]:
            f = byid[hit[1]]
            src = f"'{f.field}' ({b.names.get(f.provenance.file_id or '', 'answer')} {f.provenance.locator})"
        elif hit:
            fid, line = hit[2].split(":")[1], hit[2].rsplit(":", 1)[1]
            src = f"quote ({b.names.get(fid, fid)} line {line})"
        rows.append({"target": f"{d['sheet']}!{d['cell']}", "field": d["label"], "source": src,
                     "value": _short(hit[0]) if hit else None, "via": d.get("via") if hit else None,
                     "status": "missing" if not hit else ("auto" if d.get("confident", True) else "confirm"),
                     "reason": d.get("reason") or ""})
    tabs = _tabs(slots)
    for x, pt in enumerate(prop.get("tables") or []):
        if x >= len(tabs):
            break
        t = tabs[x]
        st = b.tables[pt["source"]] if pt.get("source") is not None and pt["source"] < len(b.tables) else None
        from openpyxl.utils import get_column_letter
        for col, d in pt["cols"].items():
            rng = f"{t.sheet}!{get_column_letter(int(col))}{t.rows[0]}:{get_column_letter(int(col))}{t.rows[-1]}"
            if d.get("via") == "sequence":
                rows.append({"target": rng, "field": d["header"], "source": "row number", "value": "1, 2, 3…",
                             "via": "sequence", "status": "auto" if pt.get("confident", True) else "confirm", "reason": ""})
                continue
            hi = d.get("hi")
            if st is None or hi is None or hi >= len(st.headers):
                rows.append({"target": rng, "field": d["header"], "source": None, "value": None, "via": None,
                             "status": "missing", "reason": ""})
                continue
            sample = next((r[hi] for r in st.rows if r[hi] not in (None, "")), None)
            n = min(len(st.rows), len(t.rows))
            rows.append({"target": rng, "field": d["header"], "source": f"column '{st.headers[hi]}' ({st.name}, {n} rows)",
                         "value": _short(sample) if sample is not None else None, "via": d.get("via"),
                         "status": "auto" if d.get("confident", True) else "confirm", "reason": d.get("reason") or ""})
    return rows


def _template_source(state: dict) -> tuple[Path, str]:
    if state.get("template_file_id"):
        f = upload.get(state["template_file_id"])
        return f.path, f.filename
    if state.get("template_id"):
        meta = templates.get(state["template_id"])
        return templates.path(state["template_id"]), meta["name"] + ".xlsx"
    raise InvalidOperation("No template selected.")


def run(run: runs.Run, ws: dict, state: dict) -> dict:
    from .generate import _filename, _step
    prompt, answers = state["prompt"], state.get("answers") or {}
    src, tpl_name = _template_source(state)
    b = _step(run, "read_sources", lambda: sources.load(state["source_file_ids"], prompt),
              lambda b: f"{len(b.facts)} facts, {len(b.tables)} tables")
    sources.answer_figures(b, answers.get("figures") or "", keep_text=True)
    if state.get("mock"):                     # mock data the user asked for and reviewed: re-validated cell by cell
        def _mock() -> tuple[list[Fill], list[str], str]:
            if answers.get("mock_apply", "apply") != "apply":
                return [], ["Mock data was declined; the template is unchanged."], "mock"
            cells, mw = mock_fill.parse(src, answers.get("mock_cells") or "")
            out = [Fill(s, c, "mock", v, "number" if isinstance(v, (int, float)) else "string", None,
                        "mock:ai_generated", "mock") for s, c, v in cells]
            if out:
                mw.append(f"{len(out)} cell(s) contain MOCK data generated by the AI on your request: fictional, "
                          "not real data.")
            return out, mw, config.active_model() or "mock"
        fills, warns, mode = _step(run, "mock_data", _mock, lambda o: f"{len(o[0])} mock cell(s)")
    else:
        fills, warns, mode = _map(run, state, src, b)
    meta = service.create_from_path(src, _filename(state, Path(tpl_name).stem), "template", ws["id"],
                                    extra={"template_file_id": state.get("template_file_id"),
                                           "template_id": state.get("template_id"), "llm_mode": mode,
                                           "mock_data": bool(state.get("mock") and fills)})
    return _write(run, ws, state, b, meta, fills, warns, mode)


def _map(run: runs.Run, state: dict, src: Path, b: sources.SourceBundle) -> tuple[list[Fill], list[str], str]:
    from .generate import _step
    slots = _step(run, "inspect_template", lambda: templates.detect_slots(src),
                  lambda s: f"{len(s.kv)} label slot(s), {sum(1 for t in s.tables if t.rows)} input table(s)")
    if state.get("proposal"):                 # mapping decided at analyze and confirmed by the user: no new LLM call
        prop = state["proposal"]
        fills, warns, mode = _step(run, "map_sources", lambda: (*fills_from(slots, b, prop), prop["mode"]),
                                   lambda o: f"{len(o[0])} cell(s) mapped (confirmed plan, {o[2]})")
    else:
        fills, warns, mode = _step(run, "map_sources", lambda: plan_fill(slots, b, config.llm_mode() == "azure"),
                                   lambda o: f"{len(o[0])} cell(s) mapped ({o[2]})")
    return fills, warns, mode


def _write(run: runs.Run, ws: dict, state: dict, b: sources.SourceBundle, meta: dict, fills: list[Fill],
           warns: list[str], mode: str) -> dict:
    from .generate import _step
    wb_id, rev, changed = meta["id"], 1, []
    report = [f.report() for f in fills]
    if fills:
        actions = [f.action() for f in fills]
        chunks = [actions[i:i + config.MAX_ACTIONS_PER_BATCH] for i in range(0, len(actions), config.MAX_ACTIONS_PER_BATCH)]
        try:
            for k, chunk in enumerate(chunks):
                res = _step(run, "apply_workbook_edits", lambda c=chunk: service.apply_operations(wb_id, OperationBatch(
                    base_revision=rev, operation_id=f"fill-{new_id('op')}-{k}", actions=c, origin="system",
                    note="mock data (AI-generated on request)" if state.get("mock") else "template fill from sources")), lambda r: f"revision {r.revision}, {len(r.changed_ranges)} range(s)")
                rev, changed = res.revision, changed + res.changed_ranges
                warns += res.validation.warnings
        except Exception as e:                # validation or OOXML failure: keep revision 1, report honestly
            msg = getattr(e, "message", str(e))
            run.emit("validation_failed", {"errors": [msg] + list(getattr(e, "extra", {}).get("errors", []))[:10]})
            return {"errors": [msg], "workbook_id": wb_id}
    else:
        warns.append("No mock value was written; the workbook is the unchanged template." if state.get("mock") else
                     "No template cell matched the sources; the workbook is the unchanged template.")
    warnings = list(dict.fromkeys(state.get("warnings", []) + b.warnings + warns))
    service.update_metadata(wb_id, warnings=warnings, fill_report=report[:2000])
    cur = workspaces.get(ws["id"])
    cur["workbook_id"] = wb_id
    cur["runs"] = (cur.get("runs") or [])[-20:] + [run.id]
    workspaces.save(cur)
    run.emit("workbook_revision", {"workbook_id": wb_id, "revision": rev, "changed_ranges": changed[:200]})
    return {"workbook_id": wb_id, "revision": rev, "warnings": warnings, "llm_mode": mode, "filled": len(fills)}
