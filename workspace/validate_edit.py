"""Validation L0-L4 for a value-edit revision (spec 07). Failure means the revision is not published."""
from __future__ import annotations

import io
import re
import warnings
import zipfile
from html import unescape
from typing import Any
from xml.etree import ElementTree as ET

from xlsx_template import ooxml

from .ooxml_edit import SST_PART, EditStats

ALLOWED_TOUCH = re.compile(r"^(xl/worksheets/[^/]+\.xml|xl/sharedStrings\.xml|xl/workbook\.xml)$")
T_RE = re.compile(r"<t(?:\s[^>]*)?>(.*?)</t>", re.S)


def shared_strings(z: zipfile.ZipFile) -> list[str]:
    if SST_PART not in z.namelist():
        return []
    xml = z.read(SST_PART).decode("utf-8")
    return [unescape("".join(T_RE.findall(si or ""))) for si in re.findall(r"<si>(.*?)</si>|<si/>", xml, re.S)]


def read_cells(xml: str, sst: list[str]) -> dict[str, tuple[Any, str | None, str | None]]:
    """ref -> (value, style, formula). Formula cells report value None (caches may be stripped)."""
    out = {}
    _, inner, _ = ooxml.split_sheet_data(xml)
    for m in ooxml.CELL_RE.finditer(inner):
        a, body = ooxml.attrs(m.group(1)), m.group(2) or ""
        ref = a.get("r")
        if ref is None:
            continue
        fm = re.search(r"<f\b[^>]*?(?:/>|>(.*?)</f>)", body, re.S)
        if fm:
            out[ref] = (None, a.get("s"), unescape(fm.group(0)))
            continue
        vm = re.search(r"<v(?:\s[^>]*)?>(.*?)</v>", body, re.S)
        t = a.get("t", "n")
        if t == "inlineStr":
            val: Any = unescape("".join(T_RE.findall(body))) if "<is>" in body else None
        elif vm is None:
            val = None
        elif t == "s":
            val = sst[int(vm.group(1))]
        elif t == "b":
            val = vm.group(1).strip() in ("1", "true")
        elif t in ("str", "e"):
            val = unescape(vm.group(1))
        else:
            val = float(vm.group(1))
        out[ref] = (val, a.get("s"), None)
    return out


def _same(a: Any, b: Any) -> bool:
    if isinstance(a, bool) or isinstance(b, bool):
        return a is b or a == b and type(a) is type(b)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) <= 1e-9 * max(1.0, abs(float(a)))
    return a == b


def validate(src: str, dst: str, stats: EditStats) -> dict:
    errors: list[str] = []
    checks: dict[str, int] = {}

    # L0 zip integrity + XML well-formed
    try:
        zd = zipfile.ZipFile(dst)
    except zipfile.BadZipFile:
        return {"pass": False, "level": "L0", "errors": ["output is not a valid zip"], "warnings": [], "checks": {}}
    with zipfile.ZipFile(src) as zs, zd:
        names_s, names_d = zs.namelist(), zd.namelist()
        if zd.testzip() is not None:
            errors.append("L0: corrupt zip member")
        if len(names_d) != len(set(names_d)):
            errors.append("L0: duplicate zip members")
        for n in names_d:
            if n.endswith((".xml", ".rels")):
                try:
                    ET.fromstring(zd.read(n))
                except ET.ParseError as e:
                    errors.append(f"L0: {n} is not well-formed XML ({e})")
        checks["L0_parts"] = len(names_d)

        # L1 untouched members byte-equal, touched within allowlist
        if names_s != names_d:
            errors.append("L1: zip member list changed")
        for n in names_s:
            if n not in names_d:
                continue
            a, b = zs.read(n), zd.read(n)
            if a != b:
                if not ALLOWED_TOUCH.match(n):
                    errors.append(f"L1: {n} changed but is not an editable part")
                elif n not in stats.touched_parts:
                    errors.append(f"L1: {n} changed without being reported")
        checks["L1_untouched_equal"] = len(names_s) - len(stats.touched_parts)

        # L2 structure outside sheetData, formulas, styles, workbook metadata
        sheets = ooxml.sheet_parts(zs)
        sst_s, sst_d = shared_strings(zs), shared_strings(zd)
        if sst_d[: len(sst_s)] != sst_s:
            errors.append("L2: existing shared strings were modified or renumbered")
        if "xl/workbook.xml" in stats.touched_parts:
            strip = lambda x: re.sub(r"<calcPr\b[^>]*/>|<calcPr\b[^>]*>.*?</calcPr>", "", x, flags=re.S)
            if strip(zs.read("xl/workbook.xml").decode()) != strip(zd.read("xl/workbook.xml").decode()):
                errors.append("L2: workbook.xml changed outside calcPr")
        formulas = 0
        edited = {s for (s, _) in stats.intended}
        for name, part in sheets.items():
            if part not in stats.touched_parts:
                continue
            xs, xd = zs.read(part).decode("utf-8"), zd.read(part).decode("utf-8")
            bs, _, as_ = ooxml.split_sheet_data(xs)
            bd, _, ad = ooxml.split_sheet_data(xd)
            dim = lambda x: re.sub(r'(<dimension\b[^>]*\bref=")[^"]*"', r'\1"', x)
            if dim(bs) != dim(bd) or as_ != ad:
                errors.append(f"L2: {name}: content outside sheetData changed")
            cs, cd = read_cells(xs, sst_s), read_cells(xd, sst_d)
            targets = {ref for (s, ref) in stats.intended if s == name}
            for ref, (val, style, f) in cs.items():
                if f is not None:
                    formulas += 1
                    if ref not in cd or cd[ref][2] != f:
                        errors.append(f"L2: {name}!{ref}: formula changed")
                if ref not in targets and ref in cd and cd[ref][1] != style:
                    errors.append(f"L2: {name}!{ref}: style changed")
                if ref in targets and ref in cd and cd[ref][1] != style:
                    errors.append(f"L2: {name}!{ref}: target style changed")

            # L3 edit semantics: targets hold intended values, everything else unchanged
            if name in edited:
                for ref in targets:
                    want = stats.intended[(name, ref)]
                    got = cd.get(ref, (None, None, None))[0]
                    if not _same(got, want):
                        errors.append(f"L3: {name}!{ref}: expected {want!r}, found {got!r}")
            for ref, (val, _, f) in cs.items():
                if ref in targets or f is not None:
                    continue
                if not _same(cd.get(ref, (None, None, None))[0], val):
                    errors.append(f"L3: {name}!{ref}: untargeted value changed")
            for ref in cd:
                if ref not in cs and ref not in targets and cd[ref][0] is not None:
                    errors.append(f"L3: {name}!{ref}: unexpected new value")
        checks["L2_formulas_checked"] = formulas
        checks["L3_targets"] = len(stats.intended)

    # L4 independent parser opens it (load only, never save)
    if not errors:
        try:
            from openpyxl import load_workbook
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                wb = load_workbook(io.BytesIO(open(dst, "rb").read()))
            checks["L4_sheets"] = len(wb.sheetnames)
            wb.close()
        except Exception as e:  # noqa: BLE001
            errors.append(f"L4: openpyxl could not open the result ({type(e).__name__}: {e})")
    level = "L4" if not errors else min((e[:2] for e in errors), default="L4")
    return {"pass": not errors, "level": level, "errors": errors[:50], "warnings": list(stats.warnings), "checks": checks}
