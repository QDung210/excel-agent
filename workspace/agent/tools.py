"""The agent's tool allowlist (spec 06). Each tool is a closure over one turn's context.

Safety rules enforced here, not in the prompt:
- writes only via WorkbookService.apply_operations (revision guard, L0-L4 validation, new revision);
- edits of more than AUTO_APPLY_MAX_CELLS cells become a pending plan that the user must approve;
- a numeric value may be written only if it is grounded: typed by the user in this thread, a source fact of the
  workbook's workspace, or a value already read from the workbook in this turn;
- text starting with "=" is refused (the editor cannot write formulas yet) instead of being stored as text;
- a conversion/scaling of existing numbers (e.g. VND -> USD at a rate the user typed) is computed in code by
  transform_range, never by the model, so the derived values are exact and traceable to grounded inputs;
- no filesystem, shell or network access; reads are bounded (MAX_READ_CELLS, source chunk size).
"""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from langchain_core.tools import StructuredTool
from pydantic import BaseModel, Field

from .. import config, ooxml_edit, runs, service, sources, templates, upload, validate_edit, workspaces
from ..errors import InvalidOperation, NotFound, WorkspaceError
from ..schemas import ClearValue, OperationBatch, SetValue, parse_range
from . import store

AUTO_APPLY_MAX_CELLS = 25
MAX_PLAN_CELLS = 1000
MAX_SOURCE_CHARS = 6000
CONFIRM_RE = re.compile(r"^\s*(?:đồng ý|dong y|ok(?:ay)?|yes|apply|áp dụng|ap dung|xác nhận|xac nhan|confirm|được|duoc)\b",
                        re.I)


@dataclass
class Ctx:
    run: runs.Run
    workbook_id: str
    revision: int
    thread_id: str
    message: str
    user_text: str                      # every user message of the thread (number grounding)
    active_sheet: str | None = None
    vi: bool = False
    allowed: set[float] = field(default_factory=set)
    texts: list[str] = field(default_factory=list)
    mutations: int = 0
    pending_plan_id: str | None = None
    tool_errors: list[str] = field(default_factory=list)
    new_workbook_id: str | None = None
    _books: dict = field(default_factory=dict)
    _grounding_loaded: bool = False
    _sources: Any = None

    def t(self, vi: str, en: str) -> str:
        return vi if self.vi else en


# ---------------------------------------------------------------- helpers
def numbers_in(text: str) -> set[float]:
    out: set[float] = set()
    for tok in sources.NUM_TOKEN_RE.findall(text or ""):
        v, unit = sources.parse_number(tok)
        if v is not None:
            out.add(round(v, 9))
            if unit == "%":
                out.add(round(v * 100, 9))
    return out


def _book(ctx: Ctx, data_only: bool):
    key = (ctx.revision, data_only)
    if key not in ctx._books:
        from openpyxl import load_workbook
        ctx._books[key] = load_workbook(service.revision_path(ctx.workbook_id, ctx.revision), data_only=data_only)
    return ctx._books[key]


def _sheet(ctx: Ctx, name: str | None):
    wb = _book(ctx, False)
    name = (name or ctx.active_sheet or wb.sheetnames[0]).strip().strip("'")
    for n in wb.sheetnames:
        if n.lower() == name.lower():
            return n
    raise InvalidOperation(f"Sheet {name!r} does not exist. Sheets: {wb.sheetnames}")


def _plain(v: Any) -> Any:
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:10] if isinstance(v, dt.datetime) and v.time() == dt.time() else v.isoformat()
    if isinstance(v, dt.time):
        return v.isoformat()
    return v


def _workspace_sources(ctx: Ctx) -> sources.SourceBundle | None:
    if ctx._sources is None:
        ws_id = service.get_metadata(ctx.workbook_id).get("workspace_id")
        ids = []
        if ws_id:
            try:
                ids = [f["id"] for f in workspaces.get(ws_id)["files"] if f.get("role") == "source"]
            except NotFound:
                ids = []
        ctx._sources = sources.load(ids) if ids else False
    return ctx._sources or None


def _load_grounding(ctx: Ctx) -> None:
    if not ctx._grounding_loaded:
        ctx._grounding_loaded = True
        ctx.allowed |= numbers_in(ctx.user_text)
        b = _workspace_sources(ctx)
        if b:
            ctx.allowed |= b.allowed_numbers()
            ctx.texts += list(b.texts.values())


def _grounded(ctx: Ctx, v: float) -> bool:
    _load_grounding(ctx)
    return round(float(v), 9) in ctx.allowed


def _expand(ref: str) -> list[str]:
    from openpyxl.utils import get_column_letter
    r1, c1, r2, c2 = parse_range(ref)
    return [f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def _cell_count(ref: str) -> int:
    r1, c1, r2, c2 = parse_range(ref)
    return (r2 - r1 + 1) * (c2 - c1 + 1)


class _Tool:
    """Wraps a tool body: emits tool_started/tool_completed, logs (no values), turns errors into tool output."""

    def __init__(self, ctx: Ctx, name: str):
        self.ctx, self.name = ctx, name

    def __call__(self, fn):
        def wrapped(**kw):
            ctx = self.ctx
            ctx.run.emit("tool_started", {"tool": self.name, "args": _args_summary(kw)})
            try:
                out = fn(**kw)
                status = "error" if out.get("error") else "ok"
            except WorkspaceError as e:
                out, status = {"error": e.message, "code": e.code, **e.extra}, "error"
            if status == "error":
                ctx.tool_errors.append(f"{self.name}: {out.get('error')}"[:300])
            summary = out.get("summary") or out.get("error") or "ok"
            ctx.run.emit("tool_completed", {"tool": self.name, "status": status, "summary": str(summary)[:200]})
            store.log(ctx.thread_id, {"turn": ctx.run.id, "tool": self.name, "status": status,
                                      "revision": ctx.revision, "args": _args_summary(kw)})
            out.pop("summary", None)
            return json.dumps(out, ensure_ascii=False, default=str)[:12000]
        return wrapped


def _args_summary(kw: dict) -> dict:
    """Addresses and counts only; cell values and source text are not logged."""
    out = {}
    for k, v in kw.items():
        if k == "edits" and isinstance(v, list):
            out["edits"] = len(v)
            out["cells"] = [str(getattr(e, "cell", None) or (e.get("cell") if isinstance(e, dict) else ""))
                            for e in v][:20]
        elif k in ("sheet", "range", "file_id", "plan_id", "start_line", "max_lines", "revision", "op", "factor"):
            out[k] = v
    return out


# ---------------------------------------------------------------- arg schemas
class Edit(BaseModel):
    op: Literal["set", "clear"] = Field("set", description="'set' writes a value, 'clear' empties the cell(s)")
    sheet: str | None = Field(None, description="Sheet name; defaults to the active sheet")
    cell: str = Field(description="A1 address. For op='clear' a range like A2:C9 is allowed.")
    value: str | float | int | bool | None = Field(None, description="Value for op='set'. Never a formula.")
    value_type: Literal["auto", "string", "number", "bool", "date"] = Field(
        "auto", description="'date' expects ISO YYYY-MM-DD; 'auto' turns plain numerals into numbers")


class ProposeArgs(BaseModel):
    edits: list[Edit] = Field(min_length=1, max_length=MAX_PLAN_CELLS)
    reason: str = Field("", max_length=300, description="One line: why these edits")


class TransformArgs(BaseModel):
    range: str = Field(description="A1 range of the numbers to transform, e.g. G9:G20 (header cells are skipped)")
    sheet: str | None = Field(None, description="Sheet name; defaults to the active sheet")
    op: Literal["multiply", "divide", "add", "subtract"] = Field(description="e.g. 'divide' by the rate for VND->USD")
    factor: float = Field(description="The rate/amount exactly as the user typed it, e.g. 25000")
    round_digits: int | None = Field(2, ge=0, le=10, description="Decimals to round the results to; null = no rounding")
    reason: str = Field("", max_length=300)


class RangeArgs(BaseModel):
    range: str = Field(description="A1 range, e.g. A1:F20 (at most 400 cells)")
    sheet: str | None = Field(None, description="Sheet name; defaults to the active sheet")


class SourceArgs(BaseModel):
    file_id: str | None = Field(None, description="Omit to list the workspace's source files")
    start_line: int = Field(1, ge=1)
    max_lines: int = Field(80, ge=1, le=200)


class PlanArgs(BaseModel):
    plan_id: str


class BuildArgs(BaseModel):
    brief: str = Field(max_length=2000, description="What the new workbook should contain")


class NoArgs(BaseModel):
    pass


class ValidateArgs(BaseModel):
    revision: int | None = Field(None, ge=1)


# ---------------------------------------------------------------- tool bodies
def make_tools(ctx: Ctx) -> list[StructuredTool]:

    @_Tool(ctx, "inspect_workbook")
    def inspect_workbook() -> dict:
        wb = _book(ctx, False)
        sheets = []
        for ws in wb.worksheets:
            formulas = 0
            for row in ws.iter_rows(max_row=min(ws.max_row, 5000)):
                formulas += sum(1 for c in row if isinstance(c.value, str) and c.value.startswith("="))
            sheets.append({"name": ws.title, "dimensions": ws.dimensions, "max_row": ws.max_row,
                           "max_col": ws.max_column, "hidden": ws.sheet_state != "visible",
                           "merged_ranges": len(ws.merged_cells.ranges), "formulas": formulas})
        slots = templates.detect_slots(service.revision_path(ctx.workbook_id, ctx.revision))
        meta = service.get_metadata(ctx.workbook_id)
        prev = service.get_preview(ctx.workbook_id, ctx.revision)
        return {"workbook": meta["name"], "revision": ctx.revision, "sheets": sheets,
                "input_slots": {"kv": [f"{k.sheet}!{k.cell} ({k.label})" for k in slots.kv][:40],
                                "tables": slots.summary()["tables"][:10]},
                "warnings": prev.get("warnings", [])[:20],
                "editor_limits": "set/clear cell values only; no formulas, styles, row/column insert, charts",
                "summary": f"{len(sheets)} sheet(s)"}

    @_Tool(ctx, "read_range")
    def read_range(range: str, sheet: str | None = None) -> dict:
        name = _sheet(ctx, sheet)
        n = _cell_count(range)
        if n > config.MAX_READ_CELLS:
            raise InvalidOperation(f"Range has {n} cells; read at most {config.MAX_READ_CELLS} at a time.")
        f_ws, v_ws = _book(ctx, False)[name], _book(ctx, True)[name]
        cells, empty = [], 0
        for a in _expand(range):
            raw, val = f_ws[a].value, v_ws[a].value
            if raw is None:
                empty += 1
                continue
            item = {"cell": a, "value": _plain(val)}
            if isinstance(raw, str) and raw.startswith("="):
                item["formula"] = raw
            cells.append(item)
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                ctx.allowed.add(round(float(val), 9))
            elif isinstance(val, str):
                ctx.texts.append(val)
        return {"sheet": name, "range": range.upper(), "cells": cells, "empty_cells": empty,
                "summary": f"{len(cells)} non-empty cell(s)"}

    @_Tool(ctx, "read_source")
    def read_source(file_id: str | None = None, start_line: int = 1, max_lines: int = 80) -> dict:
        b = _workspace_sources(ctx)
        if not b:
            return {"error": ctx.t("Workbook này không có file nguồn txt/csv.", "This workbook has no txt/csv sources.")}
        if not file_id:
            files = [{"file_id": k, "filename": b.names[k], "lines": len(v.splitlines())} for k, v in b.texts.items()]
            return {"files": files, "facts": len(b.facts), "summary": f"{len(files)} source file(s)"}
        if file_id not in b.texts:
            raise NotFound("Source file is not attached to this workbook's workspace.")
        lines = b.texts[file_id].splitlines()
        chunk, size = [], 0
        for i in range(start_line - 1, min(len(lines), start_line - 1 + max_lines)):
            size += len(lines[i])
            if size > MAX_SOURCE_CHARS:
                break
            chunk.append({"line": i + 1, "text": lines[i]})
        for c in chunk:
            ctx.allowed |= numbers_in(c["text"])
            ctx.texts.append(c["text"])
        facts = [{"id": f.id, "field": f.field, "value": f.value, "locator": f.provenance.locator}
                 for f in b.facts if f.provenance.file_id == file_id][:60]
        return {"file": b.names[file_id], "total_lines": len(lines), "lines": chunk, "facts": facts,
                "summary": f"{len(chunk)} line(s)"}

    def _actions(edits: list[Edit]) -> list[SetValue | ClearValue]:
        acts: list[SetValue | ClearValue] = []
        bad: list[str] = []
        for e in edits:
            sheet = _sheet(ctx, e.sheet)
            if e.op == "clear":
                cells = _expand(e.cell) if ":" in e.cell else [e.cell]
                acts += [ClearValue(type="clear_value", sheet=sheet, cell=c) for c in cells]
                continue
            if ":" in e.cell:
                raise InvalidOperation("op='set' takes one cell; list each cell separately.")
            if isinstance(e.value, str) and e.value.strip().startswith("="):
                raise InvalidOperation(ctx.t("Trình sửa chưa hỗ trợ ghi công thức; không ghi gì.",
                                             "Writing formulas is not supported by the editor yet; nothing written."))
            act = SetValue(type="set_value", sheet=sheet, cell=e.cell, value=e.value, value_type=e.value_type)
            kind, v = ooxml_edit.coerce(act)
            if kind == "number" and not _grounded(ctx, v):
                bad.append(f"{sheet}!{act.cell}={e.value}")
            elif kind == "date" and not _date_in_user_text(ctx, str(e.value)):
                bad.append(f"{sheet}!{act.cell}={e.value}")
            acts.append(act)
        if bad:
            raise InvalidOperation(ctx.t(
                "Các giá trị số/ngày sau không có nguồn (người dùng chưa nhập, không có trong file nguồn hay "
                "workbook); hãy hỏi người dùng thay vì tự đặt: ",
                "These numbers/dates are not grounded (not typed by the user, not in the sources or the workbook); "
                "ask the user instead of inventing them: ") + ", ".join(bad[:10]) + ctx.t(
                ". Nếu đây là giá trị tính từ số trong workbook (quy đổi tiền tệ, nhân/chia theo tỷ lệ), dùng "
                "transform_range để code tự tính thay vì tự điền số.",
                ". If these are computed from workbook numbers (currency conversion, scaling by a rate), call "
                "transform_range so the code computes them instead of writing them yourself."),
                ungrounded=bad[:20])
        if len({(a.sheet, a.cell) for a in acts}) > MAX_PLAN_CELLS:
            raise InvalidOperation(f"At most {MAX_PLAN_CELLS} cells per plan.")
        return acts

    def _diff(acts) -> list[dict]:
        out = []
        for a in acts[:200]:
            before = _book(ctx, False)[a.sheet][a.cell].value
            out.append({"cell": f"{a.sheet}!{a.cell}", "before": _plain(before),
                        "after": getattr(a, "value", None)})
        return out

    def _commit(plan: dict) -> dict:
        res = service.apply_operations(ctx.workbook_id, OperationBatch(
            base_revision=plan["base_revision"], operation_id=plan["id"], origin="agent",
            note=("agent: " + (plan.get("reason") or "edit"))[:300], actions=plan["actions"]))
        plan.update(status="applied", revision=res.revision)
        store.save_plan(plan)
        ctx.revision, ctx.mutations = res.revision, ctx.mutations + 1
        ctx.run.emit("workbook_revision", {"workbook_id": ctx.workbook_id, "revision": res.revision,
                                           "changed_ranges": res.changed_ranges})
        return {"applied": True, "plan_id": plan["id"], "revision": res.revision,
                "changed_cells": res.changed_ranges[:50], "validation_warnings": res.validation.warnings[:10],
                "message": ctx.t(f"Đã cập nhật {len(res.changed_ranges)} ô (revision {res.revision}).",
                                 f"Updated {len(res.changed_ranges)} cell(s) (revision {res.revision})."),
                "summary": f"revision {res.revision}"}

    @_Tool(ctx, "propose_workbook_edits")
    def propose_workbook_edits(edits: list[Edit], reason: str = "") -> dict:
        edits = [e if isinstance(e, Edit) else Edit(**e) for e in edits]
        return _plan(_actions(edits), reason)

    @_Tool(ctx, "transform_range")
    def transform_range(range: str, op: str, factor: float, sheet: str | None = None, round_digits: int | None = 2,
                        reason: str = "") -> dict:
        name = _sheet(ctx, sheet)
        if _cell_count(range) > MAX_PLAN_CELLS:
            raise InvalidOperation(f"At most {MAX_PLAN_CELLS} cells per plan.")
        f = float(factor)
        _load_grounding(ctx)
        if not (_grounded(ctx, f) or (f and round(1 / f, 9) in ctx.allowed)):
            raise InvalidOperation(ctx.t(
                f"Hệ số {factor:g} không có trong tin nhắn của người dùng hay file nguồn; hãy hỏi tỷ giá/hệ số.",
                f"The factor {factor:g} was not given by the user or the sources; ask for the rate."),
                ungrounded=[f"factor={factor:g}"])
        if op == "divide" and f == 0:
            raise InvalidOperation("Cannot divide by 0.")
        f_ws = _book(ctx, False)[name]
        acts, skipped_formulas, formats = [], [], set()
        for a in _expand(range):
            raw = f_ws[a].value
            if isinstance(raw, str) and raw.startswith("="):
                skipped_formulas.append(a)
                continue
            if not isinstance(raw, (int, float)) or isinstance(raw, bool):
                continue
            x = float(raw)
            y = {"multiply": x * f, "divide": x / f, "add": x + f, "subtract": x - f}[op]
            y = round(y, round_digits) if round_digits is not None else y
            y = int(y) if float(y).is_integer() else y
            acts.append(SetValue(type="set_value", sheet=name, cell=a, value=y, value_type="number"))
            formats.add(f_ws[a].number_format)
        if not acts:
            return {"error": ctx.t(
                "Vùng này không có ô số nào để đổi" + (f" (các ô {', '.join(skipped_formulas[:10])} là công thức; "
                                                     "hãy đổi các ô đầu vào của công thức)" if skipped_formulas else "")
                + ".", "No numeric cells to transform in this range" + (
                    f" ({', '.join(skipped_formulas[:10])} are formulas; transform their input cells instead)"
                    if skipped_formulas else "") + "."), "skipped_formula_cells": skipped_formulas[:50]}
        out = _plan(acts, reason or f"{op} {range} by {factor:g}")
        out["computed_by"] = f"code: value {dict(multiply='*', divide='/', add='+', subtract='-')[op]} {factor:g}" + (
            f", rounded to {round_digits} decimals" if round_digits is not None else "")
        if skipped_formulas:
            out["skipped_formula_cells"] = skipped_formulas[:50]
            out["formula_note"] = ctx.t("Ô công thức không bị ghi đè; chúng tự tính lại từ ô đầu vào.",
                                        "Formula cells were not overwritten; they recalculate from their inputs.")
        cur = [x for x in formats if any(k in x for k in ("₫", "VND", "đ", "$", "€", "¥"))]
        if cur:
            out["number_format_warning"] = ctx.t(
                f"Định dạng số của các ô vẫn là {cur[0]!r} (trình sửa chưa đổi được định dạng); hãy đổi định dạng "
                "trong Excel hoặc ghi chú đơn vị ở tiêu đề cột.",
                f"The cells keep the number format {cur[0]!r} (the editor cannot change formats); change it in "
                "Excel or note the unit in the column header.")
        hidden = [x for x in formats if x in ("0", "#,##0") or (x.startswith("#,##0") and ".0" not in x and "₫" in x)]
        if hidden and any(isinstance(a.value, float) for a in acts):
            out["decimals_hidden_warning"] = ctx.t(
                f"Giá trị đã lưu đủ số lẻ nhưng định dạng {hidden[0]!r} hiển thị số nguyên (vd 15.6 hiện là 16).",
                f"Values are stored with decimals but the format {hidden[0]!r} displays whole numbers "
                "(e.g. 15.6 shows as 16).")
        return out

    def _plan(acts: list, reason: str) -> dict:
        diff = _diff(acts)
        plan = store.create_plan(workbook_id=ctx.workbook_id, thread_id=ctx.thread_id, turn_id=ctx.run.id,
                                 base_revision=ctx.revision, reason=reason[:300],
                                 actions=[a.model_dump() for a in acts], diff=diff)
        n = len({(a.sheet, a.cell) for a in acts})
        if n <= AUTO_APPLY_MAX_CELLS:
            return _commit(plan)
        errors = _dry_run(ctx, acts)
        if errors:
            plan["status"] = "invalid"
            store.save_plan(plan)
            return {"error": "The edit would fail validation; nothing was changed.", "errors": errors[:10]}
        ctx.pending_plan_id = plan["id"]
        ctx.run.emit("confirmation_required", {"plan_id": plan["id"], "base_revision": ctx.revision,
                                               "actions": plan["actions"][:200], "diff": diff, "cells": n,
                                               "reason": reason[:300]})
        return {"applied": False, "awaiting_user_confirmation": True, "plan_id": plan["id"], "cells": n,
                "message": ctx.t(f"Kế hoạch sửa {n} ô cần bạn xác nhận (bấm Áp dụng).",
                                 f"This plan edits {n} cells and needs your confirmation (press Apply)."),
                "summary": f"{n} cells awaiting confirmation"}

    @_Tool(ctx, "apply_workbook_edits")
    def apply_workbook_edits(plan_id: str) -> dict:
        plan = store.get_plan(plan_id)
        if plan["workbook_id"] != ctx.workbook_id or plan["thread_id"] != ctx.thread_id:
            raise NotFound("Plan not found in this conversation.")
        if plan["status"] != "pending":
            return {"error": f"Plan is {plan['status']}; nothing to apply."}
        if not CONFIRM_RE.search(ctx.message):
            return {"error": "Only apply a pending plan after the user explicitly confirms it in this message."}
        return _commit(plan)

    @_Tool(ctx, "validate_workbook")
    def validate_workbook(revision: int | None = None) -> dict:
        rev = revision or ctx.revision
        path = service.revision_path(ctx.workbook_id, rev)
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        errs = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=min(ws.max_row, 5000)):
                for c in row:
                    if isinstance(c.value, str) and c.value in ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A",
                                                                "#NUM!", "#NULL!"):
                        errs.append(f"{ws.title}!{c.coordinate} {c.value}")
        prev = service.get_preview(ctx.workbook_id, rev)
        entry = next((r for r in service.get_metadata(ctx.workbook_id)["revisions"] if r["revision"] == rev), {})
        return {"revision": rev, "opens": True, "cached_error_cells": errs[:30],
                "stale_formulas": prev.get("stale_formulas", 0), "warnings": prev.get("warnings", [])[:20],
                "origin": entry.get("origin"), "changed_count": entry.get("changed_count"),
                "note": "Cached values of formulas outside the edit are recomputed by Excel on open.",
                "summary": "ok" if not errs else f"{len(errs)} error value(s)"}

    @_Tool(ctx, "build_workbook")
    def build_workbook(brief: str) -> dict:
        from .. import builder, clarify, generate, spec_gen
        b = _workspace_sources(ctx) or sources.SourceBundle()
        sources.prompt_numbers(b, ctx.user_text)
        spec, warns, mode = spec_gen.build_spec(b, brief, {})
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "new.xlsx"
            stats = builder.build(spec, out, clarify.is_vi(brief))
            errors, vwarns = generate.validate_generated(out, stats)
            if errors:
                return {"error": "Generated workbook failed validation", "errors": errors[:10]}
            meta = service.create_from_path(out, generate._filename({}, spec.title), "generated",
                                            service.get_metadata(ctx.workbook_id).get("workspace_id"),
                                            extra={"warnings": warns + vwarns, "llm_mode": mode})
        ctx.new_workbook_id = meta["id"]
        ctx.run.emit("workbook_created", {"workbook_id": meta["id"], "revision": 1})
        return {"workbook_id": meta["id"], "sheets": meta["sheets"], "warnings": (warns + vwarns)[:10],
                "message": ctx.t(f"Đã tạo workbook mới '{meta['name']}'.", f"Created new workbook '{meta['name']}'."),
                "summary": f"{len(meta['sheets'])} sheet(s)"}

    @_Tool(ctx, "run_template_extractor")
    def run_template_extractor() -> dict:
        from .. import extractor_adapter
        if not extractor_adapter.available():
            return {"error": ctx.t("Tạo template trống cần model Azure OpenAI; đang ở mock mode nên bị tắt.",
                                   "Making a blank template needs the Azure OpenAI model; it is disabled in mock mode.")}
        meta = service.get_metadata(ctx.workbook_id)
        f = upload.store(service.revision_path(ctx.workbook_id, ctx.revision).read_bytes(),
                         f"{meta['name']}.xlsx"[:120], upload.XLSX_EXT)        # a copy; the revision is untouched
        run = extractor_adapter.start(file_id=f.id)
        return {"run_id": run.id, "message": ctx.t(
            "Đang tạo template trống từ workbook này; khi xong nó sẽ xuất hiện trong thư viện template.",
            "Making a blank template from this workbook; it will appear in the template gallery when done."),
            "summary": f"extract run {run.id}"}

    specs = [
        (inspect_workbook, "inspect_workbook", NoArgs,
         "Summarise the workbook: sheets, sizes, formulas, detected input slots, unsupported features. Call first."),
        (read_range, "read_range", RangeArgs, "Read values and formulas of an A1 range (max 400 cells)."),
        (read_source, "read_source", SourceArgs,
         "List the workspace's txt/csv sources (no file_id) or read numbered lines + extracted facts of one."),
        (propose_workbook_edits, "propose_workbook_edits", ProposeArgs,
         f"Set or clear cell values. Up to {AUTO_APPLY_MAX_CELLS} cells are applied immediately as a new revision; "
         "larger edits wait for the user's confirmation. Numbers must come from the user, the sources or the "
         "workbook (never a value you computed). Formulas, styles and inserting rows are not supported."),
        (transform_range, "transform_range", TransformArgs,
         "Convert or scale the numbers of a range by a factor the user gave (currency conversion like VND->USD: "
         "op='divide', factor=25000; percentage uplift; unit change). The code computes the new values and writes "
         f"them like propose_workbook_edits (<= {AUTO_APPLY_MAX_CELLS} cells immediately, else after confirmation). "
         "Formula cells are left untouched. Use this instead of computing numbers yourself."),
        (apply_workbook_edits, "apply_workbook_edits", PlanArgs,
         "Apply a pending plan from this conversation, only after the user explicitly confirmed it."),
        (validate_workbook, "validate_workbook", ValidateArgs,
         "Check a revision: opens, error values in formulas, preview warnings."),
        (build_workbook, "build_workbook", BuildArgs,
         "Create a NEW workbook from the workspace sources and a brief (does not change the current one)."),
        (run_template_extractor, "run_template_extractor", NoArgs,
         "Only when the user asks to turn this populated workbook into a blank template (needs the Azure OpenAI model)."),
    ]
    return [StructuredTool.from_function(func=fn, name=name, description=desc, args_schema=schema)
            for fn, name, schema, desc in specs]


def _date_in_user_text(ctx: Ctx, value: str) -> bool:
    from .. import builder
    _load_grounding(ctx)
    d = builder.to_date(value)
    if not d:
        return False
    hay = " ".join([ctx.user_text, *ctx.texts])
    return any(s in hay for s in (d.strftime("%Y-%m-%d"), d.strftime("%d/%m/%Y"), f"{d.day}/{d.month}/{d.year}"))


def _dry_run(ctx: Ctx, acts: list) -> list[str]:
    """Apply to a temp copy and validate, without creating a revision."""
    src = service.revision_path(ctx.workbook_id, ctx.revision)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        stats = ooxml_edit.apply_actions(str(src), tmp, acts)
        report = validate_edit.validate(str(src), tmp, stats)
        return [] if report["pass"] else list(report["errors"])
    except WorkspaceError as e:
        return [e.message]
    finally:
        Path(tmp).unlink(missing_ok=True)


def apply_plan(plan_id: str, base_revision: int):
    """User approval from the UI (POST /api/agent/plans/{id}/apply)."""
    plan = store.get_plan(plan_id)
    if plan["status"] != "pending":
        raise InvalidOperation(f"Plan is {plan['status']}.")
    if base_revision != plan["base_revision"]:
        from ..errors import RevisionConflict
        cur = service.get_metadata(plan["workbook_id"])["current_revision"]
        raise RevisionConflict(f"The plan was made for revision {plan['base_revision']}; ask the agent again.",
                               current_revision=cur)
    res = service.apply_operations(plan["workbook_id"], OperationBatch(
        base_revision=plan["base_revision"], operation_id=plan["id"], origin="agent",
        note=("agent (approved): " + (plan.get("reason") or "edit"))[:300], actions=plan["actions"]))
    plan.update(status="applied", revision=res.revision)
    store.save_plan(plan)
    return res


def reject_plan(plan_id: str) -> dict:
    plan = store.get_plan(plan_id)
    if plan["status"] == "pending":
        plan["status"] = "rejected"
        store.save_plan(plan)
    return {"ok": True, "status": plan["status"]}

