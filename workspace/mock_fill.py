"""Explicitly requested mock data for a template (no sources attached).

Only runs when the prompt itself asks for mock / sample / fake data. The AI reads the template grid and proposes
values for value cells (never labels or formulas). The proposal is shown to the user as editable 'CELL: value'
lines; generate re-validates every cell and writes it with via="mock", and the workbook is marked as mock data.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from . import config

MOCK_RE = re.compile(
    r"\bmock|\bdummy|\bfake\b|sample\s+data|demo\s+data|test\s+data|placeholder\s+data|"
    r"giả\s*lập|dữ\s*liệu\s*(mẫu|giả|ảo|demo|test)|data\s*(mẫu|giả)|bịa|"
    r"샘플|가상|예시\s*데이터|더미|테스트\s*데이터|ダミー|サンプル|仮データ", re.I)
LINE_RE = re.compile(r"^\s*(?:(?P<sheet>[^!:]+)!)?(?P<cell>[A-Z]{1,3}[1-9]\d{0,6})\s*:\s?(?P<value>.*)$")
MAX_CELLS = 300
MAX_TEXT = 500
MAX_ANSWER = 30000

SYSTEM = (
    "You fill an Excel form template with MOCK data that the user explicitly asked for. You get the template grid "
    "(cell: current text) plus the empty merged blocks. Decide which cells are VALUES (the content of the form: "
    "titles of the request, names, dates, numbers, descriptions, O/X answers, table rows) and which are LABELS "
    "(field names, section titles, column headers, instructions). Propose a new value for value cells only: replace "
    "existing example values and fill empty value cells next to / below their labels, so the result reads like one "
    "coherent, realistic record matching the user's request. Never touch labels or column headers. Write in the "
    "template's language unless the user asks otherwise. Use obviously fictional names, companies, phone numbers and "
    "ids (never real organisations or people). Keep dates and numbers consistent with each other. "
    'Answer JSON only: {"cells": {"<sheet>!<A1>": "<value>", ...}}. Numbers may be JSON numbers.')


def wants_mock(prompt: str) -> bool:
    return bool(MOCK_RE.search(prompt or ""))


def _grid(path: Path) -> tuple[dict, dict]:
    """{sheet: {cell: text}} of non-empty cells, and {sheet: {"formulas": set, "covered": set, "anchors": {cell: ref}}}."""
    wb = load_workbook(path, data_only=False)
    cells: dict[str, dict[str, str]] = {}
    info: dict[str, dict[str, Any]] = {}
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        formulas, covered, anchors = set(), set(), {}
        for rng in ws.merged_cells.ranges:
            c1, r1, c2, r2 = range_boundaries(str(rng))
            anchor = ws.cell(r1, c1).coordinate
            anchors[anchor] = str(rng)
            covered.update(ws.cell(r, c).coordinate for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)
                           if (r, c) != (r1, c1))
        cells[ws.title] = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None or c.coordinate in covered:
                    continue
                if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
                    formulas.add(c.coordinate)
                    continue
                cells[ws.title][c.coordinate] = str(c.value)
        info[ws.title] = {"formulas": formulas, "covered": covered, "anchors": anchors,
                          "max_row": ws.max_row, "max_col": ws.max_column}
    return cells, info


def _check(info: dict, sheet: str, cell: str) -> str | None:
    """Why a cell may not be written, or None when it is a valid target."""
    if sheet not in info:
        return f"unknown sheet '{sheet}'"
    s = info[sheet]
    if cell in s["formulas"]:
        return f"{cell} holds a formula"
    if cell in s["covered"]:
        return f"{cell} is inside a merged range (not its top-left cell)"
    c1, r1, _, _ = range_boundaries(cell + ":" + cell)
    if r1 > s["max_row"] + 5 or c1 > s["max_col"] + 2:
        return f"{cell} is outside the template's used area"
    return None


def _one_line(v: Any) -> str:
    return str(v).replace("\r\n", "\n").replace("\n", "\\n")[:MAX_TEXT]


def propose(path: Path, prompt: str) -> tuple[str, int, list[str]]:
    """Ask the model for mock values -> (editable 'CELL: value' text, cell count, warnings)."""
    cells, info = _grid(path)
    single = len(cells) == 1
    lines = [f"User request: {prompt}", "Template grid (cell: current text; formulas omitted):"]
    for sheet, cs in cells.items():
        lines.append(f"## sheet {sheet}")
        lines += [f"{sheet}!{a}: {_one_line(v)[:120]}" for a, v in cs.items()]
        empty = [ref for a, ref in info[sheet]["anchors"].items() if a not in cs and a not in info[sheet]["formulas"]]
        if empty:
            lines.append("Empty merged blocks (candidate value cells, anchor:range): " +
                         ", ".join(f"{sheet}!{r.split(':')[0]}:{r}" for r in empty[:300]))
    llm = config.azure_chat(4000, 120, model_kwargs={"response_format": {"type": "json_object"}})
    raw = llm.invoke([("system", SYSTEM), ("user", "\n".join(lines))]).content
    data = json.loads(raw if isinstance(raw, str) else json.dumps(raw))
    out, warnings = [], []
    for key, v in (data.get("cells") or {}).items():
        sheet, _, cell = str(key).rpartition("!")
        sheet = sheet.strip("'") or next(iter(cells))
        cell = cell.strip().upper().replace("$", "")
        if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", cell) or v is None or isinstance(v, (dict, list)):
            continue
        if why := _check(info, sheet, cell):
            warnings.append(f"Mock value for {sheet}!{cell} skipped: {why}.")
            continue
        out.append(f"{cell if single else f'{sheet}!{cell}'}: {_one_line(v)}")
        if len(out) >= MAX_CELLS:
            break
    return "\n".join(out), len(out), warnings[:10]


def parse(path: Path, text: str) -> tuple[list[tuple[str, str, Any]], list[str]]:
    """User-reviewed 'CELL: value' lines -> [(sheet, cell, value)] for valid targets, plus warnings."""
    _, info = _grid(path)
    first = next(iter(info), "")
    out: dict[tuple[str, str], Any] = {}
    warnings: list[str] = []
    for line in (text or "").splitlines()[:MAX_CELLS * 2]:
        m = LINE_RE.match(line)
        if not m or not m["value"].strip():
            continue
        sheet, cell = (m["sheet"] or first).strip().strip("'"), m["cell"]
        if why := _check(info, sheet, cell):
            warnings.append(f"Mock value for {sheet}!{cell} skipped: {why}.")
            continue
        v: Any = m["value"].strip().replace("\\n", "\n")[:MAX_TEXT]
        if re.fullmatch(r"-?\d+(\.\d+)?", v) and not (len(v) > 1 and v.startswith("0") and "." not in v):
            v = float(v) if "." in v else int(v)
        out[(sheet, cell)] = v
        if len(out) >= MAX_CELLS:
            break
    return [(s, c, v) for (s, c), v in out.items()], warnings[:10]
