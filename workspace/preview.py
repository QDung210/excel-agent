"""Canonical .xlsx -> FortuneSheet sheet DTO (read-only adapter; never writes the workbook).

The preview is an approximation: see docs/compatibility-matrix.md. Features the grid cannot
show are reported as warnings; the exported file is always the canonical backend revision.
"""
from __future__ import annotations

import colorsys
import datetime as dt
import io
import re
import warnings
import zipfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from . import config, formula_eval

INDEXED = ["000000", "FFFFFF", "FF0000", "00FF00", "0000FF", "FFFF00", "FF00FF", "00FFFF",
           "000000", "FFFFFF", "FF0000", "00FF00", "0000FF", "FFFF00", "FF00FF", "00FFFF",
           "800000", "008000", "000080", "808000", "800080", "008080", "C0C0C0", "808080",
           "9999FF", "993366", "FFFFCC", "CCFFFF", "660066", "FF8080", "0066CC", "CCCCFF",
           "000080", "FF00FF", "FFFF00", "00FFFF", "800080", "800000", "008080", "0000FF",
           "00CCFF", "CCFFFF", "CCFFCC", "FFFF99", "99CCFF", "FF99CC", "CC99FF", "FFCC99",
           "3366FF", "33CCCC", "99CC00", "FFCC00", "FF9900", "FF6600", "666699", "969696",
           "003366", "339966", "003300", "333300", "993300", "993366", "333399", "333333"]
BORDER_STYLE = {"thin": 1, "hair": 2, "dotted": 3, "dashed": 4, "dashDot": 5, "dashDotDot": 6, "double": 7,
                "medium": 8, "mediumDashed": 9, "mediumDashDot": 10, "mediumDashDotDot": 11,
                "slantDashDot": 12, "thick": 13}
H_ALIGN = {"center": 0, "centerContinuous": 0, "left": 1, "right": 2}
V_ALIGN = {"center": 0, "top": 1, "bottom": 2}


# ---------------------------------------------------------------- colors
def _theme_colors(z: zipfile.ZipFile) -> list[str]:
    try:
        xml = z.read("xl/theme/theme1.xml").decode("utf-8")
    except KeyError:
        return []
    scheme = re.search(r"<a:clrScheme\b.*?</a:clrScheme>", xml, re.S)
    if not scheme:
        return []
    vals = {}
    for tag in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
                "hlink", "folHlink"):
        m = re.search(rf"<a:{tag}>.*?(?:srgbClr val|lastClr)=\"([0-9A-Fa-f]{{6}})\".*?</a:{tag}>", scheme.group(0), re.S)
        vals[tag] = m.group(1).upper() if m else "000000"
    # Excel theme index order swaps dark/light pairs
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
             "hlink", "folHlink"]
    return [vals[t] for t in order]


def _tint(hex6: str, tint: float) -> str:
    if not tint:
        return hex6
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0), 1), s)
    return "".join(f"{round(x * 255):02X}" for x in (r, g, b))


def _color(c: Any, theme: list[str]) -> str | None:
    if c is None:
        return None
    try:
        if c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) >= 6:
            hex6 = c.rgb[-6:]
        elif c.type == "theme" and c.theme is not None and c.theme < len(theme):
            hex6 = theme[c.theme]
        elif c.type == "indexed" and c.indexed is not None and c.indexed < len(INDEXED):
            hex6 = INDEXED[c.indexed]
        else:
            return None
        return "#" + _tint(hex6.upper(), float(c.tint or 0))
    except (AttributeError, ValueError, TypeError):
        return None


# ---------------------------------------------------------------- number display
def _fmt_number(v: float, fmt: str) -> str:
    f = (fmt or "General").split(";")[0]
    if f == "General" or not f:
        if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
            return str(int(v))
        return f"{v:.10g}"
    clean = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.|_.|\*.', "", f)
    pct = "%" in clean
    if pct:
        v = v * 100
    dec = 0
    m = re.search(r"\.([0#]+)", clean)
    if m:
        dec = len(m.group(1))
    grouped = "," in clean.split(".")[0]
    s = f"{v:,.{dec}f}" if grouped else f"{v:.{dec}f}"
    prefix = "".join(re.findall(r'"([^"]*)"', f.split("0")[0].split("#")[0]))
    sym = re.search(r"\[\$([^-\]]+)", f)
    lead = (sym.group(1) if sym and f.index("[$") < max(f.find("0"), 0) else "") + prefix
    trail = "".join(re.findall(r'"([^"]*)"', f[max(f.rfind("0"), f.rfind("#")):]))
    if sym and f.index("[$") > max(f.rfind("0"), 0):
        trail += sym.group(1)
    return f"{lead}{s}{'%' if pct else ''}{trail}"


def _fmt_date(v: dt.datetime | dt.date | dt.time, fmt: str) -> str:
    f = (fmt or "").lower()
    if isinstance(v, dt.time):
        return v.strftime("%H:%M:%S" if "s" in f else "%H:%M")
    if isinstance(v, dt.datetime) and ("h" in f) and v.time() != dt.time(0):
        return v.strftime("%Y-%m-%d %H:%M")
    d = v.date() if isinstance(v, dt.datetime) else v
    if "mmm" in f:
        return d.strftime("%d %b %Y") if f.startswith("d") else d.strftime("%b %d, %Y")
    if f.startswith("d") or f.startswith("[$-") and "/d" not in f[:6]:
        return d.strftime("%d/%m/%Y")
    if f.startswith("m"):
        return d.strftime("%m/%d/%Y")
    return d.isoformat()


def _serial(v: dt.datetime | dt.date | dt.time) -> float:
    if isinstance(v, dt.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if not isinstance(v, dt.datetime):
        v = dt.datetime(v.year, v.month, v.day)
    d = v - dt.datetime(1899, 12, 30)
    return d.days + d.seconds / 86400


# ---------------------------------------------------------------- sheet conversion
def _border_side(side: Any, theme: list[str]) -> dict | None:
    if side is None or not side.style:
        return None
    return {"style": BORDER_STYLE.get(side.style, 1), "color": _color(side.color, theme) or "#000000"}


def _sheet_warnings(z: zipfile.ZipFile, part: str, name: str) -> list[str]:
    out = []
    try:
        xml = z.read(part).decode("utf-8")
    except KeyError:
        return out
    if "x14:dataValidation" in xml:
        out.append(f"{name}: dropdowns stored as Excel extensions (x14) are preserved in the file but not shown in the grid.")
    elif "<dataValidation" in xml:
        out.append(f"{name}: data validation rules are preserved in the file but not enforced in the grid.")
    if "<conditionalFormatting" in xml or "x14:conditionalFormatting" in xml:
        out.append(f"{name}: conditional formatting is preserved in the file but not rendered in the grid.")
    if "<drawing " in xml or "<drawing>" in xml:
        out.append(f"{name}: charts/images are preserved in the file but not shown in the grid.")
    if "<legacyDrawing" in xml:
        out.append(f"{name}: comments/notes are preserved in the file but not shown in the grid.")
    if "<tableParts" in xml:
        out.append(f"{name}: Excel tables are preserved; table styling/filters are not shown in the grid.")
    if "<sheetProtection" in xml:
        out.append(f"{name}: sheet protection is preserved in the file but not enforced in the grid.")
    return out


def build_preview(path: str) -> dict:
    with open(path, "rb") as f:
        data = f.read()
    from xlsx_template import ooxml
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        theme = _theme_colors(z)
        parts = ooxml.sheet_parts(z)
        names = z.namelist()
        sheet_warn = {n: _sheet_warnings(z, p, n) for n, p in parts.items()}
        global_warn = []
        if any(n.startswith("xl/pivotTables/") for n in names):
            global_warn.append("Pivot tables are preserved in the file but not shown in the grid.")
        if any(n.startswith("xl/externalLinks/") for n in names):
            global_warn.append("External links are preserved; their values are not refreshed.")
        if any(n.startswith("xl/chartsheets/") for n in names):
            global_warn.append("Chart sheets are preserved in the file but not shown.")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(io.BytesIO(data))
        wb_vals = load_workbook(io.BytesIO(data), data_only=True)

    sheets_out: list[dict] = []
    budget = config.MAX_PREVIEW_CELLS
    stale = 0
    grid: dict[str, dict[tuple[int, int], Any]] = {}        # values for the formula evaluator
    pending: dict[str, dict[tuple[int, int], str]] = {}     # formulas without a cached result
    pending_cells: list[tuple[str, int, int, dict, str]] = []
    for order, ws in enumerate(wb.worksheets):
        sheet_grid = grid.setdefault(ws.title, {})
        wv = wb_vals[ws.title]
        celldata = []
        border_info = []
        max_r, max_c = ws.max_row or 1, ws.max_column or 1
        truncated = False
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                has_style = cell.has_style and (cell.fill.fill_type or cell.border.left.style or cell.border.right.style
                                                or cell.border.top.style or cell.border.bottom.style or cell.font.b)
                if cell.value is None and not has_style:
                    continue
                if budget <= 0:
                    truncated = True
                    break
                budget -= 1
                r, c = cell.row - 1, cell.column - 1
                v: dict[str, Any] = {}
                fmt = cell.number_format or "General"
                val = cell.value
                if isinstance(val, str) and val.startswith("=") and cell.data_type == "f":
                    v["f"] = val
                    cached = wv.cell(row=cell.row, column=cell.column).value
                    if cached is None:
                        pending.setdefault(ws.title, {})[(cell.row, cell.column)] = val
                        pending_cells.append((ws.title, cell.row, cell.column, v, fmt))
                    val = cached
                elif cell.data_type == "f":           # ArrayFormula / DataTableFormula objects
                    v["f"] = "=" + str(getattr(val, "text", val)).lstrip("=")
                    val = wv.cell(row=cell.row, column=cell.column).value
                if val is not None:
                    sheet_grid[(cell.row, cell.column)] = (_serial(val) if isinstance(val, (dt.datetime, dt.date))
                                                           else val)
                    if isinstance(val, bool):
                        v.update(v=val, m="TRUE" if val else "FALSE", ct={"fa": "General", "t": "b"})
                    elif isinstance(val, (int, float)):
                        v.update(v=val, m=_fmt_number(float(val), fmt), ct={"fa": fmt, "t": "n"})
                    elif isinstance(val, (dt.datetime, dt.date, dt.time)):
                        v.update(v=_serial(val), m=_fmt_date(val, fmt), ct={"fa": fmt, "t": "d"})
                    else:
                        s = str(val)
                        v.update(v=s, m=s, ct={"fa": "@" if fmt == "@" else "General", "t": "s" if fmt == "@" else "g"})
                else:
                    v["ct"] = {"fa": fmt, "t": "g"}
                font = cell.font
                if font is not None:
                    if font.b:
                        v["bl"] = 1
                    if font.i:
                        v["it"] = 1
                    if font.u:
                        v["un"] = 1
                    if font.strike:
                        v["cl"] = 1
                    if font.sz and float(font.sz) != 11:
                        v["fs"] = float(font.sz)
                    if font.name:
                        v["ff"] = font.name
                    fc = _color(font.color, theme)
                    if fc and fc != "#000000":
                        v["fc"] = fc
                if cell.fill is not None and cell.fill.fill_type == "solid":
                    bg = _color(cell.fill.fgColor, theme)
                    if bg:
                        v["bg"] = bg
                al = cell.alignment
                if al is not None:
                    if al.horizontal in H_ALIGN:
                        v["ht"] = H_ALIGN[al.horizontal]
                    v["vt"] = V_ALIGN.get(al.vertical or "bottom", 2)
                    if al.wrap_text:
                        v["tb"] = 2
                b = cell.border
                if b is not None:
                    sides = {k: _border_side(getattr(b, n), theme) for k, n in
                             (("l", "left"), ("r", "right"), ("t", "top"), ("b", "bottom"))}
                    sides = {k: s for k, s in sides.items() if s}
                    if sides:
                        border_info.append({"rangeType": "cell", "value": {"row_index": r, "col_index": c, **sides}})
                celldata.append({"r": r, "c": c, "v": v})
            if truncated:
                break

        merge = {}
        for mr in ws.merged_cells.ranges:
            r0, c0 = mr.min_row - 1, mr.min_col - 1
            merge[f"{r0}_{c0}"] = {"r": r0, "c": c0, "rs": mr.max_row - mr.min_row + 1, "cs": mr.max_col - mr.min_col + 1}
        # FortuneSheet expects the merge marker on the anchor cell too
        anchors = {(m["r"], m["c"]): m for m in merge.values()}
        for cd in celldata:
            m = anchors.get((cd["r"], cd["c"]))
            if m:
                cd["v"]["mc"] = dict(m)
        seen = {(cd["r"], cd["c"]) for cd in celldata}
        for m in merge.values():
            if (m["r"], m["c"]) not in seen:
                celldata.append({"r": m["r"], "c": m["c"], "v": {"mc": dict(m)}})
            for rr in range(m["r"], m["r"] + m["rs"]):
                for cc in range(m["c"], m["c"] + m["cs"]):
                    if (rr, cc) != (m["r"], m["c"]):
                        celldata.append({"r": rr, "c": cc, "v": {"mc": {"r": m["r"], "c": m["c"]}}})
        columnlen, colhidden = {}, {}
        for key, dim in ws.column_dimensions.items():
            try:
                from openpyxl.utils import column_index_from_string
                lo = column_index_from_string(key)
            except ValueError:
                continue
            hi = max(dim.max or lo, lo)
            for c in range(lo, min(hi, lo + 200) + 1):
                if dim.width:
                    columnlen[str(c - 1)] = round(dim.width * 7 + 5)
                if dim.hidden:
                    colhidden[str(c - 1)] = 0
        rowlen, rowhidden = {}, {}
        for idx, dim in ws.row_dimensions.items():
            if dim.height:
                rowlen[str(idx - 1)] = round(dim.height * 4 / 3)
            if dim.hidden:
                rowhidden[str(idx - 1)] = 0
        sheet: dict[str, Any] = {
            "name": ws.title,
            "id": f"s{order}",
            "order": order,
            "status": 1 if order == wb.index(wb.active) else 0,
            "row": max(max_r + 20, 60),
            "column": max(max_c + 6, 26),
            "celldata": celldata,
            "config": {"merge": merge, "borderInfo": border_info, "columnlen": columnlen, "rowlen": rowlen,
                       "rowhidden": rowhidden, "colhidden": colhidden},
        }
        if ws.sheet_state != "visible":
            sheet["hide"] = 1
        fp = ws.freeze_panes
        if fp:
            fr, fc_ = (ws[fp].row - 1, ws[fp].column - 1)
            if fr and fc_:
                sheet["frozen"] = {"type": "rangeBoth", "range": {"row_focus": fr - 1, "column_focus": fc_ - 1}}
            elif fr:
                sheet["frozen"] = {"type": "rangeRow", "range": {"row_focus": fr - 1, "column_focus": 0}}
            elif fc_:
                sheet["frozen"] = {"type": "rangeColumn", "range": {"row_focus": 0, "column_focus": fc_ - 1}}
        w = list(sheet_warn.get(ws.title, []))
        if truncated:
            w.append(f"{ws.title}: preview truncated at {config.MAX_PREVIEW_CELLS} cells; the file itself is complete.")
        sheet["warnings"] = w
        sheet["dimensions"] = f"A1:{get_column_letter(max_c)}{max_r}"
        sheets_out.append(sheet)
    wb.close()
    wb_vals.close()

    # Formulas whose cache was stripped (after an edit) or never written: compute the common ones for display.
    if pending:
        ev = formula_eval.Workbook(grid, pending)
        for title, row, col, v, fmt in pending_cells:
            res = ev.value(title, row, col)
            if res is formula_eval.UNSUPPORTED or res is None:
                stale += 1
            elif isinstance(res, bool):
                v.update(v=res, m="TRUE" if res else "FALSE", ct={"fa": "General", "t": "b"})
            elif isinstance(res, (int, float)):
                v.update(v=res, m=_fmt_number(float(res), fmt), ct={"fa": fmt, "t": "n"})
            else:
                v.update(v=str(res), m=str(res), ct={"fa": "General", "t": "g"})

    warn = list(global_warn)
    for s in sheets_out:
        warn.extend(s["warnings"])
    computed = len(pending_cells) - stale
    if computed:
        warn.append(f"{computed} formula value(s) were computed by the preview; Excel recalculates them on open.")
    if stale:
        warn.append(f"{stale} formula cell(s) have no cached result and use functions the preview cannot compute; "
                    f"values appear after Excel recalculates.")
    if not any(s["status"] == 1 and not s.get("hide") for s in sheets_out):
        for s in sheets_out:
            if not s.get("hide"):
                s["status"] = 1
                break
    return {"sheets": sheets_out, "warnings": warn, "stale_formulas": stale, "computed_formulas": computed}
