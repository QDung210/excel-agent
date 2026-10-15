"""Template gallery (user uploads + 2 built-ins made with XlsxWriter) and input-slot detection.

Slot detection only reads the workbook (openpyxl, never saved). A slot is an empty, non-formula cell the fill step
may write with `set_value`:
- key/value slot: a text label whose right-hand neighbour is empty ("Khách hàng:" | [   ])
- table slot: a header row of >= 2 text cells followed by input rows (empty or formula cells, no constants),
  ending at the first totals/constant row.
"""
from __future__ import annotations

import re
import shutil
import threading
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import storage, upload
from .errors import InvalidOperation, NotFound

AGG_RE = re.compile(r"^=\s*(SUM|SUBTOTAL|AVERAGE|COUNT|COUNTA|MIN|MAX)\(", re.I)
TOTAL_RE = re.compile(r"^\s*(tổng|total|cộng|sum|grand total)", re.I)
MAX_TABLE_ROWS = 1000
_lock = threading.Lock()


def norm(text: str) -> str:
    s = unicodedata.normalize("NFD", str(text or "").replace("đ", "d").replace("Đ", "D"))
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


# ---------------------------------------------------------------- slots
@dataclass
class KvSlot:
    sheet: str
    cell: str
    label: str


@dataclass
class TableColumn:
    col: int
    header: str
    formula: bool


@dataclass
class TableSlot:
    sheet: str
    header_row: int
    columns: list[TableColumn]
    rows: list[int]
    prefilled_rows: int = 0               # rows below the header that already hold constants

    @property
    def inputs(self) -> list[TableColumn]:
        return [c for c in self.columns if not c.formula]


@dataclass
class Slots:
    kv: list[KvSlot] = field(default_factory=list)
    tables: list[TableSlot] = field(default_factory=list)
    hidden_sheets: list[str] = field(default_factory=list)

    def summary(self) -> dict:
        return {"kv": [{"sheet": k.sheet, "cell": k.cell, "label": k.label} for k in self.kv],
                "tables": [{"sheet": t.sheet, "header_row": t.header_row, "columns": [c.header for c in t.columns],
                            "input_rows": len(t.rows), "prefilled_rows": t.prefilled_rows} for t in self.tables],
                "hidden_sheets": list(self.hidden_sheets)}


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _text(v) -> bool:
    return isinstance(v, str) and not v.startswith("=") and v.strip() != ""


def detect_slots(path: str | Path) -> Slots:
    wb = load_workbook(path, data_only=False)          # read only; this object is never saved
    out = Slots()
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                out.hidden_sheets.append(ws.title)
                continue
            merged_tail = {(r, c) for m in ws.merged_cells.ranges for r in range(m.min_row, m.max_row + 1)
                           for c in range(m.min_col, m.max_col + 1) if (r, c) != (m.min_row, m.min_col)}
            merged_end = {(m.min_row, m.min_col): m.max_col for m in ws.merged_cells.ranges}
            cells = {k: v for k, v in ws._cells.items() if k not in merged_tail}
            val = {k: c.value for k, c in cells.items()}
            styled = {k for k, c in cells.items() if c.has_style}
            used_rows: set[int] = set()
            rows = sorted({r for r, _ in val})
            # --- tables
            for r in rows:
                if r in used_rows:
                    continue
                heads = sorted(c for (rr, c), v in val.items() if rr == r and _text(v))
                runs, cur = [], []
                for c in heads:
                    if cur and c - merged_end.get((r, cur[-1]), cur[-1]) > 1:      # gap ends a header run
                        runs.append(cur)
                        cur = []
                    cur.append(c)
                if cur:
                    runs.append(cur)
                for run in runs:
                    if len(run) < 2:
                        continue
                    t = _table_below(ws.title, r, run, val, styled)
                    if t:
                        out.tables.append(t)
                        used_rows |= {r, *t.rows, *range(r + 1, r + 1 + t.prefilled_rows)}
            # --- key/value
            in_table = {(t.header_row, c.col) for t in out.tables if t.sheet == ws.title for c in t.columns}
            for (r, c), v in sorted(val.items()):
                if r in used_rows or not _text(v) or (r, c) in in_table or len(v.strip()) > 60 or TOTAL_RE.match(v):
                    continue
                tc = merged_end.get((r, c), c) + 1
                tv = val.get((r, tc))
                if tv not in (None, "") or (r, tc) in merged_tail:
                    continue
                if v.strip().endswith(":") or (r, tc) in styled:
                    out.kv.append(KvSlot(ws.title, f"{get_column_letter(tc)}{r}", v.strip().rstrip(":").strip()))
    finally:
        wb.close()
    return out


def _table_below(sheet: str, r: int, run: list[int], val: dict, styled: set) -> TableSlot | None:
    heads = {c: str(val[(r, c)]).strip() for c in run}
    formula_cols: set[int] = set()
    rows: list[int] = []
    prefilled = 0
    rr = r + 1
    while rr <= r + MAX_TABLE_ROWS:
        vs = {c: val.get((rr, c)) for c in run}
        present = any((rr, c) in val for c in run)
        if not present:
            break
        consts = [v for v in vs.values() if v not in (None, "") and not _is_formula(v)]
        if any(_is_formula(v) and AGG_RE.match(v) for v in vs.values()) or \
                any(isinstance(v, str) and TOTAL_RE.match(v) for v in consts):
            break
        if consts:
            if rows:
                break
            prefilled += 1                 # populated template: data rows hold constants
            rr += 1
            continue
        if not any((rr, c) in styled or _is_formula(vs[c]) for c in run):
            break
        formula_cols |= {c for c, v in vs.items() if _is_formula(v)}
        rows.append(rr)
        rr += 1
    if not rows and not prefilled:
        return None
    cols = [TableColumn(c, heads[c], c in formula_cols) for c in run]
    if not rows or not any(not c.formula for c in cols):
        return TableSlot(sheet, r, cols, [], prefilled) if prefilled else None
    return TableSlot(sheet, r, cols, rows, prefilled)


# ---------------------------------------------------------------- gallery
def _dir(tpl_id: str) -> Path:
    try:
        storage.check_id(tpl_id, "tpl")
    except KeyError:
        raise NotFound("Template not found.") from None
    return storage.templates_dir() / tpl_id


def path(tpl_id: str) -> Path:
    p = _dir(tpl_id) / "template.xlsx"
    if not p.exists():
        raise NotFound("Template not found.")
    return p


def get(tpl_id: str) -> dict:
    meta = storage.read_json(_dir(tpl_id) / "meta.json")
    if not meta:
        raise NotFound("Template not found.")
    return meta


def _sheets(p: Path) -> list[str]:
    wb = load_workbook(p, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def add(data: bytes, name: str, description: str = "", origin: str = "user", builtin_key: str | None = None) -> dict:
    upload.check_xlsx(data)
    tid = storage.new_id("tpl")
    d = storage.templates_dir() / tid
    storage.atomic_write_bytes(d / "template.xlsx", data)
    try:
        sheets = _sheets(d / "template.xlsx")
    except Exception:
        shutil.rmtree(d, ignore_errors=True)
        raise InvalidOperation("The template could not be opened as a workbook.") from None
    meta = {"id": tid, "name": (name or "Template")[:120], "description": description[:300], "origin": origin,
            "sheets": sheets, "created": storage.now(), "builtin_key": builtin_key}
    storage.atomic_write_json(d / "meta.json", meta)
    return meta


def list_all() -> list[dict]:
    ensure_builtins()
    root = storage.templates_dir()
    out = [m for d in root.iterdir() if d.is_dir() and (m := storage.read_json(d / "meta.json"))] if root.exists() else []
    out.sort(key=lambda m: (m["origin"] != "builtin", -m["created"]))
    return [{k: m[k] for k in ("id", "name", "description", "origin", "sheets", "created")} for m in out]


def ensure_builtins() -> None:
    with _lock:
        root = storage.templates_dir()
        have = {m.get("builtin_key") for d in (root.iterdir() if root.exists() else []) if d.is_dir()
                for m in [storage.read_json(d / "meta.json") or {}]}
        for key, (name, desc, fn) in BUILTINS.items():
            if key in have:
                continue
            tmp = root / f".{key}.xlsx"
            tmp.parent.mkdir(parents=True, exist_ok=True)
            fn(tmp)
            try:
                add(tmp.read_bytes(), name, desc, "builtin", key)
            finally:
                tmp.unlink(missing_ok=True)


def _fmts(wb: xlsxwriter.Workbook) -> dict:
    return {"title": wb.add_format({"bold": True, "font_size": 15, "font_color": "#1F4E78"}),
            "label": wb.add_format({"bold": True}),
            "input": wb.add_format({"bg_color": "#FFF9E6", "border": 1, "border_color": "#D9D9D9"}),
            "date_in": wb.add_format({"bg_color": "#FFF9E6", "border": 1, "border_color": "#D9D9D9",
                                      "num_format": "dd/mm/yyyy"}),
            "head": wb.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#1F4E78", "border": 1,
                                   "text_wrap": True, "valign": "vcenter"}),
            "cell": wb.add_format({"border": 1, "border_color": "#D9D9D9"}),
            "num": wb.add_format({"border": 1, "border_color": "#D9D9D9", "num_format": "#,##0"}),
            "num_in": wb.add_format({"border": 1, "border_color": "#D9D9D9", "num_format": "#,##0",
                                     "bg_color": "#FFF9E6"}),
            "total": wb.add_format({"bold": True, "top": 2, "bg_color": "#DDEBF7", "border": 1, "num_format": "#,##0"})}


def _invoice(p: Path, rows: int = 12) -> None:
    wb = xlsxwriter.Workbook(str(p))
    f, ws = _fmts(wb), None
    ws = wb.add_worksheet("Hóa đơn")
    ws.write(0, 0, "HÓA ĐƠN BÁN HÀNG", f["title"])
    for i, (lab, fmt) in enumerate([("Khách hàng:", "input"), ("Địa chỉ:", "input"), ("Số hóa đơn:", "input"),
                                    ("Ngày:", "date_in")]):
        ws.write(2 + i, 0, lab, f["label"])
        ws.write_blank(2 + i, 1, None, f[fmt])
    hr = 7
    for c, h in enumerate(["STT", "Sản phẩm", "Số lượng", "Đơn giá", "Thành tiền"]):
        ws.write(hr, c, h, f["head"])
    for i in range(rows):
        r = hr + 1 + i
        ws.write_blank(r, 0, None, f["num_in"])
        ws.write_blank(r, 1, None, f["input"])
        ws.write_blank(r, 2, None, f["num_in"])
        ws.write_blank(r, 3, None, f["num_in"])
        ws.write_formula(r, 4, f'=IF(C{r + 1}="","",C{r + 1}*D{r + 1})', f["num"], "")
    tr = hr + 1 + rows
    ws.write(tr, 1, "Tổng cộng", f["total"])
    for c in (0, 2, 3):
        ws.write_blank(tr, c, None, f["total"])
    ws.write_formula(tr, 4, f"=SUM(E{hr + 2}:E{tr})", f["total"], 0)
    ws.write(tr + 1, 1, "VAT (%):", f["label"])
    ws.write_blank(tr + 1, 2, None, f["input"])
    ws.write(tr + 2, 1, "Tổng thanh toán", f["label"])
    ws.write_formula(tr + 2, 4, f"=E{tr + 1}*(1+N(C{tr + 2})/100)", f["total"], 0)
    ws.set_column(0, 0, 14)
    ws.set_column(1, 1, 30)
    ws.set_column(2, 4, 16)
    ws.freeze_panes(hr + 1, 0)
    wb.close()


def _budget(p: Path, rows: int = 15) -> None:
    wb = xlsxwriter.Workbook(str(p))
    f = _fmts(wb)
    ws = wb.add_worksheet("Ngân sách")
    ws.write(0, 0, "NGÂN SÁCH CHI PHÍ", f["title"])
    ws.write(2, 0, "Kỳ:", f["label"])
    ws.write_blank(2, 1, None, f["input"])
    hr = 4
    for c, h in enumerate(["Hạng mục", "Kế hoạch", "Thực tế", "Chênh lệch", "Ghi chú"]):
        ws.write(hr, c, h, f["head"])
    for i in range(rows):
        r = hr + 1 + i
        ws.write_blank(r, 0, None, f["input"])
        ws.write_blank(r, 1, None, f["num_in"])
        ws.write_blank(r, 2, None, f["num_in"])
        ws.write_formula(r, 3, f'=IF(AND(B{r + 1}="",C{r + 1}=""),"",C{r + 1}-B{r + 1})', f["num"], "")
        ws.write_blank(r, 4, None, f["input"])
    tr = hr + 1 + rows
    ws.write(tr, 0, "Tổng cộng", f["total"])
    for c, col in ((1, "B"), (2, "C"), (3, "D")):
        ws.write_formula(tr, c, f"=SUM({col}{hr + 2}:{col}{tr})", f["total"], 0)
    ws.write_blank(tr, 4, None, f["total"])
    ws.set_column(0, 0, 28)
    ws.set_column(1, 3, 16)
    ws.set_column(4, 4, 30)
    ws.freeze_panes(hr + 1, 0)
    wb.close()


BUILTINS = {
    "invoice_vi": ("Hóa đơn bán hàng", "Thông tin khách hàng + bảng sản phẩm, thành tiền và tổng có công thức.", _invoice),
    "budget_vi": ("Ngân sách chi phí", "Kế hoạch vs thực tế theo hạng mục, chênh lệch và tổng có công thức.", _budget),
}
